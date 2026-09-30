# -*- coding: utf-8 -*-
"""
Tela /clientes-agenda do painel (fila de autorizacao do AGENDA, 30/09).

    py -3.11 -m unittest painel_agentes.test_clientes_agenda_tela
"""
import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

_AQUI = Path(__file__).parent
_RAIZ = _AQUI.parent
for _p in (_RAIZ, _AQUI):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import marcar_clientes_agenda as mca  # noqa: E402

# "painel_agentes" e o nome da pasta: o modulo do app precisa ser carregado pelo caminho
import importlib.util  # noqa: E402
_spec = importlib.util.spec_from_file_location("painel_agentes_app", _AQUI / "painel_agentes.py")
painel_agentes = importlib.util.module_from_spec(_spec)
sys.modules[_spec.name] = painel_agentes  # o Flask acha a pasta templates/ pelo modulo registrado
_spec.loader.exec_module(painel_agentes)


class TestTelaClientesAgenda(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.addCleanup(self._tmp.cleanup)
        pasta = Path(self._tmp.name)
        self.planilha = pasta / "BD_CLIENTES.xlsx"
        banco = pasta / "dados.db"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Destinatário - Código", "Destinatário - Nome ", "ALERTAS CLIENTES"])
        ws.append([49749452000160, "VERO PANE", None])
        wb.save(self.planilha)

        conn = sqlite3.connect(banco)
        conn.execute("CREATE TABLE agendamentos_pedido (pedido TEXT, cnpj_destinatario TEXT, nome_destinatario TEXT, "
                     "cnpj_embarcador TEXT, status TEXT, data_agendada TEXT, origem TEXT)")
        conn.execute("INSERT INTO agendamentos_pedido VALUES ('PS-1', '49749452000160', 'VERO PANE', '1', "
                     "'RESPONDIDO', '30/09/2026', NULL)")
        conn.commit()
        conn.close()

        for alvo, valor in ((mca, banco),):
            p = mock.patch.object(alvo, "DB_PATH", valor)
            p.start()
            self.addCleanup(p.stop)
        config_real = painel_agentes._carregar_config()
        config = {**config_real, "clientes_agendamento": {"planilha": str(self.planilha)}}
        p = mock.patch.object(painel_agentes, "_carregar_config", return_value=config)
        p.start()
        self.addCleanup(p.stop)

        mca.processar(self.planilha, modo_teste=False)
        painel_agentes.app.config["TESTING"] = True
        self.cliente = painel_agentes.app.test_client()

    def _logar(self, nivel):
        with self.cliente.session_transaction() as sess:
            sess["nivel_acesso"] = nivel
            sess["usuario"] = "teste"

    def _alerta(self):
        wb = openpyxl.load_workbook(self.planilha, read_only=True, data_only=True)
        valor = list(wb.active.iter_rows(min_row=2, values_only=True))[0][2]
        wb.close()
        return valor

    def test_nivel_total_ve_o_pendente(self):
        self._logar("total")
        resp = self.cliente.get("/clientes-agenda")
        self.assertEqual(resp.status_code, 200)
        html = resp.get_data(as_text=True)
        self.assertIn("VERO PANE", html)
        self.assertIn("49.749.452/0001-60", html)
        self.assertIn("PS-1 · 30/09/2026 · portal ou e-mail", html)
        self.assertIn('data-acao="autorizar"', html)

    def test_outros_niveis_nao_entram(self):
        for nivel in ("operador", "leitura", "atendimento"):
            self._logar(nivel)
            self.assertEqual(self.cliente.get("/clientes-agenda").status_code, 403, nivel)

    def test_sem_login_vai_pro_login(self):
        resp = self.cliente.get("/clientes-agenda")
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/login", resp.headers["Location"])

    def test_autorizar_grava_na_planilha(self):
        self._logar("total")
        resp = self.cliente.post("/api/clientes-agenda/decidir", json={"documento": "49749452000160", "acao": "autorizar"},
                                 headers={"Origin": "http://localhost"})
        self.assertEqual(resp.status_code, 200, resp.get_data(as_text=True))
        self.assertEqual(resp.get_json()["situacao"], "AUTORIZADO")
        self.assertEqual(self._alerta(), "AGENDA")
        self.assertNotIn('data-acao="autorizar"', self.cliente.get("/clientes-agenda").get_data(as_text=True))

    def test_nao_marcar_nao_toca_na_planilha(self):
        self._logar("total")
        resp = self.cliente.post("/api/clientes-agenda/decidir", json={"documento": "49749452000160", "acao": "nao_marcar"},
                                 headers={"Origin": "http://localhost"})
        self.assertEqual(resp.status_code, 200)
        self.assertIsNone(self._alerta())
        # decidir de novo: ja nao esta pendente
        resp = self.cliente.post("/api/clientes-agenda/decidir", json={"documento": "49749452000160", "acao": "autorizar"},
                                 headers={"Origin": "http://localhost"})
        self.assertEqual(resp.status_code, 400)

    def test_acao_invalida_e_origem_estranha(self):
        self._logar("total")
        resp = self.cliente.post("/api/clientes-agenda/decidir", json={"documento": "49749452000160", "acao": "apagar"},
                                 headers={"Origin": "http://localhost"})
        self.assertEqual(resp.status_code, 400)
        resp = self.cliente.post("/api/clientes-agenda/decidir", json={"documento": "49749452000160", "acao": "autorizar"},
                                 headers={"Origin": "http://outro.site"})
        self.assertEqual(resp.status_code, 403)
        self.assertIsNone(self._alerta())


if __name__ == "__main__":
    unittest.main()
