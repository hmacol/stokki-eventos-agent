# -*- coding: utf-8 -*-
"""
Testes da calculadora de frete dedicado do portal do cliente (11/09/2026):
regras de cálculo, escolha de veículo, lista, geocodificação/km mockados,
proposta/aceite com e-mail mockado, PDF e rotas Flask.

    py -3 -m pytest portal_cliente/test_cotacao.py -q
"""
import json
import sys
from pathlib import Path
from unittest import mock

import pytest

_AQUI = Path(__file__).parent
_RAIZ = _AQUI.parent
for _p in (_RAIZ, _AQUI, _RAIZ / "roteirizacao"):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import cotacao as ct  # noqa: E402
import km_rodoviario  # noqa: E402

CNPJ = "12345678000195"
CLIENTE = {"cnpj": CNPJ, "sender_id": 1, "nome": "EMB TESTE", "cnpj_formatado": "12.345.678/0001-95", "versao": "x"}
CONFIG = {"google_maps": {"api_key": "chave"}, "email": {"remetente": "hugo@freshlogbr.com", "senha_app": "x"},
          "portal_cliente": {"cotacao": {"forcar_destino": ""}}}
REGRAS = ct.regras_de(CONFIG)

PARADA = {"cep": "01310-100", "numero": "1578", "logradouro": "Av. Paulista", "bairro": "Bela Vista",
          "cidade": "São Paulo", "uf": "SP", "referencia": "Mercado"}


# ── Cálculo puro ─────────────────────────────────────────────────────────────

def _calc(**kw):
    base = {"caixas": 80, "peso_kg": 300, "tipo_carga": "REFRIGERADO", "urgente": False, "valor_nf": None,
            "km_total": 40, "pedagio": None}
    base.update(kw)
    return ct.calcular(base, REGRAS)


def test_fiorino_dentro_da_franquia_gross_up():
    r = _calc()
    assert r["veiculo"] == "FIORINO"
    assert r["km_excedente"] == 0 and r["valor_km"] == 0
    assert r["frete"] == 650.0
    # gross-up: 650 / 0,88 = 738,64 -> imposto 88,64
    assert r["impostos"] == 88.64
    assert r["total"] == 738.64


def test_desconto_seco_e_km_arredonda_pra_cima():
    r = _calc(tipo_carga="SECO", km_total=72.4)
    assert r["desconto_seco"] == 150.0
    assert r["km_excedente"] == 8          # 7,4 km -> 8
    assert r["valor_km"] == 16.0
    assert r["frete"] == 516.0
    assert r["total"] == round(516.0 / 0.88, 2)


def test_km_exato_na_franquia_nao_cobra_excedente():
    assert ct.km_excedente(65.0, 65) == 0
    assert ct.km_excedente(65.01, 65) == 1
    assert ct.km_excedente(66.0, 65) == 1


def test_ad_valorem_e_urgencia():
    r = _calc(valor_nf=10000, urgente=True)
    assert r["ad_valorem"] == 50.0
    assert r["urgencia"] == round((650 + 50) * 0.40, 2)   # 280
    assert r["frete_liquido"] == 980.0
    assert r["total"] == round(980.0 / 0.88, 2)


def test_pedagio_fora_do_gross_up():
    a = _calc()
    b = _calc(pedagio=23.5)
    assert b["impostos"] == a["impostos"]
    assert b["total"] == round(a["total"] + 23.5, 2)


def test_pedagio_entra_na_linha_unica_do_frete():
    r = _calc(tipo_carga="SECO", km_total=72.4, pedagio=23.5)
    # frete base 650 - 150 (seco) + 16 (8 km) = 516, + pedágio 23,50
    assert r["frete_com_pedagio"] == 539.5


def test_composicao_do_cliente_junta_frete_desconto_km_e_pedagio():
    """Hugo, 11/09: o cliente vê UMA linha com frete base, desconto de carga
    seca, km adicional e pedágio já somados -- só ad valorem, urgência e
    impostos aparecem separados."""
    r = _calc(tipo_carga="SECO", km_total=72.4, pedagio=23.5, valor_nf=10000, urgente=True)
    r["km_volta"] = 30.0
    linhas = ct.linhas_composicao(r)
    rotulos = [k for k, _ in linhas]
    assert len(linhas) == 4
    assert rotulos[0].startswith("Frete dedicado Fiorino · 72,4 km (ida e volta)") and "pedágio incluso" in rotulos[0]
    assert linhas[0][1] == ct._fmt_brl(r["frete_com_pedagio"])
    assert "Ad valorem" in rotulos[1] and "Urgência same-day (+40%)" == rotulos[2] and "Impostos (12%)" == rotulos[3]
    # nenhuma linha entrega base, desconto ou km adicional separados
    assert not any("base" in k.lower() or "Desconto" in k or "Km adicional" in k for k in rotulos)
    # e a soma das linhas fecha o total
    soma = r["frete_com_pedagio"] + r["ad_valorem"] + r["urgencia"] + r["impostos"]
    assert round(soma, 2) == r["total"]


def test_composicao_detalhada_do_interno_abre_tudo():
    r = _calc(tipo_carga="SECO", km_total=72.4, pedagio=23.5)
    rotulos = [k for k, _ in ct.linhas_composicao(r, detalhado=True)]
    assert any("Frete base" in k for k in rotulos) and "Desconto carga seca" in rotulos
    assert any(k.startswith("Km adicional") for k in rotulos) and any(k.startswith("Pedágio") for k in rotulos)


def test_escolha_de_veiculo_por_caixas_e_peso():
    assert ct.escolher_veiculo(100, 550, REGRAS)["codigo"] == "FIORINO"
    assert ct.escolher_veiculo(101, 100, REGRAS)["codigo"] == "VAN_HR"     # caixas estouram a Fiorino
    assert ct.escolher_veiculo(10, 551, REGRAS)["codigo"] == "VAN_HR"      # peso estoura a Fiorino
    assert ct.escolher_veiculo(400, 1300, REGRAS)["codigo"] == "VAN_HR"
    assert ct.escolher_veiculo(401, 100, REGRAS)["codigo"] == "VUC"
    assert ct.escolher_veiculo(600, 2000, REGRAS)["codigo"] == "VUC"
    with pytest.raises(ct.ForaDaTabela):
        ct.escolher_veiculo(601, 100, REGRAS)
    with pytest.raises(ct.ForaDaTabela):
        ct.escolher_veiculo(10, 2001, REGRAS)
    with pytest.raises(ct.ErroCotacao):
        ct.escolher_veiculo(0, 0, REGRAS)


def test_vuc_franquia_120_e_km_250():
    r = _calc(caixas=500, peso_kg=1800, km_total=130)
    assert r["veiculo"] == "VUC"
    assert r["km_excedente"] == 10 and r["valor_km"] == 25.0
    assert r["frete"] == 1325.0


def test_config_sobrescreve_regras():
    cfg = {"portal_cliente": {"cotacao": {"impostos": 0.10, "veiculos": [
        {"codigo": "X", "nome": "X", "base": 100, "franquia_km": 10, "km_adicional": 1, "caixas_max": 5, "peso_max_kg": 5}]}}}
    regras = ct.regras_de(cfg)
    assert regras["impostos"] == 0.10 and regras["ad_valorem"] == 0.005
    assert [v["codigo"] for v in regras["veiculos"]] == ["X"]


# ── Entrada ───────────────────────────────────────────────────────────────────

def test_normalizar_paradas_exige_numero_e_cep_ou_endereco():
    with pytest.raises(ct.ErroCotacao):
        ct.normalizar_paradas([], REGRAS)
    with pytest.raises(ct.ErroCotacao, match="número"):
        ct.normalizar_paradas([{**PARADA, "numero": ""}], REGRAS)
    with pytest.raises(ct.ErroCotacao, match="CEP"):
        ct.normalizar_paradas([{"numero": "1"}], REGRAS)
    p = ct.normalizar_paradas([PARADA], REGRAS)[0]
    assert p["endereco"] == "Av. Paulista, 1578, Bela Vista, São Paulo - SP, 01310-100"


def test_normalizar_paradas_completa_pelo_cep(monkeypatch):
    monkeypatch.setattr(ct, "consultar_cep", lambda cep: {"cep": "01310-100", "logradouro": "Av. Paulista",
                                                          "bairro": "Bela Vista", "cidade": "São Paulo", "uf": "SP"})
    p = ct.normalizar_paradas([{"cep": "01310100", "numero": "10"}], REGRAS)[0]
    assert p["cidade"] == "São Paulo" and p["cep"] == "01310-100"


def test_ler_lista_xlsx_e_csv():
    import openpyxl, io
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["CEP", "Número", "Complemento", "Referência"])
    ws.append(["01310-100", 1578, "Loja 2", "Mercado"])
    ws.append([None, None, None, None])
    buf = io.BytesIO(); wb.save(buf)
    paradas = ct.ler_lista(buf.getvalue(), "lista.xlsx")
    assert paradas == [{"cep": "01310-100", "numero": "1578", "complemento": "Loja 2", "referencia": "Mercado"}]
    csv_txt = "cep;numero;destinatario\n04538-132;1000;Cliente A\n".encode("utf-8")
    paradas = ct.ler_lista(csv_txt, "lista.csv")
    assert paradas == [{"cep": "04538-132", "numero": "1000", "referencia": "Cliente A"}]
    with pytest.raises(ct.ErroCotacao, match="CEP"):
        ct.ler_lista("numero;x\n1;2\n".encode(), "a.csv")


def test_modelo_lista_abre():
    import openpyxl, io
    ws = openpyxl.load_workbook(io.BytesIO(ct.modelo_lista())).active
    assert ws.cell(1, 1).value == "CEP"


# ── Routes API com pedágio (mock do POST) ─────────────────────────────────────

def _fake_post(corpo, api_key, timeout, field_mask):
    assert "TOLLS" in corpo["extraComputations"] and "tollInfo" in field_mask
    n = len(corpo["intermediates"]) + 1
    return {"routes": [{"distanceMeters": 12000 * n, "legs": [{"distanceMeters": 12000} for _ in range(n)],
                        "travelAdvisory": {"tollInfo": {"estimatedPrice": [{"currencyCode": "BRL", "units": "12", "nanos": 500000000}]}}}]}


def test_calcular_trajeto_google_soma_pernas_e_pedagio():
    with mock.patch.object(km_rodoviario, "_post_com_mascara", side_effect=_fake_post):
        r = km_rodoviario.calcular_trajeto_google((0, 0), [(1, 1), (2, 2)], "k")
    assert r.km_total == 24.0 and r.km_ida == 24.0 and r.km_volta == 0.0
    assert r.pernas_km == (12.0, 12.0) and r.pedagio_brl == 12.5


def test_calcular_trajeto_google_com_retorno():
    """voltar=True acrescenta a perna de volta ao ponto de origem: ela entra no
    km_total (e no pedágio) mas NÃO na lista de pernas das entregas."""
    with mock.patch.object(km_rodoviario, "_post_com_mascara", side_effect=_fake_post):
        r = km_rodoviario.calcular_trajeto_google((0, 0), [(1, 1), (2, 2)], "k", voltar=True)
    assert r.pernas_km == (12.0, 12.0)        # só as 2 entregas
    assert r.km_ida == 24.0 and r.km_volta == 12.0 and r.km_total == 36.0


def test_cotacao_pede_o_trajeto_com_retorno():
    """A cotação SEMPRE manda voltar=True (config considerar_retorno)."""
    chamadas = []

    class _Falso:
        km_total, km_ida, km_volta, pernas_km, pedagio_brl, fonte = 90.0, 60.0, 30.0, (60.0,), 5.0, "GOOGLE_ROUTES"

        def como_dict(self):
            return {"km_total": self.km_total, "km_ida": self.km_ida, "km_volta": self.km_volta,
                    "pernas_km": list(self.pernas_km), "pedagio_brl": self.pedagio_brl, "fonte": self.fonte}

    def fake(origem, paradas, api_key, timeout=12, voltar=False):
        chamadas.append({"origem": origem, "voltar": voltar})
        return _Falso()

    with mock.patch.object(km_rodoviario, "calcular_trajeto", side_effect=fake):
        d = ct.calcular_trajeto([{"lat": -23.5, "lng": -46.6}], REGRAS, CONFIG)
    assert chamadas[0]["voltar"] is True
    assert tuple(chamadas[0]["origem"]) == tuple(REGRAS["origem_coords"])
    assert d["km_total"] == 90.0 and d["km_volta"] == 30.0


def test_calcular_trajeto_cai_em_linha_reta():
    with mock.patch.object(km_rodoviario, "_post_com_mascara", side_effect=RuntimeError("403")):
        r = km_rodoviario.calcular_trajeto((-23.5, -46.6), [(-23.55, -46.63)], "k", voltar=True)
    assert r.fonte == "HAVERSINE" and r.pedagio_brl is None
    assert r.km_volta > 0 and r.km_total == round(r.km_ida + r.km_volta, 2)


# ── Fluxo completo com banco temporário ───────────────────────────────────────

@pytest.fixture
def ambiente(tmp_path, monkeypatch):
    monkeypatch.setattr(ct, "DB_PATH", tmp_path / "dados.db")
    monkeypatch.setattr(ct, "PASTA_PDF", tmp_path / "cotacoes")
    import chamados
    monkeypatch.setattr(chamados, "DB_PATH", tmp_path / "dados.db")
    monkeypatch.setattr(ct, "resolver_coordenadas", lambda paradas, config: [p.update({"lat": -23.56, "lng": -46.65}) or p for p in paradas])
    def _trajeto(paradas, regras, config):
        """60 km até a 1ª entrega, 10 km entre as seguintes, 20 km de volta."""
        pernas = [60.0] + [10.0] * (len(paradas) - 1)
        ida, volta = round(sum(pernas), 2), 20.0
        return {"km_total": round(ida + volta, 2), "km_ida": ida, "km_volta": volta,
                "pernas_km": pernas, "pedagio_brl": 10.0, "fonte": "GOOGLE_ROUTES"}
    monkeypatch.setattr(ct, "calcular_trajeto", _trajeto)
    enviados = []
    monkeypatch.setattr(ct, "enviar_email", lambda dest, assunto, corpo, cfg, **kw: enviados.append({"para": dest, "assunto": assunto, "corpo": corpo, **kw}) or True)
    conn = ct.conectar()
    conn.execute("CREATE TABLE IF NOT EXISTS interno (cnpj_embarcador TEXT, apelido TEXT, nome_remetente TEXT, stkkc_id INTEGER, sender_id INTEGER, email TEXT, notificar_email INTEGER)")
    conn.execute("INSERT INTO interno VALUES (?, 'EMB TESTE', 'EMB TESTE', 98, 1, 'cliente@x.com', 1)", (CNPJ,))
    conn.commit()
    yield conn, enviados
    conn.close()


def _entrada(**kw):
    e = {"caixas": 50, "peso_kg": 200, "tipo_carga": "SECO", "urgente": False, "valor_nf": "1.000,00", "paradas": [PARADA]}
    e.update(kw)
    return e


def test_cotar_grava_e_calcula(ambiente):
    conn, _ = ambiente
    cot = ct.cotar(conn, CLIENTE, _entrada(), CONFIG, "cliente")
    assert cot["numero"] == "COT-00001" and cot["status"] == "SIMULADA"
    r = cot["resultado"]
    assert r["veiculo"] == "FIORINO" and r["km_total"] == 80.0 and r["km_excedente"] == 15
    # o retorno ao galpão (20 km) está DENTRO do km cobrado -- Hugo, 11/09
    assert r["km_ida"] == 60.0 and r["km_volta"] == 20.0
    assert r["valor_km"] == 30.0 and r["ad_valorem"] == 5.0 and r["pedagio"] == 10.0
    assert cot["entrada"]["valor_nf"] == 1000.0
    assert cot["entrada"]["paradas"][0]["km_perna"] == 60.0
    assert ct.listar(conn, CNPJ)[0]["id"] == cot["id"]
    assert ct.buscar(conn, cot["id"], "outro") is None


def test_cotar_fora_da_tabela_registra(ambiente):
    conn, _ = ambiente
    with pytest.raises(ct.ForaDaTabela, match="COT-00001"):
        ct.cotar(conn, CLIENTE, _entrada(caixas=700), CONFIG)
    assert ct.listar(conn, CNPJ)[0]["status"] == "FORA_DA_TABELA"


def test_proposta_aceite_e_pdf(ambiente):
    conn, enviados = ambiente
    cot = ct.cotar(conn, CLIENTE, _entrada(), CONFIG)
    cot = ct.enviar_proposta(conn, cot, ["a@b.com", "c@d.com"], CONFIG, "https://x/aceite/tok", "https://x/portal")
    assert cot["status"] == "PROPOSTA_ENVIADA" and cot["proposta_para"] == "a@b.com, c@d.com"
    assert enviados[-1]["para"] == ["a@b.com", "c@d.com"] and "https://x/aceite/tok" in enviados[-1]["corpo"]
    pdf = Path(_RAIZ / cot["pdf_path"])
    assert pdf.exists() and pdf.read_bytes()[:4] == b"%PDF"

    token = ct.gerar_token_aceite("segredo", cot)
    assert ct.validar_token_aceite("segredo", conn, token, REGRAS)["estado"] == "ok"
    assert ct.validar_token_aceite("outro", conn, token, REGRAS)["estado"] == "invalido"

    cot = ct.aceitar(conn, cot, "Fulano", "e-mail", CONFIG, ["cliente@x.com"])
    assert cot["status"] == "ACEITA" and cot["aceita_por"] == "Fulano" and cot["chamado_id"]
    assuntos = [e["assunto"] for e in enviados]
    assert any(a.startswith("[ACEITE]") for a in assuntos) and any("aprovada" in a for a in assuntos)
    assert ct.validar_token_aceite("segredo", conn, token, REGRAS)["estado"] == "aceita"
    # aceitar de novo é idempotente e não reenvia
    n = len(enviados)
    assert ct.aceitar(conn, cot, "Outro", "portal", CONFIG, [])["aceita_por"] == "Fulano"
    assert len(enviados) == n
    # chamado aberto no atendimento
    import chamados
    ch = chamados.buscar_chamado(conn, cot["chamado_id"])
    assert ch and "COT-00001" in ch["assunto"]


def test_email_da_proposta_so_total_e_botao(ambiente):
    """Hugo, 11/09: o e-mail é objetivo -- valor final e botão de aprovar.
    Entregas e composição só no PDF anexo."""
    conn, enviados = ambiente
    cot = ct.cotar(conn, CLIENTE, _entrada(valor_nf=10000), CONFIG)
    ct.enviar_proposta(conn, cot, ["a@b.com"], CONFIG, "https://x/aceite/tok", "")
    corpo = enviados[-1]["corpo"]
    assert ct._fmt_brl(cot["resultado"]["total"]) in corpo
    assert "Aprovar proposta" in corpo and "https://x/aceite/tok" in corpo
    for fora in ("Frete base", "Frete dedicado", "Desconto carga seca", "Km adicional", "Ad valorem",
                 "Impostos", "Pedágio", "Av. Paulista"):
        assert fora not in corpo, f"e-mail da proposta não deveria trazer '{fora}'"
    assert enviados[-1]["anexos"], "a proposta em PDF continua anexa"


def test_forcar_destino_redireciona(ambiente):
    conn, enviados = ambiente
    cfg = {**CONFIG, "portal_cliente": {"cotacao": {"forcar_destino": "hugo@freshlogbr.com"}}}
    cot = ct.cotar(conn, CLIENTE, _entrada(), cfg)
    ct.enviar_proposta(conn, cot, ["a@b.com"], cfg, "https://x", "")
    assert enviados[-1]["para"] == ["hugo@freshlogbr.com"]


def test_pdf_muitas_entregas_pagina(ambiente):
    conn, _ = ambiente
    cot = ct.cotar(conn, CLIENTE, _entrada(paradas=[{**PARADA, "referencia": f"Cliente {i}"} for i in range(25)]), CONFIG)
    import cotacao_pdf
    from pypdf import PdfReader
    import io
    pdf = cotacao_pdf.gerar(cot, REGRAS)
    assert len(PdfReader(io.BytesIO(pdf)).pages) >= 2


# ── Rotas Flask ───────────────────────────────────────────────────────────────

@pytest.fixture
def cliente_http(ambiente, monkeypatch):
    conn, enviados = ambiente
    import app as portal
    monkeypatch.setattr(portal.auth, "sessao_valida", lambda c, cnpj, v: dict(CLIENTE))
    monkeypatch.setattr(portal, "_CONFIG", CONFIG)
    import cotacao_web
    portal.app.config["TESTING"] = True
    with portal.app.test_client() as tc:
        with tc.session_transaction() as s:
            s["cnpj"] = CNPJ; s["v"] = "x"
        yield tc, conn, enviados


def test_rotas_calcular_proposta_aceite(cliente_http, monkeypatch):
    tc, conn, enviados = cliente_http
    import cotacao_web  # noqa
    r = tc.get("/cotacao")
    assert r.status_code == 200 and "Calcular frete" in r.get_data(as_text=True)

    r = tc.post("/api/cotacao/calcular", json=_entrada(), headers={"Origin": "http://localhost"})
    assert r.status_code == 200, r.get_data(as_text=True)
    cot = r.get_json()["cotacao"]
    assert cot["resultado"]["veiculo"] == "FIORINO"

    r = tc.post("/api/cotacao/calcular", json=_entrada(caixas=999), headers={"Origin": "http://localhost"})
    assert r.status_code == 422 and r.get_json()["fora_da_tabela"]

    r = tc.get("/api/cotacao/historico")
    assert len(r.get_json()["cotacoes"]) == 2

    r = tc.get(f"/cotacao/{cot['id']}/pdf")
    assert r.status_code == 200 and r.data[:4] == b"%PDF"

    r = tc.post(f"/api/cotacao/{cot['id']}/proposta", json={"emails": ["a@b.com"]}, headers={"Origin": "http://localhost"})
    assert r.status_code == 200 and r.get_json()["cotacao"]["status"] == "PROPOSTA_ENVIADA"
    link = [l for l in enviados[-1]["corpo"].split('"') if "/cotacao/aceite/" in l][0]
    token = link.rsplit("/", 1)[1]

    # aceite pelo link do e-mail, sem sessão
    with tc.session_transaction() as s:
        s.clear()
    r = tc.get(f"/cotacao/aceite/{token}")
    assert r.status_code == 200 and "Aprovar a proposta" in r.get_data(as_text=True)
    r = tc.post(f"/cotacao/aceite/{token}", data={"nome": "Fulano"})
    assert r.status_code == 200 and "aprovada" in r.get_data(as_text=True).lower()
    assert ct.buscar(conn, cot["id"])["status"] == "ACEITA"
    r = tc.get(f"/cotacao/aceite/{token}")
    assert "já aprovada" in r.get_data(as_text=True)
    r = tc.get("/cotacao/aceite/tokeninvalido")
    assert "Link inválido" in r.get_data(as_text=True)


def test_rotas_cep_e_lista(cliente_http, monkeypatch):
    tc, conn, _ = cliente_http
    monkeypatch.setattr(ct, "consultar_cep", lambda cep: {"cep": "01310-100", "logradouro": "Av. Paulista", "bairro": "Bela Vista", "cidade": "São Paulo", "uf": "SP"} if cep == "01310100" else None)
    assert tc.get("/api/cotacao/cep/01310-100").get_json()["cidade"] == "São Paulo"
    assert tc.get("/api/cotacao/cep/99999999").status_code == 404
    assert tc.get("/api/cotacao/cep/123").status_code == 400
    import io
    csv = io.BytesIO("CEP;Numero\n01310-100;10\n".encode())
    r = tc.post("/api/cotacao/lista", data={"arquivo": (csv, "l.csv")}, headers={"Origin": "http://localhost"})
    assert r.status_code == 200 and r.get_json()["paradas"][0]["cidade"] == "São Paulo"
    assert tc.get("/api/cotacao/modelo-lista").status_code == 200


def test_equipe_leitura_nao_aceita(cliente_http, monkeypatch):
    tc, conn, _ = cliente_http
    import app as portal
    r = tc.post("/api/cotacao/calcular", json=_entrada(), headers={"Origin": "http://localhost"})
    cot = r.get_json()["cotacao"]
    monkeypatch.setattr(portal, "_cliente_da_equipe", lambda cnpj: dict(CLIENTE, equipe=True))
    with tc.session_transaction() as s:
        s.clear(); s["equipe"] = {"usuario": "leitor", "nivel": "leitura"}; s["cnpj_equipe"] = CNPJ
    assert tc.post(f"/api/cotacao/{cot['id']}/aceitar", json={}, headers={"Origin": "http://localhost"}).status_code == 403
    assert tc.get("/cotacao").status_code == 200
