# -*- coding: utf-8 -*-
"""
portal_cliente/app.py

Portal do cliente (embarcador): acompanhamento do dia a dia dos pedidos
-- pedido do Hugo, 08/09. Publicado em app.freshhub.com.br/cliente pelo
Caddy (`handle_path /cliente/*`, ver infra/Caddyfile-cliente), app Flask
separado do painel interno (mesmo padrão de resposta_insucesso/app.py e
nucleo/api_motorista.py), reaproveitando o MESMO venv/config.yaml/
dados.db do repositório.

Acesso (decisão do Hugo, 08/09): CNPJ + PIN de 6 dígitos, com primeiro
acesso e "esqueci o PIN" por link no e-mail cadastrado na tabela
`interno` -- ver auth_cliente.py. Dados: dados_cliente.py (VUUPT
filtrada por sender_id + tabelas locais).

COMO RODAR (dev):
    py -3 portal_cliente/app.py            # http://127.0.0.1:8074
    (PORTAL_CLIENTE_DEV=1 desliga o cookie Secure pra testar em http)
COMO RODAR (produção -- ver infra/portal-cliente.service):
    waitress-serve --host=127.0.0.1 --port=8074 app:app
"""
import io
import logging
import os
import sys
from datetime import date, datetime, timedelta
from functools import wraps
from pathlib import Path
from urllib.parse import urlparse

_RAIZ = Path(__file__).parent.parent
_AQUI = Path(__file__).parent
# ORDEM IMPORTA (mesmo achado de resposta_insucesso/app.py): a raiz antes
# de insucesso_entrega/, que tem uma cópia velha de expedir_pedidos.py.
sys.path.insert(0, str(_RAIZ))
sys.path.insert(1, str(_RAIZ / "insucesso_entrega"))

import yaml
from flask import Flask, Response, abort, g, jsonify, redirect, render_template, request, send_file, session, url_for
from jinja2 import ChoiceLoader, FileSystemLoader
from werkzeug.middleware.proxy_fix import ProxyFix

import auth_cliente as auth
import dados_cliente as dados
import envio_pedidos as envios
from email_utils import enviar_email, envelope_html
from fingerprint_aguardando_resposta import buscar_pendentes_por_grupo
import aplicar_resposta_insucesso as logica_insucesso

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
logger = logging.getLogger("portal_cliente")


def _carregar_config() -> dict:
    with open(_RAIZ / "config.yaml", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


_CONFIG = _carregar_config()
_CFG = _CONFIG.get("portal_cliente", {}) or {}
_SECRET = _CFG.get("secret_key")
if not _SECRET:
    raise RuntimeError("portal_cliente.secret_key ausente no config.yaml -- gere uma string aleatória forte.")
_URL_BASE = (_CFG.get("url_base") or "https://app.freshhub.com.br/cliente").rstrip("/")
_DEV = os.environ.get("PORTAL_CLIENTE_DEV") == "1"

app = Flask(__name__)
app.wsgi_app = ProxyFix(app.wsgi_app, x_prefix=1, x_proto=1, x_for=1, x_host=1)
# Reaproveita o partial de paleta do painel (_paleta_cores.html) -- única
# fonte dos tokens de cor claro/escuro, pra não divergir do resto.
app.jinja_loader = ChoiceLoader([
    FileSystemLoader(str(_AQUI / "templates")),
    FileSystemLoader(str(_RAIZ / "painel_agentes" / "templates")),
])
app.secret_key = _SECRET
app.config.update(
    SESSION_COOKIE_NAME="portal_cliente",
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SECURE=not _DEV,
    SESSION_COOKIE_SAMESITE="Lax",
    PERMANENT_SESSION_LIFETIME=timedelta(days=14),
    # 64 KB bastava pros formulários; a máscara de envio recebe XMLs/ZIPs
    # de NF-e (08/09) -- um lote de 100 XMLs zipados fica na casa de 1-2 MB.
    MAX_CONTENT_LENGTH=30 * 1024 * 1024,
)

ACOES_VALIDAS = ("manter", "reagendar", "cancelar")
# Níveis do painel interno que podem entrar no portal em nome de um
# cliente (decisão do Hugo, 08/09, item 2: "os dois" sobem XML). Leitura
# só olha.
_NIVEIS_EQUIPE_ENVIA = ("total", "operador")


# ── Sessão ─────────────────────────────────────────────────────────────────────

def _cliente_da_equipe(cnpj: str) -> dict | None:
    conn = auth.conectar()
    try:
        emb = auth.buscar_embarcador(conn, cnpj)
    finally:
        conn.close()
    if not emb:
        return None
    return {"cnpj": emb["cnpj"], "cnpj_formatado": auth.formatar_cnpj(emb["cnpj"]), "sender_id": emb["sender_id"],
            "nome": emb["nome"], "versao": "equipe", "equipe": True}


@app.before_request
def _carregar_cliente():
    g.cliente = None
    g.equipe = session.get("equipe")
    if g.equipe:
        # Equipe Fresh Log logada (usuário do painel) escolhendo um cliente.
        cnpj = session.get("cnpj_equipe")
        g.cliente = _cliente_da_equipe(cnpj) if cnpj else None
        if cnpj and not g.cliente:
            session.pop("cnpj_equipe", None)
        return
    cnpj, versao = session.get("cnpj"), session.get("v")
    if cnpj and versao:
        conn = auth.conectar()
        try:
            g.cliente = auth.sessao_valida(conn, cnpj, versao)
        finally:
            conn.close()
        if not g.cliente:
            session.clear()


def requer_cliente(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not g.cliente:
            if request.path.startswith("/api/"):
                return jsonify({"erro": "Sessão expirada. Entre de novo."}), 401
            if g.get("equipe"):
                return redirect(url_for("equipe_cliente"))
            return redirect(url_for("login", proximo=request.script_root + request.full_path.rstrip("?")))
        return f(*args, **kwargs)
    return wrapper


def _quem_envia() -> str:
    """Carimbo de quem fez a ação (portal_envios.enviado_por)."""
    if g.get("equipe"):
        return f"equipe:{g.equipe.get('usuario', '?')}"
    return "cliente"


def _exige_pode_enviar():
    """Cliente sempre pode; equipe só nos níveis total/operador."""
    if g.get("equipe") and g.equipe.get("nivel") not in _NIVEIS_EQUIPE_ENVIA:
        abort(Response(jsonify({"erro": "Seu usuário só tem acesso de leitura."}).get_data(), status=403,
                       mimetype="application/json"))


def exige_mesma_origem(f):
    """Anti-CSRF por Origin/Referer nos POSTs (mesmo padrão do painel)."""
    @wraps(f)
    def wrapper(*args, **kwargs):
        origem = request.headers.get("Origin") or request.headers.get("Referer")
        if origem:
            host = urlparse(origem).netloc
            if host and host != request.host:
                abort(403)
        return f(*args, **kwargs)
    return wrapper


def _proximo_seguro(valor: str | None) -> str:
    """Só caminho relativo dentro do próprio app (anti open-redirect), SEMPRE
    com o prefixo do Caddy (/cliente) na frente -- redirect("/") sem o
    prefixo manda o navegador pra raiz do site (achado no 1º teste em
    produção, 08/09)."""
    raiz = request.script_root or ""
    if valor and valor.startswith("/") and not valor.startswith("//") and "\\" not in valor:
        if raiz and not (valor == raiz or valor.startswith(raiz + "/")):
            valor = raiz + valor
        return valor
    return url_for("inicio")


@app.context_processor
def _globais():
    return {"cliente": g.get("cliente"), "equipe": g.get("equipe"), "ano": date.today().year}


# ── Acesso da equipe Fresh Log (em nome de um cliente) ─────────────────────────

def _nivel_equipe(usuario: str, senha: str) -> str | None:
    """Mesmos pares usuário/senha do painel interno (painel_agentes.* no
    config.yaml) -- reaproveita as credenciais que a equipe já tem."""
    import hmac
    cfg = _carregar_config().get("painel_agentes", {}) or {}
    if not usuario or not senha:
        return None
    pares = (("usuario", "senha", "total"), ("usuario_operador", "senha_operador", "operador"),
             ("usuario_leitura", "senha_leitura", "leitura"))
    for ch_u, ch_s, nivel in pares:
        u, s = cfg.get(ch_u), cfg.get(ch_s)
        if u and s and hmac.compare_digest(usuario, str(u)) and hmac.compare_digest(senha, str(s)):
            return nivel
    return None


@app.route("/equipe", methods=["GET", "POST"])
def equipe_login():
    erro = None
    if request.method == "POST":
        usuario = (request.form.get("usuario") or "").strip()
        nivel = _nivel_equipe(usuario, request.form.get("senha") or "")
        if nivel is None:
            erro = "Usuário ou senha incorretos."
        else:
            session.clear()
            session["equipe"] = {"usuario": usuario, "nivel": nivel}
            session.permanent = True
            logger.info(f"equipe login usuario={usuario} nivel={nivel}")
            return redirect(url_for("equipe_cliente"))
    return render_template("equipe.html", erro=erro, modo="login")


@app.route("/equipe/cliente", methods=["GET", "POST"])
def equipe_cliente():
    if not g.get("equipe"):
        return redirect(url_for("equipe_login"))
    conn = auth.conectar()
    try:
        embarcadores = auth.listar_embarcadores(conn)
    finally:
        conn.close()
    if request.method == "POST":
        cnpj = auth.normalizar_cnpj(request.form.get("cnpj"))
        if _cliente_da_equipe(cnpj):
            session["cnpj_equipe"] = cnpj
            proximo = request.form.get("proximo") or ""
            return redirect(_proximo_seguro(proximo) if proximo else url_for("inicio"))
    return render_template("equipe.html", modo="cliente", embarcadores=embarcadores,
                           cnpj_atual=session.get("cnpj_equipe", ""))


# ── Login / acesso ─────────────────────────────────────────────────────────────

@app.route("/login", methods=["GET", "POST"])
def login():
    if g.cliente:
        return redirect(url_for("inicio"))
    erro = None
    cnpj_digitado = ""
    if request.method == "POST":
        cnpj_digitado = (request.form.get("cnpj") or "").strip()
        pin = (request.form.get("pin") or "").strip()
        conn = auth.conectar()
        try:
            conta = auth.autenticar(conn, cnpj_digitado, pin)
            session.clear()
            session["cnpj"] = conta["cnpj"]
            session["v"] = auth.versao_conta(conta)
            session.permanent = True
            logger.info(f"login ok cnpj={conta['cnpj']} sender_id={conta.get('sender_id')}")
            return redirect(_proximo_seguro(request.form.get("proximo")))
        except auth.AutenticacaoInvalida as e:
            erro = e.mensagem
        finally:
            conn.close()
    return render_template("login.html", erro=erro, cnpj=cnpj_digitado,
                           proximo=request.values.get("proximo", ""))


@app.route("/sair")
def sair():
    session.clear()
    return redirect(url_for("login"))


def _mascarar_email(e: str) -> str:
    usuario, _, dominio = e.partition("@")
    if len(usuario) <= 2:
        return usuario[:1] + "***@" + dominio
    return usuario[:2] + "***@" + dominio


@app.route("/primeiro-acesso", methods=["GET", "POST"])
def primeiro_acesso():
    """Primeiro acesso E 'esqueci o PIN': manda o link de definição de PIN
    pro(s) e-mail(s) da tabela interno."""
    erro = None
    cnpj_digitado = ""
    if request.method == "POST":
        cnpj_digitado = (request.form.get("cnpj") or "").strip()
        conn = auth.conectar()
        try:
            emb = auth.buscar_embarcador(conn, cnpj_digitado)
            if not emb:
                erro = "CNPJ não encontrado no nosso cadastro. Fale com a Fresh Log pra liberar o acesso."
            elif not emb["emails"]:
                erro = "Esse CNPJ não tem e-mail cadastrado. Fale com a Fresh Log pra liberar o acesso."
            else:
                token = auth.gerar_token_definir_pin(_SECRET, conn, emb["cnpj"])
                link = f"{_URL_BASE}/definir-pin/{token}"
                if _enviar_link_pin(emb, link):
                    return render_template("mensagem.html", titulo="Link enviado",
                                           texto="Enviamos um link pra você definir o PIN de acesso. "
                                                 "Ele vale por 24 horas.",
                                           detalhes=[_mascarar_email(e) for e in emb["emails"]],
                                           voltar=url_for("login"))
                erro = "Não conseguimos enviar o e-mail agora. Tente de novo em alguns minutos."
        finally:
            conn.close()
    return render_template("primeiro_acesso.html", erro=erro, cnpj=cnpj_digitado)


def _enviar_link_pin(emb: dict, link: str) -> bool:
    corpo = envelope_html(
        f"<p>Olá, <strong>{emb['nome']}</strong>.</p>"
        f"<p>Use o botão abaixo pra definir o PIN de acesso ao portal de acompanhamento de entregas "
        f"da Fresh Log (CNPJ {auth.formatar_cnpj(emb['cnpj'])}).</p>"
        f"<p style='margin:24px 0'><a href='{link}' style='background:#0EA575;color:#fff;padding:12px 22px;"
        f"border-radius:7px;text-decoration:none;font-weight:700'>Definir meu PIN</a></p>"
        f"<p style='color:#6B7280;font-size:13px'>O link vale por 24 horas e só pode ser usado uma vez. "
        f"Se não foi você que pediu, ignore este e-mail.</p>"
        f"<p style='color:#6B7280;font-size:12px'>Se o botão não abrir, copie este endereço: {link}</p>",
        rodape="Fresh Log · Portal de acompanhamento de entregas",
    )
    if _DEV:
        logger.info(f"[dev] link de PIN para {emb['emails']}: {link}")
    return enviar_email(emb["emails"], "Fresh Log · Defina seu PIN de acesso", corpo, _CONFIG.get("email", {}))


@app.route("/definir-pin/<token>", methods=["GET", "POST"])
def definir_pin(token):
    conn = auth.conectar()
    try:
        v = auth.validar_token_definir_pin(_SECRET, conn, token)
        if v["estado"] != "ok":
            textos = {
                "expirado": "Esse link já venceu (vale 24 horas). Peça um novo em “Primeiro acesso / esqueci o PIN”.",
                "usado": "Esse link já foi usado. Se precisar trocar o PIN de novo, peça um novo link.",
                "invalido": "Link inválido.",
            }
            return render_template("mensagem.html", titulo="Link não vale mais", texto=textos[v["estado"]],
                                   voltar=url_for("primeiro_acesso"), voltar_rotulo="Pedir novo link"), 410
        erro = None
        if request.method == "POST":
            pin, conf = (request.form.get("pin") or "").strip(), (request.form.get("pin2") or "").strip()
            try:
                auth.validar_pin_formato(pin)
                if pin != conf:
                    raise ValueError("Os dois PINs não conferem.")
                conta = auth.definir_pin(conn, v["cnpj"], pin)
                session.clear()
                session["cnpj"] = conta["cnpj"]
                session["v"] = auth.versao_conta(conta)
                session.permanent = True
                logger.info(f"PIN definido cnpj={conta['cnpj']}")
                return redirect(url_for("inicio"))
            except ValueError as e:
                erro = str(e)
        return render_template("definir_pin.html", embarcador=v["embarcador"],
                               cnpj_formatado=auth.formatar_cnpj(v["cnpj"]), erro=erro)
    finally:
        conn.close()


# ── Telas ──────────────────────────────────────────────────────────────────────

def _data_da_query() -> date:
    bruto = request.args.get("data", "")
    try:
        return date.fromisoformat(bruto) if bruto else date.today()
    except ValueError:
        return date.today()


@app.route("/")
@requer_cliente
def inicio():
    aba = request.args.get("aba") or "acompanhamento"
    return render_template("acompanhamento.html", data_inicial=_data_da_query().isoformat(),
                           hoje=date.today().isoformat(), aba_inicial=aba if aba in ("acompanhamento", "envios") else "acompanhamento",
                           pode_enviar=not (g.get("equipe") and g.equipe.get("nivel") not in _NIVEIS_EQUIPE_ENVIA))


@app.route("/saude")
def saude():
    return jsonify({"ok": True, "agora": datetime.now().isoformat(timespec="seconds")})


# ── API ────────────────────────────────────────────────────────────────────────

@app.route("/api/dia")
@requer_cliente
def api_dia():
    data_alvo = _data_da_query()
    forcar = request.args.get("atualizar") == "1"
    try:
        return jsonify(dados.montar_dia(g.cliente["sender_id"], data_alvo, _CONFIG, forcar=forcar))
    except Exception as e:
        logger.exception(f"api_dia falhou sender_id={g.cliente['sender_id']} data={data_alvo}")
        return jsonify({"erro": f"Não foi possível carregar os pedidos agora ({type(e).__name__}). Tente de novo."}), 502


@app.route("/api/canhoto/<int:service_id>")
@requer_cliente
def api_canhoto(service_id):
    token = _CONFIG.get("vuupt_api", {}).get("token", "")
    servico = dados.buscar_servico(token, service_id)
    if not servico or servico.get("sender_id") != g.cliente["sender_id"]:
        abort(404)
    checklist_id = dados.checklist_id_do_servico(servico)
    if not checklist_id:
        return render_template("mensagem.html", titulo="Comprovante ainda não disponível",
                               texto="O motorista ainda não enviou a foto do canhoto deste pedido, "
                                     "ou ela ainda está sendo processada. Tente de novo mais tarde.",
                               voltar=url_for("inicio")), 404
    pdf = dados.baixar_canhoto_pdf(token, checklist_id, servico.get("code") or str(service_id))
    if not pdf:
        return render_template("mensagem.html", titulo="Comprovante indisponível",
                               texto="Não conseguimos gerar o comprovante agora. Tente de novo em alguns minutos.",
                               voltar=url_for("inicio")), 502
    nome = f"comprovante_{(servico.get('code') or str(service_id)).lstrip('#')}.pdf"
    return send_file(pdf, mimetype="application/pdf", as_attachment=False, download_name=nome, max_age=0)


@app.route("/api/responder", methods=["POST"])
@requer_cliente
@exige_mesma_origem
def api_responder():
    """Decisão do embarcador sobre insucessos aguardando resposta -- mesmo
    caminho da página /insucesso (aplicar_decisao aplica por grupo
    sender_id + failed_reason_id)."""
    corpo = request.get_json(silent=True) or {}
    try:
        failed_reason_id = int(corpo.get("failed_reason_id"))
    except (TypeError, ValueError):
        return jsonify({"erro": "Motivo inválido."}), 400
    acao = corpo.get("acao")
    if acao not in ACOES_VALIDAS:
        return jsonify({"erro": "Ação inválida."}), 400

    sender_id = g.cliente["sender_id"]
    if not buscar_pendentes_por_grupo(sender_id, failed_reason_id):
        return jsonify({"erro": "Esses pedidos já foram respondidos.", "resolvido": True}), 409

    nova_data, hora = None, None
    if acao == "reagendar":
        try:
            nova_data = date.fromisoformat((corpo.get("data") or "").strip())
            if nova_data < date.today():
                raise ValueError
        except ValueError:
            return jsonify({"erro": "Informe uma data válida, a partir de hoje."}), 400
        hora = (corpo.get("hora") or "").strip()
        try:
            datetime.strptime(hora, "%H:%M")
        except ValueError:
            return jsonify({"erro": "Informe um horário válido (HH:MM)."}), 400

    try:
        resultados = logica_insucesso.aplicar_decisao(sender_id, failed_reason_id, acao, nova_data, _CONFIG,
                                                      hora_pedida=hora)
    except Exception as e:
        logger.exception("aplicar_decisao falhou")
        return jsonify({"erro": f"Não foi possível aplicar a resposta ({type(e).__name__})."}), 502
    logger.info(f"resposta insucesso cnpj={g.cliente['cnpj']} motivo={failed_reason_id} acao={acao} -> {resultados}")
    dados.invalidar_cache(sender_id)
    return jsonify({"ok": True, "resultados": [{**r, "data": r["data"].isoformat() if r.get("data") else None}
                                               for r in resultados]})


# ── Máscara de envio de pedidos (XML → fila → Stokki) ──────────────────────────
# Pedido do Hugo, 08/09: ver envio_pedidos.py (dados/validação) e
# enviar_stokki.py (worker que cria na Stokki respeitando a trava de sessão).

def _json_erro_envio(e: Exception, status: int = 400):
    return jsonify({"erro": str(e)}), status


@app.route("/api/envios")
@requer_cliente
def api_envios_listar():
    conn = envios.conectar()
    try:
        lista = envios.listar_envios(conn, g.cliente["cnpj"])
        cfg = envios.config_stokki_cliente(conn, g.cliente["cnpj"], _CONFIG)
    except envios.ErroEnvio as e:
        return _json_erro_envio(e)
    finally:
        conn.close()
    try:
        from stokki.sessao_uso import em_uso
        fila = em_uso()
    except Exception:
        fila = None
    return jsonify({"envios": lista, "resumo": envios.resumo_envios(lista), "dias": envios.DIAS_LISTAGEM,
                    "regra_xml": cfg["regra_xml"], "regra_xml_rotulo": envios.REGRAS_XML.get(cfg["regra_xml"], ""),
                    "envio_ativo": cfg["envio_ativo"], "client_id": cfg["client_id"], "fila": fila,
                    "atualizado_em": datetime.now().strftime("%H:%M")})


@app.route("/api/envios/analisar", methods=["POST"])
@requer_cliente
@exige_mesma_origem
def api_envios_analisar():
    """Recebe .xml/.zip, lê cada NF-e, valida (item 9) e devolve a prévia
    -- nada entra na fila ainda. Os XMLs válidos ficam num temporário
    identificado por token até o cliente confirmar."""
    _exige_pode_enviar()
    arquivos = request.files.getlist("arquivos")
    if not arquivos:
        return jsonify({"erro": "Selecione pelo menos um arquivo: XML da NF-e (ou ZIP com XMLs) ou planilha .xlsx no modelo Fresh Log."}), 400
    itens, rejeitados, vistos = [], [], set()
    n_planilhas = 0
    conn = envios.conectar()
    try:
        cfg = envios.config_stokki_cliente(conn, g.cliente["cnpj"], _CONFIG)
        for f in arquivos:
            try:
                partes = envios.expandir_upload(f.filename or "", f.read())
            except envios.ErroEnvio as e:
                rejeitados.append({"arquivo": f.filename or "arquivo", "erro": str(e)})
                continue
            for nome, conteudo in partes:
                if envios.e_planilha(nome, conteudo):
                    # Planilha no modelo Fresh Log (09/09): vários pedidos por
                    # arquivo; cada pedido vira um item da prévia com token
                    # próprio (JSON temporário) apontando pra planilha original.
                    n_planilhas += 1
                    try:
                        pedidos, rej = envios.ler_planilha(conteudo, nome, g.cliente["cnpj"])
                    except envios.ErroEnvio as e:
                        rejeitados.append({"arquivo": nome, "erro": str(e)})
                        continue
                    rejeitados.extend(rej)
                    token_plan = envios.guardar_temporario(conteudo, Path(nome).suffix.lstrip(".") or "xlsx")
                    for ped in pedidos:
                        rotulo = f"Pedido {ped['referencia']}"
                        if ped["chave_nfe"] in vistos:
                            rejeitados.append({"arquivo": nome, "rotulo": rotulo, "erro": "Repetido dentro do mesmo envio."})
                            continue
                        vistos.add(ped["chave_nfe"])
                        v = envios.validar_item(conn, ped, g.cliente["cnpj"])
                        erros_sku, avisos_sku = envios.validar_skus(conn, ped, cfg)
                        if not v["ok"] or erros_sku:
                            rejeitados.append({"arquivo": nome, "rotulo": rotulo, "erro": " ".join(v["erros"] + erros_sku)})
                            continue
                        token = envios.guardar_temporario_pedido(ped, token_plan)
                        item = {k: v_ for k, v_ in ped.items() if k not in ("emitente_cnpj",)}
                        item.update({"token": token, "avisos": v["avisos"] + avisos_sku, "rotulo": rotulo,
                                     "destinatario_doc_formatado": envios.formatar_documento(ped["destinatario_doc"])})
                        itens.append(item)
                    continue
                try:
                    nfe = envios.ler_nfe(conteudo, nome)
                except envios.ErroEnvio as e:
                    rejeitados.append({"arquivo": nome, "erro": str(e)})
                    continue
                nfe["origem"] = envios.ORIGEM_XML
                if nfe["chave_nfe"] in vistos:
                    rejeitados.append({"arquivo": nome, "numero_nf": nfe["numero_nf"], "rotulo": f"NF {nfe['numero_nf']}",
                                       "erro": "Repetida dentro do mesmo envio."})
                    continue
                vistos.add(nfe["chave_nfe"])
                v = envios.validar_item(conn, nfe, g.cliente["cnpj"])
                if not v["ok"]:
                    rejeitados.append({"arquivo": nome, "numero_nf": nfe["numero_nf"], "rotulo": f"NF {nfe['numero_nf']}",
                                       "erro": " ".join(v["erros"])})
                    continue
                token = envios.guardar_temporario(conteudo)
                item = {k: v_ for k, v_ in nfe.items() if k not in ("emitente_cnpj",)}
                item.update({"token": token, "avisos": v["avisos"], "rotulo": f"NF {nfe['numero_nf']}",
                             "destinatario_doc_formatado": envios.formatar_documento(nfe["destinatario_doc"])})
                itens.append(item)
        destinatarios = envios.info_destinatarios(conn, g.cliente["cnpj"], itens, _CONFIG)
    except envios.ErroEnvio as e:
        return _json_erro_envio(e)
    finally:
        conn.close()
    logger.info(f"analisar cnpj={g.cliente['cnpj']} por={_quem_envia()} validos={len(itens)} rejeitados={len(rejeitados)} planilhas={n_planilhas}")
    return jsonify({"itens": itens, "rejeitados": rejeitados, "destinatarios": destinatarios,
                    "regra_xml": cfg["regra_xml"], "regra_xml_rotulo": envios.REGRAS_XML.get(cfg["regra_xml"], ""),
                    "envio_ativo": cfg["envio_ativo"], "hoje": date.today().isoformat()})


@app.route("/api/envios/modelo-planilha")
@requer_cliente
def api_envios_modelo_planilha():
    """Modelo .xlsx da Fresh Log pra importação por planilha (09/09)."""
    conteudo = envios.gerar_modelo_planilha(g.cliente.get("nome") or "")
    return send_file(io.BytesIO(conteudo), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name="modelo_pedidos_freshlog.xlsx", max_age=0)


@app.route("/api/envios/confirmar", methods=["POST"])
@requer_cliente
@exige_mesma_origem
def api_envios_confirmar():
    _exige_pode_enviar()
    corpo = request.get_json(silent=True) or {}
    itens = corpo.get("itens") or []
    if not isinstance(itens, list) or not itens:
        return jsonify({"erro": "Nenhum pedido pra confirmar."}), 400
    if len(itens) > 200:
        return jsonify({"erro": "Envie no máximo 200 notas por vez."}), 400
    conn = envios.conectar()
    try:
        cfg = envios.config_stokki_cliente(conn, g.cliente["cnpj"], _CONFIG)
        criados = envios.confirmar_envios(conn, g.cliente["cnpj"], itens, _quem_envia(), _CONFIG, cfg["regra_xml"])
    except envios.ErroEnvio as e:
        return _json_erro_envio(e)
    except Exception as e:
        logger.exception("confirmar_envios falhou")
        return jsonify({"erro": f"Não foi possível registrar os pedidos ({type(e).__name__})."}), 500
    finally:
        conn.close()
    logger.info(f"confirmar cnpj={g.cliente['cnpj']} por={_quem_envia()} n={len(criados)}")
    n = len(criados)
    return jsonify({"ok": True, "criados": criados,
                    "mensagem": f"{n} pedido{'s' if n > 1 else ''} recebido{'s' if n > 1 else ''} com sucesso. "
                                f"A criação na Stokki pode levar alguns instantes -- acompanhe o status na lista abaixo."})


@app.route("/api/envios/<int:envio_id>/acao", methods=["POST"])
@requer_cliente
@exige_mesma_origem
def api_envios_acao(envio_id):
    """Cancelar / tirar da rota / reagendar / reenviar (item 12). O que
    não dá pra aplicar sozinho vira solicitação pra operação + e-mail."""
    _exige_pode_enviar()
    corpo = request.get_json(silent=True) or {}
    tipo = corpo.get("tipo")
    conn = envios.conectar()
    try:
        envio = envios.buscar_envio(conn, envio_id, g.cliente["cnpj"])
        if not envio:
            return jsonify({"erro": "Pedido não encontrado."}), 404
        resultado = envios.aplicar_acao(conn, envio, tipo, corpo, _quem_envia())
    except envios.ErroEnvio as e:
        return _json_erro_envio(e)
    finally:
        conn.close()
    logger.info(f"acao cnpj={g.cliente['cnpj']} envio={envio_id} tipo={tipo} por={_quem_envia()} -> {resultado}")
    if resultado.get("precisa_operacao"):
        _avisar_operacao_solicitacao(envio, tipo, corpo)
    return jsonify({"ok": True, **resultado})


def _avisar_operacao_solicitacao(envio: dict, tipo: str, corpo: dict) -> None:
    email_cfg = _CONFIG.get("email", {}) or {}
    destino = email_cfg.get("email_atendimento") or email_cfg.get("email_responsavel")
    if not destino:
        return
    rotulo = envios.ROTULOS_SOLICITACAO.get(tipo, tipo)
    detalhe = ""
    if tipo == "reagendar":
        detalhe = f"<p>Nova data: <b>{'/'.join(reversed(str(corpo.get('data') or '').split('-')))} "
        detalhe += f"{corpo.get('hora_inicio') or ''}–{corpo.get('hora_fim') or ''}</b></p>"
    if corpo.get("motivo"):
        detalhe += f"<p>Motivo informado: {corpo['motivo'][:300]}</p>"
    html = envelope_html(
        f"<p><b>{g.cliente['nome']}</b> (CNPJ {g.cliente['cnpj_formatado']}) pediu pelo portal: <b>{rotulo}</b>.</p>"
        f"<p>Pedido {envio.get('codigo_pedido') or '(código ainda não identificado)'} · NF {envio.get('numero_nf')} · "
        f"{envio.get('destinatario_nome')} · situação no portal: {envios.ROTULOS_STATUS.get(envio['status'], envio['status'])}.</p>"
        f"{detalhe}<p>Ao concluir, marque a solicitação como atendida: "
        f"<code>portal_cliente/gerenciar_clientes.py solicitacoes</code>.</p>",
        rodape="Fresh Log · Portal do cliente · solicitação", cor_acento="#F5A623")
    enviar_email([destino], f"[Portal] {rotulo} · NF {envio.get('numero_nf')} · {g.cliente['nome']}", html, email_cfg)


@app.route("/api/envios/<int:envio_id>/xml")
@app.route("/api/envios/<int:envio_id>/arquivo")
@requer_cliente
def api_envios_xml(envio_id):
    """Arquivo original do envio: o XML da NF-e ou a planilha de onde o
    pedido foi lido (09/09)."""
    conn = envios.conectar()
    try:
        envio = envios.buscar_envio(conn, envio_id, g.cliente["cnpj"])
    finally:
        conn.close()
    if not envio:
        abort(404)
    caminho = envios.caminho_xml(envio)
    if not caminho.is_file():
        abort(404)
    if (envio.get("origem") or envios.ORIGEM_XML) == envios.ORIGEM_PLANILHA:
        tipos = {".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xls": "application/vnd.ms-excel"}
        return send_file(caminho, mimetype=tipos.get(caminho.suffix.lower(), "application/octet-stream"), as_attachment=True,
                         download_name=caminho.name.split("_", 3)[-1] if caminho.name.count("_") >= 3 else caminho.name, max_age=0)
    return send_file(caminho, mimetype="application/xml", as_attachment=True, download_name=f"{envio['chave_nfe']}.xml", max_age=0)


@app.route("/api/destinatarios", methods=["POST"])
@requer_cliente
@exige_mesma_origem
def api_destinatarios_gravar():
    """Horário de recebimento (e flag de agendamento) de um destinatário --
    gravado uma vez por CNPJ/CPF e repassado à roteirização."""
    _exige_pode_enviar()
    corpo = request.get_json(silent=True) or {}
    conn = envios.conectar()
    try:
        envios.gravar_destinatario(conn, g.cliente["cnpj"], corpo.get("documento", ""), corpo.get("nome", ""),
                                   corpo.get("horario_inicio", ""), corpo.get("horario_fim", ""),
                                   bool(corpo.get("requer_agendamento")), _CONFIG)
    except envios.ErroEnvio as e:
        return _json_erro_envio(e)
    finally:
        conn.close()
    return jsonify({"ok": True})


# ── Exportação ─────────────────────────────────────────────────────────────────

_COLUNAS_XLSX = [
    ("codigo", "Pedido"), ("nf", "NF"), ("destinatario", "Destinatário"), ("endereco_completo", "Endereço"),
    ("volumes", "Volumes"), ("situacao_rotulo", "Situação"), ("motorista", "Motorista"), ("rota", "Rota"),
    ("ordem", "Parada"), ("detalhe", "Previsão / conclusão"), ("concluido_em", "Concluído às"),
    ("motivo", "Motivo do insucesso"), ("janela_atendimento", "Janela de atendimento"),
    ("telefone_destinatario", "Telefone do destinatário"), ("criado_em", "Data do pedido"), ("observacoes", "Observações"),
]


@app.route("/exportar.xlsx")
@requer_cliente
def exportar_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    data_alvo = _data_da_query()
    d = dados.montar_dia(g.cliente["sender_id"], data_alvo, _CONFIG)
    wb = Workbook()
    ws = wb.active
    ws.title = data_alvo.strftime("%d-%m-%Y")
    ws.append([rotulo for _, rotulo in _COLUNAS_XLSX])
    for c in ws[1]:
        c.font = Font(bold=True)
    for p in d["pedidos"] + d.get("agendados_futuros", []):
        ws.append([p.get(chave, "") if p.get(chave) is not None else "" for chave, _ in _COLUNAS_XLSX])
    for i, (chave, rotulo) in enumerate(_COLUNAS_XLSX, start=1):
        largura = max([len(rotulo)] + [len(str(p.get(chave) or "")) for p in d["pedidos"]] or [10])
        ws.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 10), 60)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nome = f"pedidos_{g.cliente['cnpj']}_{data_alvo.isoformat()}.xlsx"
    return send_file(buf, as_attachment=True, download_name=nome,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ── Atendimento (chat + chamados), 09/09 -- rotas em chamados_web.py ──────────
import chamados_web
chamados_web.registrar(app, requer_cliente=requer_cliente, exige_mesma_origem=exige_mesma_origem, config=_CONFIG)


@app.errorhandler(404)
def _404(_e):
    if request.path.startswith("/api/"):
        return jsonify({"erro": "Não encontrado."}), 404
    return render_template("mensagem.html", titulo="Página não encontrada", texto="", voltar=url_for("inicio")), 404


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(_CFG.get("porta") or 8074), debug=False)
