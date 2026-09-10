# -*- coding: utf-8 -*-
"""
Testes da importação de pedidos por planilha no portal do cliente
(pedido do Hugo, 09/09/2026): modelo, leitura/agrupamento, validações,
confirmação (fila) e lote simulado do worker.

    py -3 -m pytest portal_cliente/test_envio_planilha.py -q
"""
import io
import json
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest

_AQUI = Path(__file__).parent
_RAIZ = _AQUI.parent
for _p in (_RAIZ, _AQUI):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import envio_pedidos as ep  # noqa: E402

CNPJ = "12345678000195"
DEST = "98765432000110"
AMANHA = (date.today() + timedelta(days=1)).strftime("%d/%m/%Y")


def _planilha(linhas: list[dict], cabecalho: list[str] | None = None) -> bytes:
    """Monta um .xlsx com o cabeçalho do modelo (ou um custom) e as linhas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    cab = cabecalho or [c[1] for c in ep.COLUNAS_PLANILHA]
    chaves = [c[0] for c in ep.COLUNAS_PLANILHA]
    ws.append(cab)
    for l in linhas:
        ws.append([l.get(k, "") for k in chaves] if not cabecalho else [l.get(k, "") for k in cabecalho])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _linha(**kw) -> dict:
    base = {"referencia": "PED-1", "destinatario_doc": "98.765.432/0001-10", "destinatario_nome": "Mercado Teste",
            "cep": "01310-100", "logradouro": "Av. Paulista", "numero": "1578", "bairro": "Bela Vista", "cidade": "São Paulo",
            "uf": "SP", "sku": "SKU-A", "quantidade": 10, "valor_unitario": "12,50", "data_expedicao": AMANHA}
    base.update(kw)
    return base


@pytest.fixture
def ambiente(tmp_path, monkeypatch):
    monkeypatch.setattr(ep, "DB_PATH", tmp_path / "dados.db")
    monkeypatch.setattr(ep, "_RAIZ", tmp_path)   # xml_path é relativo à raiz
    monkeypatch.setattr(ep, "PASTA_XMLS", tmp_path / "portal_envios")
    monkeypatch.setattr(ep, "PASTA_TEMP", tmp_path / "portal_envios" / "_temporarios")
    # horário do destinatário: não escrever na tabela real da roteirização
    import regras.complexidade_entrega as ce
    monkeypatch.setattr(ce, "definir_ajuste_manual", lambda *a, **k: None)
    monkeypatch.setattr(ce, "carregar_ajustes_manuais", lambda: {})
    conn = ep.conectar()
    conn.execute("CREATE TABLE IF NOT EXISTS interno (cnpj_embarcador TEXT, apelido TEXT, nome_remetente TEXT, stkkc_id INTEGER, "
                 "sender_id INTEGER, email TEXT, notificar_email INTEGER)")
    conn.execute("INSERT INTO interno VALUES (?, 'EMB TESTE', 'EMB TESTE', 98, 1, 'x@y.z', 1)", (CNPJ,))
    conn.commit()
    yield conn
    conn.close()


# ── Modelo ─────────────────────────────────────────────────────────────────────

def test_modelo_tem_todas_as_colunas_e_e_lido_de_volta():
    conteudo = ep.gerar_modelo_planilha("Cliente X")
    wb = openpyxl.load_workbook(io.BytesIO(conteudo))
    assert wb.sheetnames == ["Pedidos", "Instruções"]
    cab = [c.value for c in wb["Pedidos"][1]]
    assert cab == [c[1] for c in ep.COLUNAS_PLANILHA]
    # as linhas de exemplo do modelo são lidas como 2 pedidos válidos
    pedidos, rejeitados = ep.ler_planilha(conteudo, "modelo.xlsx", CNPJ)
    assert rejeitados == []
    assert [p["referencia"] for p in pedidos] == ["PED-1001", "PED-1002"]
    assert pedidos[0]["itens"] == 2 and pedidos[1]["itens"] == 1


# ── Leitura ────────────────────────────────────────────────────────────────────

def test_agrupa_linhas_do_mesmo_pedido_e_calcula_valor():
    conteudo = _planilha([_linha(), _linha(sku="SKU-B", quantidade=2, valor_unitario=3, destinatario_doc="", destinatario_nome=""),
                          _linha(referencia="PED-2", destinatario_doc="123.456.789-09", destinatario_nome="Ana", sku="SKU-C", quantidade="1,5")])
    pedidos, rejeitados = ep.ler_planilha(conteudo, "p.xlsx", CNPJ)
    assert rejeitados == []
    assert len(pedidos) == 2
    p1 = pedidos[0]
    assert p1["origem"] == "planilha" and p1["referencia"] == "PED-1"
    assert p1["destinatario_doc"] == DEST and p1["destinatario_cep"] == "01310100"
    assert p1["destinatario_endereco"] == "Av. Paulista, 1578"
    assert [i["sku"] for i in p1["itens_lista"]] == ["SKU-A", "SKU-B"]
    assert p1["itens"] == 2 and p1["linhas"] == [2, 3]
    assert p1["valor_nf"] == 10 * 12.5 + 2 * 3
    assert p1["volumes"] == 1
    assert p1["data_expedicao"] == (date.today() + timedelta(days=1)).isoformat()
    assert p1["emitente_cnpj"] == CNPJ
    assert p1["chave_nfe"].startswith(f"PLANILHA-{CNPJ}-PED-1-")
    p2 = pedidos[1]
    assert p2["destinatario_doc"] == "12345678909" and p2["itens_lista"][0]["quantidade"] == 1.5


def test_cabecalho_com_apelidos_e_ordem_diferente():
    cab = ["CNPJ", "Cliente", "Pedido", "Endereço", "Nº", "Bairro", "Cidade", "Estado", "CEP", "Código", "Qtd"]
    linhas = [{"CNPJ": "98765432000110", "Cliente": "Loja", "Pedido": 77, "Endereço": "Rua A", "Nº": 5, "Bairro": "Centro",
               "Cidade": "Campinas", "Estado": "sp", "CEP": 13010000, "Código": 7891234567890, "Qtd": 3}]
    pedidos, rejeitados = ep.ler_planilha(_planilha(linhas, cab), "x.xlsx", CNPJ)
    assert rejeitados == []
    assert pedidos[0]["referencia"] == "77" and pedidos[0]["destinatario_uf"] == "SP"
    assert pedidos[0]["destinatario_cep"] == "13010000"
    assert pedidos[0]["itens_lista"][0]["sku"] == "7891234567890"
    assert pedidos[0]["data_expedicao"] == date.today().isoformat()   # em branco = hoje


def test_rejeita_pedido_com_problema_mas_mantem_os_outros():
    linhas = [
        _linha(),
        _linha(referencia="PED-2", destinatario_doc="123"),                     # doc inválido
        _linha(referencia="PED-3", quantidade=0),                                # sem item válido
        _linha(referencia="PED-4", uf="XX", cep="123"),                          # uf + cep
        _linha(referencia="PED-5", data_expedicao="01/01/2020"),                 # data passada
        _linha(referencia="PED-6", destinatario_doc="11.111.111/0001-11"),
        _linha(referencia="PED-6", destinatario_doc="22.222.222/0001-22", sku="B"),  # destinatário diferente no mesmo pedido
        {"sku": "X", "quantidade": 1},                                           # sem referência
    ]
    pedidos, rejeitados = ep.ler_planilha(_planilha(linhas), "p.xlsx", CNPJ)
    assert [p["referencia"] for p in pedidos] == ["PED-1"]
    por = {r["rotulo"]: r["erro"] for r in rejeitados}
    assert "14 ou 11 dígitos" in por["Pedido PED-2"]
    assert "maior que zero" in por["Pedido PED-3"]
    assert "UF inválida" in por["Pedido PED-4"] and "CEP inválido" in por["Pedido PED-4"]
    assert "já passou" in por["Pedido PED-5"]
    assert "destinatário diferente" in por["Pedido PED-6"]
    assert "linha 9" in por and "referência" in por["linha 9"]


def test_falta_coluna_obrigatoria_e_arquivo_nao_planilha():
    cab = ["Pedido", "CNPJ", "Cliente", "SKU", "Qtd"]
    with pytest.raises(ep.ErroEnvio, match="faltam colunas obrigatórias"):
        ep.ler_planilha(_planilha([{"Pedido": 1}], cab), "x.xlsx", CNPJ)
    with pytest.raises(ep.ErroEnvio):
        ep.ler_planilha(b"isso nao e xlsx", "x.xlsx", CNPJ)
    assert ep.e_planilha("a.XLSX") and ep.e_planilha("b.xls") and not ep.e_planilha("c.xml", b"<xml/>")


def test_expandir_upload_nao_confunde_xlsx_com_zip():
    conteudo = _planilha([_linha()])
    assert conteudo[:2] == b"PK"
    partes = ep.expandir_upload("pedidos.xlsx", conteudo)
    assert partes == [("pedidos.xlsx", conteudo)]


# ── Catálogo de SKUs ───────────────────────────────────────────────────────────

def test_validar_skus_erro_so_quando_catalogo_do_embarcador_existe(ambiente):
    conn = ambiente
    cfg = {"client_id": "98", "nome": "EMB TESTE"}
    pedido = {"itens_lista": [{"sku": "SKU-A"}, {"sku": "sku-z"}]}
    erros, avisos = ep.validar_skus(conn, pedido, cfg)
    assert not erros and avisos                         # sem tabela: só aviso
    conn.execute("CREATE TABLE wms_produtos (sku TEXT, ean TEXT, dun TEXT, descricao TEXT, embarcador TEXT, embarcador_id INTEGER, ativo INTEGER)")
    conn.execute("INSERT INTO wms_produtos VALUES ('SKU-A', '789', NULL, 'Produto A', 'EMB TESTE', 98, 1)")
    conn.execute("INSERT INTO wms_produtos VALUES ('SKU-Z', NULL, NULL, 'Produto Z', 'OUTRO', 5, 1)")
    conn.commit()
    erros, avisos = ep.validar_skus(conn, pedido, cfg)
    assert erros and "sku-z" in erros[0]
    ok, _ = ep.validar_skus(conn, {"itens_lista": [{"sku": "789"}]}, cfg)   # EAN também vale
    assert ok == []


# ── Confirmação (fila) ─────────────────────────────────────────────────────────

def _analisar_e_confirmar(conn, conteudo, nome="p.xlsx", **item_extra):
    pedidos, rejeitados = ep.ler_planilha(conteudo, nome, CNPJ)
    assert rejeitados == []
    token_plan = ep.guardar_temporario(conteudo, "xlsx")
    itens = []
    for ped in pedidos:
        v = ep.validar_item(conn, ped, CNPJ)
        assert v["ok"], v["erros"]
        tk = ep.guardar_temporario_pedido(ped, token_plan)
        itens.append({"token": tk, "horario_inicio": "08:00", "horario_fim": "17:00", **item_extra})
    return ep.confirmar_envios(conn, CNPJ, itens, "cliente", {}, "nenhuma")


def test_confirmar_grava_na_fila_com_itens_e_guarda_planilha(ambiente):
    conn = ambiente
    conteudo = _planilha([_linha(), _linha(sku="SKU-B", quantidade=2), _linha(referencia="PED-2", sku="C", quantidade=1)])
    criados = _analisar_e_confirmar(conn, conteudo, observacoes="Deixar na portaria")
    assert len(criados) == 2 and all(c["origem"] == "planilha" for c in criados)
    rows = [dict(r) for r in conn.execute("SELECT * FROM portal_envios ORDER BY id")]
    assert [r["referencia"] for r in rows] == ["PED-1", "PED-2"]
    assert all(r["status"] == "NA_FILA" and r["origem"] == "planilha" and r["regra_xml"] is None for r in rows)
    itens = json.loads(rows[0]["itens_json"])
    assert [i["sku"] for i in itens["itens"]] == ["SKU-A", "SKU-B"]
    assert itens["destinatario"]["destinatario_logradouro"] == "Av. Paulista"
    assert rows[0]["linhas_planilha"] == "2,3" and rows[0]["observacoes"] == "Deixar na portaria"
    # os 2 pedidos apontam pra MESMA planilha guardada, e o temporário sumiu
    assert rows[0]["xml_path"] == rows[1]["xml_path"] and rows[0]["xml_path"].endswith("_p.xlsx")
    assert (ep._RAIZ / rows[0]["xml_path"]).is_file()
    assert not list(ep.PASTA_TEMP.glob("*"))
    # listagem: rótulo e chip
    lista = ep.listar_envios(conn, CNPJ)
    assert lista[0]["rotulo"] == "Pedido PED-2" and lista[0]["origem"] == "planilha" and lista[0]["arquivo_ext"] == "xlsx"
    assert "itens_json" not in lista[0]
    # reenviar o mesmo pedido = duplicado
    ped, _ = ep.ler_planilha(_planilha([_linha()]), "p2.xlsx", CNPJ)
    v = ep.validar_item(conn, ped[0], CNPJ)
    assert not v["ok"] and "PED-1 já foi enviado" in v["erros"][0]


def test_caminho_xml_continua_igual(ambiente):
    """Regressão: NF-e pelo fluxo original ainda grava origem='xml', o XML em
    <cnpj>/<chave>.xml e rótulo 'NF n'."""
    conn = ambiente
    chave = "35260912345678000195550010000123451000012345"
    xml = f"""<?xml version="1.0"?><nfeProc xmlns="http://www.portalfiscal.inf.br/nfe"><NFe><infNFe Id="NFe{chave}">
      <ide><nNF>12345</nNF><serie>1</serie><tpNF>1</tpNF><dhEmi>2026-09-09T10:00:00-03:00</dhEmi></ide>
      <emit><CNPJ>{CNPJ}</CNPJ><xNome>Emitente</xNome></emit>
      <dest><CNPJ>{DEST}</CNPJ><xNome>Destino XML</xNome><enderDest><xLgr>Rua A</xLgr><nro>1</nro><xBairro>B</xBairro>
      <xMun>São Paulo</xMun><UF>SP</UF><CEP>01310100</CEP></enderDest></dest>
      <det nItem="1"><prod><cProd>X</cProd></prod></det>
      <transp><vol><qVol>2</qVol><pesoB>3.5</pesoB></vol></transp><total><ICMSTot><vNF>99.90</vNF></ICMSTot></total>
    </infNFe></NFe></nfeProc>"""
    nfe = ep.ler_nfe(xml.encode("utf-8"), "nota.xml")
    assert nfe["chave_nfe"] == chave and nfe["volumes"] == 2
    tk = ep.guardar_temporario(xml.encode("utf-8"))
    criados = ep.confirmar_envios(conn, CNPJ, [{"token": tk, "horario_inicio": "08:00", "horario_fim": "12:00"}], "cliente", {}, "muai")
    assert criados[0]["origem"] == "xml" and criados[0]["numero_nf"] == "12345"
    r = dict(conn.execute("SELECT * FROM portal_envios").fetchone())
    assert r["origem"] == "xml" and r["regra_xml"] == "muai" and r["itens_json"] is None and r["referencia"] is None
    assert r["xml_path"].endswith(f"{chave}.xml") and r["valor_nf"] == 99.9
    lista = ep.listar_envios(conn, CNPJ)
    assert lista[0]["rotulo"] == "NF 12345" and lista[0]["arquivo_ext"] == "xml" and lista[0]["documento_curto"] == "12345"


def test_xlsx_pedido_stokki_formato_do_modelo():
    conteudo = ep.xlsx_pedido_stokki([{"sku": "A", "quantidade": 3, "valor_unitario": 1.5}, {"sku": "B", "quantidade": 2.5, "valor_unitario": None}])
    ws = openpyxl.load_workbook(io.BytesIO(conteudo)).active
    assert [c.value for c in ws[1]] == ["SKU", "Quantidade", "Valor Unitário"]
    assert [c.value for c in ws[2]] == ["A", 3, 1.5]
    assert [c.value for c in ws[3]] == ["B", 2.5, None]


# ── Worker (simulado) ──────────────────────────────────────────────────────────

def test_worker_simulado_marca_planilha_como_criada(ambiente, monkeypatch, tmp_path):
    conn = ambiente
    import enviar_stokki as worker
    monkeypatch.setattr(worker, "PASTA_LOTES", tmp_path / "lotes")
    monkeypatch.setattr(worker, "_avisar_erros", lambda *a, **k: None)
    _analisar_e_confirmar(conn, _planilha([_linha(), _linha(referencia="PED-2", sku="C", quantidade=1)]))
    envios = [dict(r) for r in conn.execute("SELECT * FROM portal_envios WHERE status = 'NA_FILA'")]
    resumo = worker.processar_lote(conn, CNPJ, envios, {"portal_cliente": {}}, simular=True)
    assert resumo["criados"] == 2 and resumo["erros"] == 0
    assert [r[0] for r in conn.execute("SELECT status FROM portal_envios")] == ["CRIADO", "CRIADO"]
    gerados = sorted((tmp_path / "lotes").rglob("*.xlsx"))
    assert len(gerados) == 2
    ws = openpyxl.load_workbook(gerados[0]).active
    assert [c.value for c in ws[1]] == ["SKU", "Quantidade", "Valor Unitário"]


def test_worker_planilha_sem_itens_vira_erro(ambiente, monkeypatch, tmp_path):
    conn = ambiente
    import enviar_stokki as worker
    monkeypatch.setattr(worker, "PASTA_LOTES", tmp_path / "lotes")
    monkeypatch.setattr(worker, "_avisar_erros", lambda *a, **k: None)
    _analisar_e_confirmar(conn, _planilha([_linha()]))
    conn.execute("UPDATE portal_envios SET itens_json = '{}'")
    conn.commit()
    envios = [dict(r) for r in conn.execute("SELECT * FROM portal_envios")]
    resumo = worker.processar_lote(conn, CNPJ, envios, {"portal_cliente": {}}, simular=True)
    assert resumo["erros"] == 1
    r = conn.execute("SELECT status, erro FROM portal_envios").fetchone()
    assert r["status"] == "ERRO" and "sem itens" in r["erro"]
