# -*- coding: utf-8 -*-
"""
portal_cliente/cotacao.py

Calculadora de frete DEDICADO do portal do cliente -- pedido do Hugo,
11/09/2026. Regras (decisões do Hugo, 11/09):

  Veículo escolhido AUTOMATICAMENTE pelo menor que comporta caixas + peso:
    Fiorino   R$ 650  até  65 km, R$ 2,00/km adicional   (até 1 pallet,  550 kg)
    Van / HR  R$ 850  até 100 km, R$ 2,00/km adicional   (até 2 pallets, 1.300 kg)
    VUC       R$ 1.300 até 120 km, R$ 2,50/km adicional  (até 4 pallets, 2.000 kg)
    acima disso -> "solicite cotação" (não calcula).
  Carro SECO sai R$ 150 mais barato que refrigerado/congelado, em qualquer veículo.
  Km: rodoviário (Google Routes), galpão -> entregas na ordem informada -> galpão
      (o RETORNO sempre conta); excedente sobre a franquia SEMPRE em km inteiro
      arredondado pra cima.
  Ad valorem: 0,5% do valor da(s) NF (campo opcional -- sem NF, não entra).
  Urgência same-day: +40% sobre o frete (frete + ad valorem).
  Impostos 12% por GROSS-UP: valor final = líquido / 0,88.
  Pedágio: repassado pelo valor REAL; aqui entra a ESTIMATIVA da Routes API
      do MESMO trajeto do km -- ida e volta ao galpão --, fora do gross-up
      (não é receita de frete).
  Proposta vale 7 dias. Cotações ficam registradas (portal_cotacoes).

O que o CLIENTE vê (Hugo, 11/09): frete base, desconto de carga seca, km
adicional e pedágio numa ÚNICA linha já somada (frete_com_pedagio); só ad
valorem, urgência e impostos ficam separados -- ver linhas_composicao().
O e-mail da proposta é só o TOTAL + botão de aprovar; o detalhamento vai no
PDF anexo. O e-mail interno do aceite abre a composição toda (detalhado=True).

Tudo que é número (tabela de veículos, percentuais, desconto, validade,
endereço de origem) tem default aqui e pode ser sobrescrito no config.yaml,
seção portal_cliente.cotacao -- ver REGRAS_PADRAO.

Onde entra o que:
    calcular(entrada, regras)            -> só a conta (puro, testável)
    escolher_veiculo(caixas, peso, regras)
    cotar(conn, cliente, entrada, config, quem)  -> geocodifica, km, conta, grava
    enviar_proposta(...) / aceitar(...)  -> e-mail com PDF + botão de aceite
    consultar_cep(cep)                   -> ViaCEP (BrasilAPI de reserva)
"""
import csv
import io
import json
import logging
import math
import re
import sqlite3
import sys
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path

import requests
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer

_AQUI = Path(__file__).parent
_RAIZ = _AQUI.parent
for _p in (_RAIZ, _RAIZ / "roteirizacao", _AQUI):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

from email_utils import envelope_html, enviar_email  # noqa: E402

logger = logging.getLogger("portal_cliente.cotacao")

DB_PATH = _RAIZ / "dados" / "dados.db"
PASTA_PDF = _RAIZ / "dados" / "cotacoes"

# ── Regras (defaults; config.yaml portal_cliente.cotacao sobrescreve) ──────────

VEICULOS_PADRAO = [
    # caixas_max: o roteirizador trabalha em caixas (regras/tipo_veiculo.py:
    # VAN/HR 400, VUC 600); Fiorino não existe lá -- 100 é do Hugo (11/09),
    # com a ressalva de que varia MUITO com o tamanho da caixa. Ajustar em
    # portal_cliente.cotacao.veiculos.
    {"codigo": "FIORINO", "nome": "Fiorino", "base": 650.0, "franquia_km": 65, "km_adicional": 2.0,
     "caixas_max": 100, "peso_max_kg": 550, "pallets": 1},
    {"codigo": "VAN_HR", "nome": "Van / HR", "base": 850.0, "franquia_km": 100, "km_adicional": 2.0,
     "caixas_max": 400, "peso_max_kg": 1300, "pallets": 2},
    {"codigo": "VUC", "nome": "VUC", "base": 1300.0, "franquia_km": 120, "km_adicional": 2.5,
     "caixas_max": 600, "peso_max_kg": 2000, "pallets": 4},
]

REGRAS_PADRAO = {
    "impostos": 0.12,                 # gross-up: total = líquido / (1 - impostos)
    "ad_valorem": 0.005,              # sobre o valor da NF
    "desconto_seco": 150.0,           # R$ a menos na saída pra carga seca
    "urgencia_same_day": 0.40,        # +40% sobre o frete (frete + ad valorem)
    "validade_dias": 7,
    "considerar_retorno": True,       # km cobrado = galpão -> entregas -> galpão (Hugo, 11/09)
    "origem_endereco": "Rua Zilda, 288, Casa Verde Alta, São Paulo - SP",
    "origem_coords": [-23.497039, -46.6605211],   # regras/km_cobrado.COORD_CD
    "email_comercial": "entregas@freshlogbr.com",  # recebe o aceite
    "forcar_destino": "",             # piloto: se preenchido, TODO e-mail vai só pra cá
    "abrir_chamado_no_aceite": True,
    "max_paradas": 25,
    # Hugo, 17/09: calculadora oculta do cliente (link some, rotas dão 404).
    # A equipe Fresh Log operando em nome do cliente continua vendo. True reabre.
    "visivel_cliente": False,
    "veiculos": VEICULOS_PADRAO,
}

TIPOS_CARGA = {"SECO": "Seco", "REFRIGERADO": "Refrigerado", "CONGELADO": "Congelado"}

STATUS_SIMULADA = "SIMULADA"
STATUS_PROPOSTA_ENVIADA = "PROPOSTA_ENVIADA"
STATUS_ACEITA = "ACEITA"
STATUS_FORA_DA_TABELA = "FORA_DA_TABELA"
STATUS_ROTULO = {STATUS_SIMULADA: "Simulada", STATUS_PROPOSTA_ENVIADA: "Proposta enviada",
                 STATUS_ACEITA: "Aceita", STATUS_FORA_DA_TABELA: "Fora da tabela"}

_SALT_ACEITE = "portal-cliente-cotacao-aceite"


class ErroCotacao(Exception):
    """Erro de entrada/regra com mensagem pronta pro cliente."""


class ForaDaTabela(ErroCotacao):
    """Carga maior que o maior veículo da tabela: solicite cotação."""


def regras_de(config: dict) -> dict:
    """REGRAS_PADRAO + o que estiver em portal_cliente.cotacao no config.yaml."""
    regras = dict(REGRAS_PADRAO)
    propria = ((config or {}).get("portal_cliente", {}) or {}).get("cotacao", {}) or {}
    for k, v in propria.items():
        if v is None or v == "":
            continue
        regras[k] = v
    veiculos = []
    for v in regras.get("veiculos") or VEICULOS_PADRAO:
        d = dict(v)
        d["base"] = float(d["base"]); d["km_adicional"] = float(d["km_adicional"])
        d["franquia_km"] = float(d["franquia_km"]); d["caixas_max"] = int(d["caixas_max"])
        d["peso_max_kg"] = float(d["peso_max_kg"])
        veiculos.append(d)
    regras["veiculos"] = sorted(veiculos, key=lambda v: (v["caixas_max"], v["peso_max_kg"]))
    return regras


def regras_publicas(regras: dict) -> dict:
    """O que a tela pode mostrar (sem e-mails internos)."""
    return {
        "veiculos": [{k: v[k] for k in ("codigo", "nome", "base", "franquia_km", "km_adicional",
                                        "caixas_max", "peso_max_kg", "pallets")} for v in regras["veiculos"]],
        "tipos_carga": TIPOS_CARGA,
        "impostos": regras["impostos"], "ad_valorem": regras["ad_valorem"],
        "desconto_seco": regras["desconto_seco"], "urgencia_same_day": regras["urgencia_same_day"],
        "validade_dias": regras["validade_dias"], "origem_endereco": regras["origem_endereco"],
        "max_paradas": regras["max_paradas"], "considerar_retorno": bool(regras.get("considerar_retorno", True)),
    }


# ── A conta ──────────────────────────────────────────────────────────────────────

def _r2(v: float) -> float:
    return round(float(v) + 1e-9, 2)


def escolher_veiculo(caixas: int, peso_kg: float, regras: dict) -> dict:
    """Menor veículo que comporta caixas E peso. ForaDaTabela se nenhum."""
    if caixas <= 0 and peso_kg <= 0:
        raise ErroCotacao("Informe a quantidade de caixas ou o peso da carga.")
    for v in regras["veiculos"]:
        if caixas <= v["caixas_max"] and peso_kg <= v["peso_max_kg"]:
            return v
    maior = regras["veiculos"][-1]
    raise ForaDaTabela(f"Carga acima da capacidade do {maior['nome']} ({maior['caixas_max']} caixas / "
                       f"{maior['peso_max_kg']:.0f} kg). Solicite uma cotação pelo atendimento.")


def km_excedente(km_total: float, franquia_km: float) -> int:
    """Km além da franquia, SEMPRE inteiro arredondado pra cima (Hugo, 11/09)."""
    extra = float(km_total) - float(franquia_km)
    if extra <= 0:
        return 0
    return int(math.ceil(extra - 1e-9))


def calcular(entrada: dict, regras: dict) -> dict:
    """
    entrada: caixas, peso_kg, tipo_carga (SECO|REFRIGERADO|CONGELADO), urgente (bool),
             valor_nf (float|None), km_total (float, JÁ com ida e volta),
             pedagio (float|None)
    Devolve a composição, linha a linha, já arredondada -- o total é a soma
    das linhas arredondadas (fecha o centavo com o PDF/e-mail).
    """
    caixas = int(entrada.get("caixas") or 0)
    peso = float(entrada.get("peso_kg") or 0)
    tipo = str(entrada.get("tipo_carga") or "SECO").upper()
    if tipo not in TIPOS_CARGA:
        raise ErroCotacao("Tipo de carga inválido.")
    veiculo = escolher_veiculo(caixas, peso, regras)
    km_total = float(entrada.get("km_total") or 0)
    excedente = km_excedente(km_total, veiculo["franquia_km"])

    base = _r2(veiculo["base"])
    desconto = _r2(regras["desconto_seco"]) if tipo == "SECO" else 0.0
    valor_km = _r2(excedente * veiculo["km_adicional"])
    frete = _r2(base - desconto + valor_km)

    valor_nf = entrada.get("valor_nf")
    valor_nf = float(valor_nf) if valor_nf not in (None, "") else None
    ad_valorem = _r2(valor_nf * regras["ad_valorem"]) if valor_nf else 0.0

    urgente = bool(entrada.get("urgente"))
    urgencia = _r2((frete + ad_valorem) * regras["urgencia_same_day"]) if urgente else 0.0

    liquido = _r2(frete + ad_valorem + urgencia)
    impostos = _r2(liquido / (1.0 - regras["impostos"]) - liquido)
    total_sem_pedagio = _r2(liquido + impostos)

    pedagio = entrada.get("pedagio")
    pedagio = _r2(pedagio) if pedagio not in (None, "") else None
    total = _r2(total_sem_pedagio + (pedagio or 0.0))

    return {
        "veiculo": veiculo["codigo"], "veiculo_nome": veiculo["nome"],
        "capacidade": {"caixas_max": veiculo["caixas_max"], "peso_max_kg": veiculo["peso_max_kg"], "pallets": veiculo.get("pallets")},
        "caixas": caixas, "peso_kg": peso, "tipo_carga": tipo, "tipo_carga_rotulo": TIPOS_CARGA[tipo],
        "urgente": urgente, "valor_nf": valor_nf,
        "km_total": _r2(km_total), "franquia_km": veiculo["franquia_km"], "km_excedente": excedente,
        "km_adicional_unitario": veiculo["km_adicional"],
        "valor_base": base, "desconto_seco": desconto, "valor_km": valor_km, "frete": frete,
        "ad_valorem": ad_valorem, "ad_valorem_pct": regras["ad_valorem"],
        "urgencia": urgencia, "urgencia_pct": regras["urgencia_same_day"],
        "frete_liquido": liquido,
        "impostos": impostos, "impostos_pct": regras["impostos"],
        "total_sem_pedagio": total_sem_pedagio, "pedagio": pedagio, "total": total,
        # linha ÚNICA que o cliente vê (Hugo, 11/09): base - desconto seco
        # + km adicional + pedágio, tudo já somado.
        "frete_com_pedagio": _r2(frete + (pedagio or 0.0)),
    }


# ── Entrada: CEP, endereços, lista ───────────────────────────────────────────────

_CACHE_CEP: dict[str, dict | None] = {}


def normalizar_cep(valor) -> str:
    d = re.sub(r"\D", "", str(valor or ""))
    return d if len(d) == 8 else ""


def formatar_cep(valor) -> str:
    d = normalizar_cep(valor)
    return f"{d[:5]}-{d[5:]}" if d else str(valor or "")


def _http_json(url: str, timeout: float = 8) -> dict | None:
    try:
        r = requests.get(url, timeout=timeout, headers={"User-Agent": "FreshLog-Portal/1.0"})
        if r.status_code != 200:
            return None
        return r.json()
    except Exception as e:
        logger.warning(f"CEP: falha em {url}: {e}")
        return None


def consultar_cep(cep: str) -> dict | None:
    """{cep, logradouro, bairro, cidade, uf} ou None. ViaCEP, BrasilAPI de reserva."""
    d = normalizar_cep(cep)
    if not d:
        return None
    if d in _CACHE_CEP:
        return _CACHE_CEP[d]
    resultado = None
    v = _http_json(f"https://viacep.com.br/ws/{d}/json/")
    if v and not v.get("erro"):
        resultado = {"cep": formatar_cep(d), "logradouro": v.get("logradouro") or "", "bairro": v.get("bairro") or "",
                     "cidade": v.get("localidade") or "", "uf": v.get("uf") or ""}
    else:
        b = _http_json(f"https://brasilapi.com.br/api/cep/v1/{d}")
        if b and b.get("city"):
            resultado = {"cep": formatar_cep(d), "logradouro": b.get("street") or "", "bairro": b.get("neighborhood") or "",
                         "cidade": b.get("city") or "", "uf": b.get("state") or ""}
    if resultado is not None or v is not None:
        # só cacheia resposta definitiva (achou, ou o ViaCEP disse que não existe)
        _CACHE_CEP[d] = resultado
    return resultado


def endereco_completo(p: dict) -> str:
    partes = []
    log = (p.get("logradouro") or "").strip()
    num = str(p.get("numero") or "").strip()
    if log:
        partes.append(f"{log}, {num}" if num else log)
    elif num:
        partes.append(num)
    if p.get("bairro"):
        partes.append(str(p["bairro"]).strip())
    cid = (p.get("cidade") or "").strip()
    uf = (p.get("uf") or "").strip().upper()
    if cid:
        partes.append(f"{cid} - {uf}" if uf else cid)
    cep = normalizar_cep(p.get("cep"))
    if cep:
        partes.append(formatar_cep(cep))
    return ", ".join(partes)


def normalizar_paradas(paradas: list, regras: dict) -> list[dict]:
    """Valida e limpa a lista de entregas vinda do formulário/lista."""
    if not isinstance(paradas, list) or not paradas:
        raise ErroCotacao("Informe pelo menos um endereço de entrega.")
    if len(paradas) > int(regras.get("max_paradas") or 25):
        raise ErroCotacao(f"No máximo {regras.get('max_paradas')} entregas por cotação.")
    limpas = []
    for i, p in enumerate(paradas, start=1):
        p = p if isinstance(p, dict) else {}
        cep = normalizar_cep(p.get("cep"))
        item = {
            "ordem": i,
            "referencia": str(p.get("referencia") or "").strip()[:80],
            "cep": formatar_cep(cep) if cep else "",
            "logradouro": str(p.get("logradouro") or "").strip()[:120],
            "numero": str(p.get("numero") or "").strip()[:20],
            "complemento": str(p.get("complemento") or "").strip()[:60],
            "bairro": str(p.get("bairro") or "").strip()[:80],
            "cidade": str(p.get("cidade") or "").strip()[:80],
            "uf": str(p.get("uf") or "").strip().upper()[:2],
        }
        if not cep and not (item["logradouro"] and item["cidade"]):
            raise ErroCotacao(f"Entrega {i}: informe o CEP (ou logradouro e cidade).")
        if cep and not item["cidade"]:
            achado = consultar_cep(cep)
            if achado:
                for k in ("logradouro", "bairro", "cidade", "uf"):
                    item[k] = item[k] or achado[k]
        if not item["numero"]:
            raise ErroCotacao(f"Entrega {i}: informe o número.")
        item["endereco"] = endereco_completo(item)
        limpas.append(item)
    return limpas


_COLUNAS_LISTA = [("cep", "CEP"), ("numero", "Número"), ("complemento", "Complemento"),
                  ("referencia", "Referência / destinatário"), ("logradouro", "Logradouro (opcional)"),
                  ("bairro", "Bairro (opcional)"), ("cidade", "Cidade (opcional)"), ("uf", "UF (opcional)")]


def _chave_coluna(texto: str) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().lower()
    t = re.sub(r"[^a-z0-9]+", " ", t).strip()
    if t.startswith("cep"):
        return "cep"
    if t.startswith("num") or t == "n" or t.startswith("nr"):
        return "numero"
    if t.startswith("compl"):
        return "complemento"
    if t.startswith("ref") or t.startswith("destinat") or t.startswith("cliente") or t.startswith("nome"):
        return "referencia"
    if t.startswith("logr") or t.startswith("rua") or t.startswith("endere"):
        return "logradouro"
    if t.startswith("bairro"):
        return "bairro"
    if t.startswith("cidade") or t.startswith("munic"):
        return "cidade"
    if t in ("uf", "estado"):
        return "uf"
    return ""


def modelo_lista() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Entregas"
    ws.append([c[1] for c in _COLUNAS_LISTA])
    ws.append(["01310-100", "1578", "Loja 2", "Mercado Exemplo", "", "", "", ""])
    for i, (_, rot) in enumerate(_COLUNAS_LISTA, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, len(rot) + 4)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ler_lista(conteudo: bytes, nome_arquivo: str) -> list[dict]:
    """Planilha (.xlsx) ou .csv com as colunas do modelo -> lista de paradas
    (ainda sem validação; normalizar_paradas faz isso)."""
    nome = (nome_arquivo or "").lower()
    linhas: list[list] = []
    if nome.endswith(".csv") or nome.endswith(".txt"):
        texto = conteudo.decode("utf-8-sig", errors="replace")
        try:
            dialeto = csv.Sniffer().sniff(texto[:2048], delimiters=";,\t")
        except csv.Error:
            dialeto = csv.excel
            dialeto.delimiter = ";"
        linhas = [r for r in csv.reader(io.StringIO(texto), dialeto)]
    else:
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception:
            raise ErroCotacao("Não consegui ler o arquivo. Use o modelo .xlsx ou um .csv.")
        ws = wb.active
        linhas = [list(r) for r in ws.iter_rows(values_only=True)]
    linhas = [l for l in linhas if any(c not in (None, "") for c in l)]
    if not linhas:
        raise ErroCotacao("Arquivo vazio.")
    cabecalho = [_chave_coluna(c) for c in linhas[0]]
    if "cep" not in cabecalho:
        raise ErroCotacao("Não achei a coluna CEP no cabeçalho. Baixe o modelo pra ver o formato.")
    paradas = []
    for l in linhas[1:]:
        p = {}
        for chave, valor in zip(cabecalho, l):
            if chave and valor not in (None, ""):
                p[chave] = str(valor).strip() if not isinstance(valor, float) or not valor.is_integer() else str(int(valor))
        if p:
            paradas.append(p)
    if not paradas:
        raise ErroCotacao("A lista não tem nenhuma linha de entrega.")
    return paradas


# ── Geocodificação + km ──────────────────────────────────────────────────────────

def _api_key(config: dict) -> str:
    return ((config or {}).get("google_maps", {}) or {}).get("api_key") or ""


def resolver_coordenadas(paradas: list[dict], config: dict) -> list[dict]:
    from geocodificacao import geocodificar
    chave = _api_key(config)
    for p in paradas:
        coords = geocodificar(p["endereco"], chave) if chave else None
        if not coords:
            raise ErroCotacao(f"Não consegui localizar o endereço da entrega {p['ordem']}: {p['endereco']}. "
                              "Confira CEP e número.")
        p["lat"], p["lng"] = float(coords[0]), float(coords[1])
    return paradas


def calcular_trajeto(paradas: list[dict], regras: dict, config: dict) -> dict:
    """Galpão -> entregas na ordem -> galpão. O RETORNO sempre conta na
    quilometragem cobrada (decisão do Hugo, 11/09)."""
    import km_rodoviario
    origem = tuple(regras["origem_coords"])
    r = km_rodoviario.calcular_trajeto(origem, [(p["lat"], p["lng"]) for p in paradas], _api_key(config),
                                       voltar=bool(regras.get("considerar_retorno", True)))
    if r is None:
        raise ErroCotacao("Não foi possível calcular a distância agora. Tente de novo em instantes.")
    return r.como_dict()


# ── Banco ───────────────────────────────────────────────────────────────────────

def conectar() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS portal_cotacoes (
            id                  INTEGER PRIMARY KEY AUTOINCREMENT,
            cnpj_embarcador     TEXT NOT NULL,
            sender_id           INTEGER,
            nome_cliente        TEXT,
            status              TEXT NOT NULL,
            criado_em           TEXT NOT NULL,
            criado_por          TEXT,
            valido_ate          TEXT,
            entrada_json        TEXT NOT NULL,
            resultado_json      TEXT,
            veiculo             TEXT,
            km_total            REAL,
            total               REAL,
            proposta_enviada_em TEXT,
            proposta_para       TEXT,
            aceita_em           TEXT,
            aceita_por          TEXT,
            aceita_via          TEXT,
            chamado_id          INTEGER,
            pdf_path            TEXT
        );
        CREATE INDEX IF NOT EXISTS ix_portal_cotacoes_cnpj ON portal_cotacoes (cnpj_embarcador, criado_em);
    """)
    return conn


def _agora() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def numero(cot_id: int) -> str:
    return f"COT-{int(cot_id):05d}"


def _dict(r: sqlite3.Row | None) -> dict | None:
    if not r:
        return None
    d = dict(r)
    d["entrada"] = json.loads(d.pop("entrada_json") or "{}")
    d["resultado"] = json.loads(d.pop("resultado_json") or "null")
    d["numero"] = numero(d["id"])
    d["status_rotulo"] = STATUS_ROTULO.get(d["status"], d["status"])
    d["expirada"] = bool(d.get("valido_ate")) and d["valido_ate"] < datetime.now().strftime("%Y-%m-%d") \
        and d["status"] != STATUS_ACEITA
    return d


def buscar(conn: sqlite3.Connection, cot_id: int, cnpj: str | None = None) -> dict | None:
    sql, args = "SELECT * FROM portal_cotacoes WHERE id = ?", [cot_id]
    if cnpj:
        sql += " AND cnpj_embarcador = ?"
        args.append(cnpj)
    return _dict(conn.execute(sql, args).fetchone())


def listar(conn: sqlite3.Connection, cnpj: str, limite: int = 100) -> list[dict]:
    rows = conn.execute("SELECT * FROM portal_cotacoes WHERE cnpj_embarcador = ? ORDER BY id DESC LIMIT ?",
                        (cnpj, limite)).fetchall()
    return [_dict(r) for r in rows]


def _atualizar(conn: sqlite3.Connection, cot_id: int, **campos) -> None:
    if not campos:
        return
    sets = ", ".join(f"{k} = ?" for k in campos)
    conn.execute(f"UPDATE portal_cotacoes SET {sets} WHERE id = ?", (*campos.values(), cot_id))
    conn.commit()


# ── Orquestração ─────────────────────────────────────────────────────────────────

def _entrada_limpa(entrada: dict, regras: dict) -> dict:
    def _num(v, nome, minimo=0.0):
        if v in (None, ""):
            return 0.0
        try:
            v = float(str(v).replace(".", "").replace(",", ".")) if isinstance(v, str) and "," in str(v) else float(v)
        except (TypeError, ValueError):
            raise ErroCotacao(f"{nome}: valor inválido.")
        if v < minimo:
            raise ErroCotacao(f"{nome}: não pode ser negativo.")
        return v
    caixas = int(_num(entrada.get("caixas"), "Caixas"))
    peso = _num(entrada.get("peso_kg"), "Peso")
    valor_nf = _num(entrada.get("valor_nf"), "Valor da NF")
    tipo = str(entrada.get("tipo_carga") or "SECO").upper()
    if tipo not in TIPOS_CARGA:
        raise ErroCotacao("Tipo de carga inválido.")
    return {
        "caixas": caixas, "peso_kg": peso, "tipo_carga": tipo,
        "urgente": bool(entrada.get("urgente")),
        "valor_nf": valor_nf or None,
        "observacoes": str(entrada.get("observacoes") or "").strip()[:500],
        "paradas": normalizar_paradas(entrada.get("paradas") or [], regras),
    }


def cotar(conn: sqlite3.Connection, cliente: dict, entrada: dict, config: dict, quem: str = "cliente") -> dict:
    """Valida, geocodifica, calcula km/pedágio, aplica a tabela e GRAVA.
    Carga fora da tabela também fica registrada (status FORA_DA_TABELA),
    mas levanta ForaDaTabela pra tela mostrar 'solicite cotação'."""
    regras = regras_de(config)
    limpa = _entrada_limpa(entrada, regras)
    # Veículo primeiro: se não cabe, nem gasta geocodificação.
    try:
        escolher_veiculo(limpa["caixas"], limpa["peso_kg"], regras)
    except ForaDaTabela as e:
        cur = conn.execute(
            "INSERT INTO portal_cotacoes (cnpj_embarcador, sender_id, nome_cliente, status, criado_em, criado_por, entrada_json) "
            "VALUES (?, ?, ?, ?, ?, ?, ?)",
            (cliente["cnpj"], cliente.get("sender_id"), cliente.get("nome"), STATUS_FORA_DA_TABELA, _agora(), quem,
             json.dumps(limpa, ensure_ascii=False)))
        conn.commit()
        raise ForaDaTabela(f"{e} (registrado como {numero(cur.lastrowid)})")
    resolver_coordenadas(limpa["paradas"], config)
    trajeto = calcular_trajeto(limpa["paradas"], regras, config)
    for p, km in zip(limpa["paradas"], trajeto["pernas_km"]):
        p["km_perna"] = km
    resultado = calcular({**limpa, "km_total": trajeto["km_total"], "pedagio": trajeto["pedagio_brl"]}, regras)
    resultado["km_ida"] = trajeto["km_ida"]
    resultado["km_volta"] = trajeto["km_volta"]
    resultado["km_fonte"] = trajeto["fonte"]
    resultado["origem_endereco"] = regras["origem_endereco"]
    valido_ate = (datetime.now() + timedelta(days=int(regras["validade_dias"]))).strftime("%Y-%m-%d")
    cur = conn.execute(
        "INSERT INTO portal_cotacoes (cnpj_embarcador, sender_id, nome_cliente, status, criado_em, criado_por, valido_ate, "
        "entrada_json, resultado_json, veiculo, km_total, total) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (cliente["cnpj"], cliente.get("sender_id"), cliente.get("nome"), STATUS_SIMULADA, _agora(), quem, valido_ate,
         json.dumps(limpa, ensure_ascii=False), json.dumps(resultado, ensure_ascii=False),
         resultado["veiculo"], resultado["km_total"], resultado["total"]))
    conn.commit()
    return buscar(conn, cur.lastrowid)


# ── Token de aceite (link do e-mail) ─────────────────────────────────────────────

def gerar_token_aceite(secret: str, cot: dict) -> str:
    return URLSafeTimedSerializer(secret).dumps({"c": cot["id"], "cnpj": cot["cnpj_embarcador"]}, salt=_SALT_ACEITE)


def validar_token_aceite(secret: str, conn: sqlite3.Connection, token: str, regras: dict) -> dict:
    """{estado: ok|expirado|invalido|aceita|nao_encontrada, cotacao}"""
    try:
        dados = URLSafeTimedSerializer(secret).loads(token, salt=_SALT_ACEITE,
                                                     max_age=int(regras["validade_dias"]) * 86400)
    except SignatureExpired:
        return {"estado": "expirado", "cotacao": None}
    except BadSignature:
        return {"estado": "invalido", "cotacao": None}
    cot = buscar(conn, int(dados.get("c") or 0), dados.get("cnpj"))
    if not cot:
        return {"estado": "nao_encontrada", "cotacao": None}
    if cot["status"] == STATUS_ACEITA:
        return {"estado": "aceita", "cotacao": cot}
    if cot["expirada"]:
        return {"estado": "expirado", "cotacao": cot}
    return {"estado": "ok", "cotacao": cot}


# ── PDF + e-mails ────────────────────────────────────────────────────────────────

def _fmt_brl(v) -> str:
    if v is None:
        return "—"
    s = f"{float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


def _fmt_data(iso: str | None) -> str:
    if not iso:
        return "—"
    try:
        return datetime.strptime(iso[:10], "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return iso


def gerar_pdf(cot: dict, regras: dict, conn: sqlite3.Connection | None = None) -> Path:
    import cotacao_pdf
    PASTA_PDF.mkdir(parents=True, exist_ok=True)
    caminho = PASTA_PDF / f"{cot['numero']}.pdf"
    caminho.write_bytes(cotacao_pdf.gerar(cot, regras))
    if conn is not None:
        try:
            gravado = str(caminho.relative_to(_RAIZ))
        except ValueError:
            gravado = str(caminho)
        _atualizar(conn, cot["id"], pdf_path=gravado)
    return caminho


def _cfg_email(config: dict) -> dict:
    """Mesma caixa dos chamados (From entregas@, login hugo@)."""
    import chamados
    return chamados.cfg_email_chamados(config)


def _destinos(destinatarios: list[str], regras: dict) -> list[str]:
    forcar = (regras.get("forcar_destino") or "").strip()
    if forcar:
        logger.info(f"cotacao: forcar_destino ativo -- e-mail redirecionado de {destinatarios} pra {forcar}")
        return [forcar]
    return [d for d in destinatarios if d and "@" in d]


def _km_br(v) -> str:
    return f"{float(v or 0):.1f} km".replace(".", ",")


def frete_com_pedagio(r: dict) -> float:
    """A linha única do cliente. Recalcula pra cotação antiga (gravada antes
    de 11/09, quando o resultado_json ainda não tinha o campo)."""
    v = r.get("frete_com_pedagio")
    if v is None:
        v = _r2(float(r.get("frete") or 0) + float(r.get("pedagio") or 0))
    return v


def rotulo_trajeto(r: dict) -> str:
    """'112,5 km (ida e volta)' -- o km cobrado inclui o retorno ao galpão."""
    return _km_br(r.get("km_total")) + (" (ida e volta)" if r.get("km_volta") else " (só ida)")


def linhas_composicao(r: dict, detalhado: bool = False) -> list[tuple[str, str]]:
    """Composição do valor, (rótulo, valor) linha a linha.

    Pro CLIENTE (PDF, tela e e-mails), decisão do Hugo de 11/09: frete base,
    desconto de carga seca, km adicional e pedágio saem numa ÚNICA linha já
    somada -- só ad valorem, urgência e impostos aparecem à parte. O km e o
    pedágio dessa linha são de IDA E VOLTA ao galpão.

    detalhado=True (só no e-mail interno do aceite) abre a composição toda.
    """
    linhas: list[tuple[str, str]] = []
    if detalhado:
        linhas.append((f"Frete base {r['veiculo_nome']} (até {r['franquia_km']:.0f} km)", _fmt_brl(r["valor_base"])))
        if r.get("desconto_seco"):
            linhas.append(("Desconto carga seca", "- " + _fmt_brl(r["desconto_seco"])))
        if r.get("km_excedente"):
            linhas.append((f"Km adicional ({r['km_excedente']} km × {_fmt_brl(r['km_adicional_unitario'])})", _fmt_brl(r["valor_km"])))
        linhas.append((f"Pedágio estimado ({'ida e volta' if r.get('km_volta') else 'só ida'})",
                       _fmt_brl(r["pedagio"]) if r.get("pedagio") is not None else "não estimado"))
    else:
        rotulo = f"Frete dedicado {r['veiculo_nome']} · {rotulo_trajeto(r)}"
        if r.get("pedagio") is not None:
            rotulo += " · pedágio incluso"
        linhas.append((rotulo, _fmt_brl(frete_com_pedagio(r))))
    if r.get("ad_valorem"):
        pct = f"{float(r.get('ad_valorem_pct') or 0.005) * 100:.1f}".replace(".", ",").replace(",0", "")
        linhas.append((f"Ad valorem ({pct}% de {_fmt_brl(r['valor_nf'])})", _fmt_brl(r["ad_valorem"])))
    if r.get("urgencia"):
        linhas.append((f"Urgência same-day (+{float(r.get('urgencia_pct') or 0.40) * 100:.0f}%)", _fmt_brl(r["urgencia"])))
    linhas.append((f"Impostos ({r['impostos_pct'] * 100:.0f}%)", _fmt_brl(r["impostos"])))
    return linhas


def _tabela_html(cot: dict, detalhado: bool = False) -> str:
    r = cot["resultado"]
    tds = "".join(f'<tr><td style="padding:6px 0;color:#6B7280;font-size:13px;">{k}</td>'
                  f'<td style="padding:6px 0;text-align:right;font-size:13px;white-space:nowrap;">{v}</td></tr>'
                  for k, v in linhas_composicao(r, detalhado))
    return (f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="border-top:1px solid #E5E7EB;border-bottom:1px solid #E5E7EB;">{tds}'
            f'<tr><td style="padding:10px 0 4px;font-weight:700;font-size:15px;">Total</td>'
            f'<td style="padding:10px 0 4px;text-align:right;font-weight:700;font-size:18px;color:#141428;">{_fmt_brl(r["total"])}</td></tr></table>')


def _caixa_total_html(cot: dict) -> str:
    """O bloco que o cliente vê no e-mail da proposta: só o preço final."""
    r = cot["resultado"]
    rodape = "impostos e pedágio estimado inclusos" if r.get("pedagio") is not None else "impostos inclusos"
    return (f'<table role="presentation" width="100%" cellpadding="0" cellspacing="0" '
            f'style="border:1px solid #E5E7EB;border-radius:10px;background:#F9FAFB;margin:18px 0 4px;">'
            f'<tr><td style="padding:22px;text-align:center;">'
            f'<div style="font-size:12px;letter-spacing:.08em;color:#6B7280;text-transform:uppercase;">Valor total</div>'
            f'<div style="font-size:34px;font-weight:700;color:#141428;padding:6px 0 2px;">{_fmt_brl(r["total"])}</div>'
            f'<div style="font-size:12px;color:#6B7280;">{rodape}</div>'
            f'</td></tr></table>')


def _entregas_html(cot: dict) -> str:
    itens = "".join(f'<li style="margin:0 0 4px;font-size:13px;">{p["ordem"]}. {p["endereco"]}'
                    f'{" · " + p["referencia"] if p.get("referencia") else ""}</li>' for p in cot["entrada"]["paradas"])
    return f'<ol style="padding-left:18px;margin:8px 0 16px;">{itens}</ol>'


def enviar_proposta(conn: sqlite3.Connection, cot: dict, destinatarios: list[str], config: dict, url_aceite: str,
                    url_portal: str = "") -> dict:
    """PDF + e-mail com botão de aceite. Marca PROPOSTA_ENVIADA (não rebaixa ACEITA)."""
    regras = regras_de(config)
    if cot["status"] == STATUS_FORA_DA_TABELA or not cot.get("resultado"):
        raise ErroCotacao("Essa cotação está fora da tabela -- não há proposta pra enviar.")
    if cot["expirada"]:
        raise ErroCotacao("Essa cotação expirou. Faça uma nova.")
    destinos = _destinos(destinatarios, regras)
    if not destinos:
        raise ErroCotacao("Informe pelo menos um e-mail válido.")
    pdf = gerar_pdf(cot, regras, conn)
    r = cot["resultado"]
    # E-mail OBJETIVO (Hugo, 11/09): preço final e botão de aprovar, nada mais.
    # O detalhamento (entregas, composição, condições) fica no PDF anexo.
    corpo = envelope_html(f"""
      <h2 style="margin:0 0 6px;font-size:18px;color:#141428;">Proposta de frete dedicado {cot['numero']}</h2>
      <p style="margin:0 0 4px;font-size:13px;color:#6B7280;">Para {cot['nome_cliente']} · válida até {_fmt_data(cot['valido_ate'])}</p>
      {_caixa_total_html(cot)}
      <p style="margin:22px 0 8px;text-align:center;">
        <a href="{url_aceite}" style="display:inline-block;background:#00C896;color:#fff;padding:14px 34px;border-radius:7px;
           font-size:16px;font-weight:700;text-decoration:none;">Aprovar proposta</a></p>
      <p style="font-size:12px;color:#6B7280;text-align:center;margin:0;">Detalhamento completo no PDF anexo.
      O botão vale até {_fmt_data(cot['valido_ate'])}.
      {f'Você também pode aprovar <a href="{url_portal}">pelo portal</a>.' if url_portal else ''}</p>
    """)
    ok = enviar_email(destinos, f"Fresh Log · Proposta de frete {cot['numero']} · {_fmt_brl(r['total'])}", corpo,
                      _cfg_email(config), anexos=[(pdf, f"{cot['numero']}.pdf")])
    if not ok:
        raise ErroCotacao("Não consegui enviar o e-mail agora. Tente de novo em instantes.")
    campos = {"proposta_enviada_em": _agora(), "proposta_para": ", ".join(destinatarios)}
    if cot["status"] != STATUS_ACEITA:
        campos["status"] = STATUS_PROPOSTA_ENVIADA
    _atualizar(conn, cot["id"], **campos)
    return buscar(conn, cot["id"])


def _abrir_chamado(conn: sqlite3.Connection, cot: dict, config: dict) -> int | None:
    """Aceite vira chamado no atendimento (painel /painel/atendimento), pra
    equipe programar a coleta. Falha aqui não derruba o aceite."""
    try:
        import chamados as ch
        cliente = {"cnpj": cot["cnpj_embarcador"], "sender_id": cot.get("sender_id"), "nome": cot.get("nome_cliente")}
        r = cot["resultado"]
        # conexão própria: chamados.conectar() é quem garante as tabelas portal_chamados*
        conn_ch = ch.conectar()
        try:
            chamado = ch.criar_chamado(conn_ch, cliente, ch.ORIGEM_SISTEMA, ch.STATUS_AGUARDANDO_FL,
                                       assunto=f"Frete dedicado aceito {cot['numero']} · {r['veiculo_nome']} · {_fmt_brl(r['total'])}",
                                       area="coleta", pedido_ref=cot["numero"])
            entregas = "\n".join(f"{p['ordem']}. {p['endereco']}" for p in cot["entrada"]["paradas"])
            ch.mensagem_sistema(conn_ch, chamado, f"Cliente aceitou a proposta {cot['numero']} ({r['veiculo_nome']}, "
                                                  f"{r['km_total']:.1f} km, {_fmt_brl(r['total'])}).\nEntregas:\n{entregas}\n"
                                                  f"Programar coleta e confirmar data com o cliente.")
            return chamado["id"]
        finally:
            conn_ch.close()
    except Exception as e:
        logger.warning(f"cotacao: não abriu chamado pro aceite {cot['numero']}: {e}")
        return None


def aceitar(conn: sqlite3.Connection, cot: dict, por: str, via: str, config: dict, emails_cliente: list[str]) -> dict:
    """Marca ACEITA, avisa o comercial (com PDF) e confirma pro cliente."""
    regras = regras_de(config)
    if cot["status"] == STATUS_ACEITA:
        return cot
    if cot["status"] == STATUS_FORA_DA_TABELA or not cot.get("resultado"):
        raise ErroCotacao("Essa cotação está fora da tabela.")
    if cot["expirada"]:
        raise ErroCotacao("Essa proposta expirou. Faça uma nova cotação.")
    chamado_id = _abrir_chamado(conn, cot, config) if regras.get("abrir_chamado_no_aceite", True) else None
    _atualizar(conn, cot["id"], status=STATUS_ACEITA, aceita_em=_agora(), aceita_por=por[:120], aceita_via=via,
               chamado_id=chamado_id)
    cot = buscar(conn, cot["id"])
    r = cot["resultado"]
    try:
        pdf = gerar_pdf(cot, regras, conn)
        cfg = _cfg_email(config)
        interno = envelope_html(f"""
          <h2 style="margin:0 0 6px;font-size:18px;color:#141428;">Proposta aceita · {cot['numero']}</h2>
          <p style="margin:0 0 14px;font-size:13px;color:#6B7280;">{cot['nome_cliente']} (CNPJ {cot['cnpj_embarcador']}) ·
          aceita em {cot['aceita_em']} por {por} ({via}){f' · chamado #{chamado_id}' if chamado_id else ''}</p>
          <p style="font-size:14px;margin:0 0 10px;"><b>{r['veiculo_nome']}</b> · {r['caixas']} caixas · {r['peso_kg']:.0f} kg ·
          carga {r['tipo_carga_rotulo'].lower()}{' · same-day' if r['urgente'] else ''} · {rotulo_trajeto(r)}</p>
          {_entregas_html(cot)}{_tabela_html(cot, detalhado=True)}
          {f'<p style="font-size:13px;margin:14px 0 0;"><b>Observações do cliente:</b> {cot["entrada"].get("observacoes")}</p>' if cot['entrada'].get('observacoes') else ''}
        """, cor_acento="#F5A623")
        enviar_email(_destinos([regras["email_comercial"]], regras),
                     f"[ACEITE] Frete dedicado {cot['numero']} · {cot['nome_cliente']} · {_fmt_brl(r['total'])}",
                     interno, cfg, anexos=[(pdf, f"{cot['numero']}.pdf")])
        destinos_cliente = _destinos(emails_cliente, regras)
        if destinos_cliente:
            confirmacao = envelope_html(f"""
              <h2 style="margin:0 0 6px;font-size:18px;color:#141428;">Proposta {cot['numero']} aprovada</h2>
              <p style="font-size:14px;line-height:1.5;margin:0 0 12px;">Recebemos a sua aprovação. Nossa equipe entra em contato
              pra programar a coleta e confirmar a data das entregas.</p>
              {_caixa_total_html(cot)}
              <p style="font-size:12px;color:#6B7280;margin:14px 0 0;">Detalhamento no PDF anexo. O pedágio incluso é estimativa
              do trajeto de ida e volta e é cobrado pelo valor real.</p>
            """)
            enviar_email(destinos_cliente, f"Fresh Log · Proposta {cot['numero']} aprovada", confirmacao, cfg,
                         anexos=[(pdf, f"{cot['numero']}.pdf")])
    except Exception as e:
        logger.exception(f"cotacao: aceite {cot['numero']} gravado, mas e-mail falhou: {e}")
    return cot
