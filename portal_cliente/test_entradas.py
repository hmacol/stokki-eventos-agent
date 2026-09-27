# -*- coding: utf-8 -*-
"""
Aba Pedidos de Entrada do portal (spec
docs/superpowers/specs/2026-09-24-pedidos-de-entrada-portal-design.md).
Sem rede e sem tocar no dados.db real.

    py -3.11 -m unittest portal_cliente.test_entradas -v
"""
import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest import mock

_AQUI = Path(__file__).parent
_RAIZ = _AQUI.parent
for _p in (_RAIZ, _AQUI):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import envio_pedidos as ep  # noqa: E402
import entradas  # noqa: E402

CNPJ = "12345678000195"          # o cliente (emitente da remessa)
CNPJ_FRESHLOG = "11222333000181"  # destinatário (galpão)
OUTRO = "99887766000155"
HOJE = date.today()
AMANHA = (HOJE + timedelta(days=1)).isoformat()


def _agora_txt():
    return ep._agora()


def xml_remessa(nf="41221", emitente=CNPJ, dest=CNPJ_FRESHLOG, tp_nf="1", itens=None, chave=None):
    """NF-e de remessa pra armazenagem: emitente = cliente, destinatário =
    Fresh Log. tpNF=1 (saída do ponto de vista de quem emite)."""
    chave = chave or f"3526091234567800019555001000{nf.zfill(6)}1000000017"[:44]
    itens = itens or [("NUU001FD", "7891234567890", "MINI CX PAO DE QUEIJO", 20, "CX", 64.77)]
    dets = "".join(
        f'<det nItem="{i}"><prod><cProd>{sku}</cProd><cEAN>{ean}</cEAN><xProd>{desc}</xProd>'
        f'<qCom>{qtd}</qCom><uCom>{un}</uCom><vUnCom>{vu}</vUnCom></prod></det>'
        for i, (sku, ean, desc, qtd, un, vu) in enumerate(itens, start=1))
    return f"""<?xml version="1.0"?><nfeProc xmlns="http://www.portalfiscal.inf.br/nfe"><NFe><infNFe Id="NFe{chave}">
      <ide><nNF>{nf}</nNF><serie>1</serie><tpNF>{tp_nf}</tpNF><dhEmi>2026-09-24T10:00:00-03:00</dhEmi></ide>
      <emit><CNPJ>{emitente}</CNPJ><xNome>Cliente Teste</xNome></emit>
      <dest><CNPJ>{dest}</CNPJ><xNome>FRESHLOG</xNome><enderDest><xLgr>Rua G</xLgr><nro>1</nro><xBairro>B</xBairro>
      <xMun>São Paulo</xMun><UF>SP</UF><CEP>01310100</CEP></enderDest></dest>
      {dets}
      <transp><vol><qVol>3</qVol><pesoB>12.5</pesoB></vol></transp><total><ICMSTot><vNF>1295.40</vNF></ICMSTot></total>
    </infNFe></NFe></nfeProc>""".encode("utf-8")


class TestLerNfeEntrada(unittest.TestCase):
    def test_envios_continua_recusando_tpnf_zero(self):
        with self.assertRaises(ep.ErroEnvio):
            ep.ler_nfe(xml_remessa(tp_nf="0"), "x.xml")

    def test_permitir_entrada_aceita_tpnf_zero_e_um(self):
        for tp in ("0", "1"):
            nfe = ep.ler_nfe(xml_remessa(tp_nf=tp), "x.xml", permitir_entrada=True)
            self.assertEqual(nfe["tipo_nf"], tp)
            self.assertEqual(nfe["numero_nf"], "41221")

    def test_itens_lista_vem_do_det(self):
        nfe = ep.ler_nfe(xml_remessa(itens=[("A", "7890000000001", "Prod A", "2.5", "CX", "10.00"),
                                            ("B", "SEM GTIN", "Prod B", 1, "UN", 3)]), "x.xml")
        self.assertEqual(nfe["itens"], 2)
        self.assertEqual(nfe["itens_lista"][0], {"linha": 1, "sku": "A", "ean": "7890000000001", "descricao": "Prod A",
                                                 "quantidade": 2.5, "unidade": "CX", "valor_unitario": 10.0})
        self.assertEqual(nfe["itens_lista"][1]["ean"], "")   # 'SEM GTIN' vira vazio
        self.assertEqual(nfe["itens_lista"][1]["quantidade"], 1.0)


class BasePortal(unittest.TestCase):
    """Banco temporário com `interno` (o cliente é o stkkc 48, piloto) e o
    catálogo wms_produtos do embarcador -- o mesmo desenho de
    test_envio_planilha.ambiente, em unittest."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        raiz = Path(self._tmp.name)
        self._patches = [
            mock.patch.object(ep, "DB_PATH", raiz / "dados.db"),
            mock.patch.object(ep, "_RAIZ", raiz),
            mock.patch.object(ep, "PASTA_XMLS", raiz / "portal_envios"),
            mock.patch.object(ep, "PASTA_TEMP", raiz / "portal_envios" / "_temporarios"),
        ]
        for p in self._patches:
            p.start()
        self.conn = entradas.conectar()
        self.conn.execute("CREATE TABLE IF NOT EXISTS interno (cnpj_embarcador TEXT, apelido TEXT, nome_remetente TEXT, "
                          "stkkc_id INTEGER, sender_id INTEGER, email TEXT, notificar_email INTEGER)")
        self.conn.execute("INSERT INTO interno VALUES (?, 'CLIENTE TESTE', 'CLIENTE TESTE', 48, 1, 'c@t.com', 1)", (CNPJ,))
        self.conn.execute("INSERT INTO interno VALUES (?, 'OUTRO', 'OUTRO', 77, 2, 'o@t.com', 1)", (OUTRO,))
        self.conn.execute("CREATE TABLE IF NOT EXISTS wms_produtos (id INTEGER PRIMARY KEY, stokki_id INTEGER, sku TEXT, ean TEXT, dun TEXT, "
                          "descricao TEXT, embarcador TEXT, embarcador_id INTEGER, qtd_por_caixa REAL, unidade TEXT, ativo INTEGER, atualizado_em TEXT)")
        self.conn.execute("INSERT INTO wms_produtos VALUES (1, 900, 'NUU001FD', '7891234567890', NULL, 'MINI CX PAO DE QUEIJO', "
                          "'CLIENTE TESTE', 48, 1, 'UN', 1, '2026-09-25 10:00:00')")
        self.conn.execute("INSERT INTO wms_produtos VALUES (2, 901, 'SKU-B', NULL, NULL, 'PRODUTO B', 'CLIENTE TESTE', 48, 1, 'UN', 1, '2026-09-25 10:00:00')")
        self.conn.commit()
        self.config = {"portal_entradas": {"cnpj_freshlog": CNPJ_FRESHLOG}}

    def tearDown(self):
        self.conn.close()
        for p in self._patches:
            p.stop()
        self._tmp.cleanup()


class TestTabelas(BasePortal):
    def test_cria_tabelas_e_coluna_entradas_ativo(self):
        nomes = {r[0] for r in self.conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        self.assertIn("portal_entradas", nomes)
        self.assertIn("portal_entrada_itens", nomes)
        cols = {r[1] for r in self.conn.execute("PRAGMA table_info(portal_clientes_envio)")}
        self.assertIn("entradas_ativo", cols)

    def test_garantir_tabelas_e_idempotente_em_conexao_alheia(self):
        # o timer do WMS abre a conexão dele (wms_pedidos.conectar) e chama
        # garantir_tabelas antes de amarrar -- rodar duas vezes não quebra.
        entradas.garantir_tabelas(self.conn)
        entradas.garantir_tabelas(self.conn)
        nomes = {r[0] for r in self.conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        self.assertIn("portal_entradas", nomes)
        self.assertIn("portal_entrada_itens", nomes)


class TestValidarEntrada(BasePortal):
    def _item(self, **kw):
        nfe = entradas.ler_nfe_entrada(xml_remessa(**kw), "nota.xml")
        return nfe

    def test_nfe_valida_sem_erros_com_itens(self):
        v = entradas.validar_entrada(self.conn, self._item(), CNPJ, self.config)
        self.assertTrue(v["ok"], v["erros"])
        self.assertEqual(v["avisos"], [])

    def test_emitente_diferente_do_cliente_e_erro(self):
        v = entradas.validar_entrada(self.conn, self._item(emitente=OUTRO), CNPJ, self.config)
        self.assertFalse(v["ok"])
        self.assertIn("não é o da sua empresa", v["erros"][0])

    def test_emitente_de_outra_empresa_do_grupo_manda_trocar_o_seletor(self):
        v = entradas.validar_entrada(self.conn, self._item(emitente=OUTRO), CNPJ, self.config, {OUTRO: "OUTRO LTDA"})
        self.assertIn("Troque a empresa no seletor", v["erros"][0])

    def test_destinatario_diferente_da_freshlog_e_so_aviso(self):
        v = entradas.validar_entrada(self.conn, self._item(dest="55555555000199"), CNPJ, self.config)
        self.assertTrue(v["ok"])
        self.assertIn("destinatário da nota não é a Fresh Log", v["avisos"][0])

    def test_sem_cnpj_freshlog_no_config_nao_avisa(self):
        v = entradas.validar_entrada(self.conn, self._item(dest="55555555000199"), CNPJ, {})
        self.assertEqual(v["avisos"], [])

    def test_sku_fora_do_catalogo_e_erro(self):
        item = self._item(itens=[("NAO-EXISTE", "", "X", 1, "UN", 1)])
        v = entradas.validar_entrada(self.conn, item, CNPJ, self.config)
        self.assertFalse(v["ok"])
        self.assertIn("SKU não encontrado no catálogo", v["erros"][0])

    def test_chave_ja_anunciada_e_cancelada(self):
        item = self._item()
        entradas.garantir_tabelas(self.conn)
        self.conn.execute("INSERT INTO portal_entradas (cnpj_embarcador, origem, chave_nfe, numero_nf, data_prevista, arquivo_path, "
                          "status, criado_em, atualizado_em) VALUES (?, 'xml', ?, '41221', ?, 'x', 'ANUNCIADO', '2026-09-24 10:00:00', '2026-09-24 10:00:00')",
                          (CNPJ, item["chave_nfe"], AMANHA))
        self.conn.commit()
        v = entradas.validar_entrada(self.conn, item, CNPJ, self.config)
        self.assertFalse(v["ok"])
        self.assertIn("já anunciada em 24/09/2026 10:00 (Anunciado)", v["erros"][0])
        self.conn.execute("UPDATE portal_entradas SET status = 'CANCELADO'")
        self.conn.commit()
        v = entradas.validar_entrada(self.conn, item, CNPJ, self.config)
        self.assertTrue(v["ok"])
        self.assertIn("será reaberta", v["avisos"][0])
        self.assertEqual(v["existente"]["status"], "CANCELADO")


class TestPlanilhaEntrada(BasePortal):
    def _planilha(self, linhas, cabecalho=None):
        import io
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        cab = cabecalho or [c[1] for c in entradas.COLUNAS_PLANILHA_ENTRADA]
        chaves = [c[0] for c in entradas.COLUNAS_PLANILHA_ENTRADA]
        ws.append(cab)
        for l in linhas:
            ws.append([l.get(k, "") for k in (chaves if not cabecalho else cabecalho)])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _linha(self, **kw):
        base = {"referencia": "REM-1", "data_prevista": "/".join(reversed(AMANHA.split("-"))), "sku": "NUU001FD", "quantidade": 10,
                "unidade": "CX", "numero_nf": "", "volumes": 3, "peso_kg": "12,5", "observacoes": ""}
        base.update(kw)
        return base

    def test_modelo_tem_as_colunas_e_e_lido_de_volta(self):
        import io
        import openpyxl
        conteudo = entradas.gerar_modelo_planilha_entrada("Cliente X")
        wb = openpyxl.load_workbook(io.BytesIO(conteudo))
        self.assertEqual(wb.sheetnames, ["Entradas", "Instruções"])
        self.assertEqual([c.value for c in wb["Entradas"][1]], [c[1] for c in entradas.COLUNAS_PLANILHA_ENTRADA])
        pedidos, rejeitados = entradas.ler_planilha_entrada(conteudo, "modelo.xlsx", CNPJ)
        self.assertEqual(rejeitados, [])
        self.assertEqual([p["referencia"] for p in pedidos], ["REM-1001", "REM-1002"])

    def test_agrupa_por_referencia_e_le_data_prevista(self):
        conteudo = self._planilha([self._linha(), self._linha(sku="SKU-B", quantidade=2, data_prevista=""),
                                   self._linha(referencia="REM-2", sku="SKU-B", quantidade="1,5")])
        pedidos, rejeitados = entradas.ler_planilha_entrada(conteudo, "e.xlsx", CNPJ)
        self.assertEqual(rejeitados, [])
        self.assertEqual(len(pedidos), 2)
        p1 = pedidos[0]
        self.assertEqual(p1["origem"], "planilha")
        self.assertEqual(p1["referencia"], "REM-1")
        self.assertEqual(p1["data_prevista"], AMANHA)
        self.assertEqual([i["sku"] for i in p1["itens_lista"]], ["NUU001FD", "SKU-B"])
        self.assertEqual(p1["itens_lista"][0]["unidade"], "CX")
        self.assertEqual(p1["volumes"], 3)
        self.assertEqual(p1["peso_kg"], 12.5)
        self.assertEqual(p1["emitente_cnpj"], CNPJ)
        self.assertTrue(p1["chave_nfe"].startswith(f"PLANILHA-ENTRADA-{CNPJ}-REM-1-"))
        self.assertEqual(pedidos[1]["itens_lista"][0]["quantidade"], 1.5)

    def test_rejeita_grupo_com_problema_e_mantem_os_outros(self):
        ontem = "/".join(reversed((HOJE - timedelta(days=1)).isoformat().split("-")))
        linhas = [self._linha(),
                  self._linha(referencia="REM-2", data_prevista=""),                       # sem data
                  self._linha(referencia="REM-3", data_prevista=ontem),                    # passada
                  self._linha(referencia="REM-4", quantidade=0),                           # sem item válido
                  self._linha(referencia="REM-5"), self._linha(referencia="REM-5", data_prevista="31/12/2030"),  # datas diferentes
                  {"sku": "X", "quantidade": 1}]                                           # sem referência
        pedidos, rejeitados = entradas.ler_planilha_entrada(self._planilha(linhas), "e.xlsx", CNPJ)
        self.assertEqual([p["referencia"] for p in pedidos], ["REM-1"])
        por = {r["rotulo"]: r["erro"] for r in rejeitados}
        self.assertIn("data prevista", por["Remessa REM-2"].lower())
        self.assertIn("já passou", por["Remessa REM-3"])
        self.assertIn("maior que zero", por["Remessa REM-4"])
        self.assertIn("data prevista diferente", por["Remessa REM-5"])
        self.assertIn("linha 8", por)

    def test_cabecalho_com_apelidos(self):
        cab = ["Pedido", "Chegada", "Código", "Qtd"]
        linhas = [{"Pedido": 77, "Chegada": "/".join(reversed(AMANHA.split("-"))), "Código": "SKU-B", "Qtd": 3}]
        pedidos, rejeitados = entradas.ler_planilha_entrada(self._planilha(linhas, cab), "x.xlsx", CNPJ)
        self.assertEqual(rejeitados, [])
        self.assertEqual(pedidos[0]["referencia"], "77")
        self.assertEqual(pedidos[0]["itens_lista"][0]["sku"], "SKU-B")

    def test_falta_coluna_obrigatoria(self):
        with self.assertRaises(ep.ErroEnvio):
            entradas.ler_planilha_entrada(self._planilha([{"Pedido": 1}], ["Pedido", "SKU"]), "x.xlsx", CNPJ)


def _confirmar(self_or_conn, conn=None, itens_extra=None, **kw):
    """Helper: analisa um XML de remessa e confirma pra AMANHA."""
    conn = conn or self_or_conn.conn
    conteudo = xml_remessa(**kw)
    nfe = entradas.ler_nfe_entrada(conteudo, "nota.xml")
    v = entradas.validar_entrada(conn, nfe, CNPJ, {})
    assert v["ok"], v["erros"]
    token = ep.guardar_temporario(conteudo)
    item = {"token": token, "data_prevista": AMANHA, "observacoes": "portão 2"}
    item.update(itens_extra or {})
    return entradas.confirmar_entradas(conn, CNPJ, [item], "cliente", {})


class TestConfirmar(BasePortal):
    def test_grava_entrada_itens_e_move_o_arquivo(self):
        criados = _confirmar(self)
        self.assertEqual(len(criados), 1)
        e = dict(self.conn.execute("SELECT * FROM portal_entradas").fetchone())
        self.assertEqual((e["status"], e["stokki_status"], e["origem"]), ("ANUNCIADO", "NA_FILA", "xml"))
        self.assertEqual(e["numero_nf"], "41221")
        self.assertEqual(e["data_prevista"], AMANHA)
        self.assertEqual(e["observacoes"], "portão 2")
        self.assertTrue(e["arquivo_path"].endswith(f"{e['chave_nfe']}.xml"))
        self.assertTrue((ep._RAIZ / e["arquivo_path"]).is_file())
        self.assertFalse(list(ep.PASTA_TEMP.glob("*")))
        itens = [dict(r) for r in self.conn.execute("SELECT * FROM portal_entrada_itens ORDER BY linha")]
        self.assertEqual([(i["sku"], i["quantidade"], i["unidade"]) for i in itens], [("NUU001FD", 20.0, "CX")])

    def test_data_prevista_obrigatoria_e_nao_passada(self):
        for ruim in ("", None, (HOJE - timedelta(days=1)).isoformat(), "31/12/2030"):
            with self.assertRaises(ep.ErroEnvio):
                _confirmar(self, itens_extra={"data_prevista": ruim})
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM portal_entradas").fetchone()[0], 0)
        _confirmar(self, itens_extra={"data_prevista": HOJE.isoformat()})   # hoje vale

    def test_reanunciar_cancelada_regrava_a_mesma_linha(self):
        _confirmar(self)
        self.conn.execute("UPDATE portal_entradas SET status = 'CANCELADO', stokki_status = 'CRIADO', stokki_id = 9")
        self.conn.execute("INSERT INTO portal_entrada_itens (entrada_id, linha, sku, quantidade) VALUES (1, 99, 'VELHO', 1)")
        self.conn.commit()
        _confirmar(self, itens_extra={"data_prevista": (HOJE + timedelta(days=3)).isoformat()})
        rows = [dict(r) for r in self.conn.execute("SELECT * FROM portal_entradas")]
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]["status"], rows[0]["stokki_status"], rows[0]["stokki_id"]), ("ANUNCIADO", "NA_FILA", None))
        self.assertEqual(rows[0]["data_prevista"], (HOJE + timedelta(days=3)).isoformat())
        skus = [r[0] for r in self.conn.execute("SELECT sku FROM portal_entrada_itens")]
        self.assertEqual(skus, ["NUU001FD"])

    def test_planilha_confirmada_guarda_a_planilha_uma_vez(self):
        import io
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([c[1] for c in entradas.COLUNAS_PLANILHA_ENTRADA])
        d = "/".join(reversed(AMANHA.split("-")))
        ws.append(["REM-1", d, "NUU001FD", 10, "CX", "", 3, "12,5", ""])
        ws.append(["REM-2", d, "SKU-B", 1, "UN", "", "", "", ""])
        buf = io.BytesIO()
        wb.save(buf)
        conteudo = buf.getvalue()
        pedidos, rej = entradas.ler_planilha_entrada(conteudo, "e.xlsx", CNPJ)
        self.assertEqual(rej, [])
        tk_plan = ep.guardar_temporario(conteudo, "xlsx")
        itens = [{"token": ep.guardar_temporario_pedido(p, tk_plan), "data_prevista": p["data_prevista"]} for p in pedidos]
        criados = entradas.confirmar_entradas(self.conn, CNPJ, itens, "equipe:hugo", {})
        self.assertEqual([c["referencia"] for c in criados], ["REM-1", "REM-2"])
        rows = [dict(r) for r in self.conn.execute("SELECT * FROM portal_entradas ORDER BY id")]
        self.assertEqual(rows[0]["arquivo_path"], rows[1]["arquivo_path"])
        self.assertTrue(rows[0]["arquivo_path"].endswith("_e.xlsx"))
        self.assertEqual(rows[0]["enviado_por"], "equipe:hugo")


class TestListar(BasePortal):
    def test_lista_com_rotulos_atrasada_e_pode_cancelar(self):
        _confirmar(self)
        self.conn.execute("UPDATE portal_entradas SET data_prevista = ?", ((HOJE - timedelta(days=2)).isoformat(),))
        self.conn.commit()
        lista = entradas.listar_entradas(self.conn, CNPJ)
        e = lista[0]
        self.assertEqual(e["rotulo"], "NF 41221")
        self.assertEqual(e["status_rotulo"], "Anunciado")
        self.assertTrue(e["atrasada"])
        self.assertTrue(e["pode_cancelar"])
        self.assertEqual(e["data_prevista_br"], (HOJE - timedelta(days=2)).strftime("%d/%m/%Y"))
        self.assertEqual([i["sku"] for i in e["itens_lista"]], ["NUU001FD"])
        self.assertEqual(e["conferencia"], [])   # sem wms_recebimento_id ainda
        self.assertNotIn("arquivo_path", e)
        r = entradas.resumo_entradas(lista)
        self.assertEqual((r["anunciados"], r["atrasados"], r["divergencias"], r["chegaram_hoje"]), (1, 1, 0, 0))

    def test_conferencia_do_galpao_agrupa_por_sku(self):
        # NF com o mesmo SKU em duas linhas (embalagens diferentes): a aba
        # mostra as linhas como vieram e a conferência do galpão por SKU.
        _confirmar(self, itens=[("NUU001FD", "7891234567890", "MINI CX", 20, "CX", 64.77),
                                ("NUU001FD", "17891234567897", "MINI CX (DUN)", 5, "CX", 64.77)])
        self.conn.execute("CREATE TABLE wms_recebimentos (id INTEGER PRIMARY KEY, id_stokki INTEGER, codigo TEXT, estado TEXT, situacao TEXT, "
                          "observacao_divergencia TEXT DEFAULT '', encerrado_em TEXT DEFAULT '')")
        self.conn.execute("CREATE TABLE wms_recebimento_itens (id INTEGER PRIMARY KEY, recebimento_id INTEGER, linha INTEGER, sku TEXT, "
                          "qtd_embalagem REAL, qtd_un REAL, produto_id INTEGER, qtd_enderecada REAL DEFAULT 0, falta_un REAL DEFAULT 0)")
        self.conn.execute("INSERT INTO wms_recebimentos VALUES (7, 2497, '#PE-2497', 'DIVERGENCIA', 'Recebido', 'chegou avariado', '2026-09-25 11:00:00')")
        self.conn.execute("INSERT INTO wms_recebimento_itens VALUES (1, 7, 1, 'NUU001FD', 25, 25, 1, 22, 3)")
        self.conn.execute("UPDATE portal_entradas SET wms_recebimento_id = 7, status = 'DIVERGENCIA', stokki_id = 2497, stokki_codigo = '#PE-2497'")
        self.conn.commit()
        e = entradas.listar_entradas(self.conn, CNPJ)[0]
        self.assertEqual(len(e["itens_lista"]), 2)
        self.assertEqual(e["conferencia"], [{"sku": "NUU001FD", "descricao": "MINI CX PAO DE QUEIJO", "unidade": "UN",
                                             "anunciado": 25.0, "recebido": 22.0, "falta": 3.0}])
        self.assertEqual(e["observacao_divergencia"], "chegou avariado")
        self.assertFalse(e["pode_cancelar"])

    def test_janela_de_30_dias_mas_abertas_sempre_aparecem(self):
        _confirmar(self)
        self.conn.execute("UPDATE portal_entradas SET criado_em = '2026-01-01 10:00:00'")
        self.conn.commit()
        self.assertEqual(len(entradas.listar_entradas(self.conn, CNPJ)), 1)     # ANUNCIADO: aparece
        self.conn.execute("UPDATE portal_entradas SET status = 'ENDERECADO'")
        self.conn.commit()
        self.assertEqual(entradas.listar_entradas(self.conn, CNPJ), [])         # fechada e velha: some


class TestCancelar(BasePortal):
    def _cliente(self):
        return {"cnpj": CNPJ, "sender_id": 1, "nome": "CLIENTE TESTE"}

    def test_cancelar_anunciada_sem_pe(self):
        _confirmar(self)
        e = entradas.buscar_entrada(self.conn, 1, CNPJ)
        r = entradas.cancelar_entrada(self.conn, e, "cliente", self._cliente(), {})
        self.assertTrue(r["aplicado"])
        self.assertIsNone(r["chamado_id"])
        e = entradas.buscar_entrada(self.conn, 1, CNPJ)
        self.assertEqual(e["status"], "CANCELADO")
        self.assertEqual(e["cancelado_por"], "cliente")

    def test_cancelar_com_pe_abre_chamado(self):
        _confirmar(self)
        self.conn.execute("UPDATE portal_entradas SET stokki_status = 'CRIADO', stokki_id = 2497, stokki_codigo = '#PE-2497'")
        self.conn.commit()
        chamado_falso = mock.Mock()
        chamado_falso.criar_chamado.return_value = {"id": 55}
        chamado_falso.conectar.return_value = mock.MagicMock()
        with mock.patch.object(entradas, "_chamados", return_value=chamado_falso):
            r = entradas.cancelar_entrada(self.conn, entradas.buscar_entrada(self.conn, 1, CNPJ), "cliente", self._cliente(), {})
        self.assertTrue(r["aplicado"])
        self.assertEqual(r["chamado_id"], 55)
        self.assertIn("#PE-2497", chamado_falso.criar_chamado.call_args.kwargs["assunto"])
        self.assertEqual(entradas.buscar_entrada(self.conn, 1, CNPJ)["chamado_id"], 55)

    def test_cancelar_tira_do_wms_so_se_esperado_sem_enderecamento(self):
        _confirmar(self)
        self.conn.execute("CREATE TABLE wms_recebimentos (id INTEGER PRIMARY KEY, id_stokki INTEGER, codigo TEXT, estado TEXT, atualizado_em TEXT)")
        self.conn.execute("CREATE TABLE wms_recebimento_itens (id INTEGER PRIMARY KEY, recebimento_id INTEGER, linha INTEGER, sku TEXT, "
                          "qtd_embalagem REAL, qtd_un REAL, produto_id INTEGER, qtd_enderecada REAL DEFAULT 0, falta_un REAL DEFAULT 0)")
        self.conn.execute("INSERT INTO wms_recebimentos VALUES (7, 2497, '#PE-2497', 'ESPERADO', '')")
        self.conn.execute("INSERT INTO wms_recebimento_itens VALUES (1, 7, 1, 'NUU001FD', 20, 20, 1, 0, 0)")
        self.conn.execute("UPDATE portal_entradas SET wms_recebimento_id = 7")
        self.conn.commit()
        entradas.cancelar_entrada(self.conn, entradas.buscar_entrada(self.conn, 1, CNPJ), "cliente", self._cliente(), {})
        self.assertEqual(self.conn.execute("SELECT estado FROM wms_recebimentos WHERE id = 7").fetchone()[0], "CANCELADO")

    def test_cancelar_fora_de_anunciado_e_recusado(self):
        _confirmar(self)
        for st in ("CHEGOU", "ENDERECADO", "DIVERGENCIA", "CANCELADO"):
            self.conn.execute("UPDATE portal_entradas SET status = ?", (st,))
            self.conn.commit()
            with self.assertRaises(ep.ErroEnvio):
                entradas.cancelar_entrada(self.conn, entradas.buscar_entrada(self.conn, 1, CNPJ), "cliente", self._cliente(), {})

    def test_cancelar_enquanto_worker_envia_e_recusado(self):
        _confirmar(self)
        self.conn.execute("UPDATE portal_entradas SET stokki_status = 'ENVIANDO'")
        self.conn.commit()
        with self.assertRaises(ep.ErroEnvio):
            entradas.cancelar_entrada(self.conn, entradas.buscar_entrada(self.conn, 1, CNPJ), "cliente", self._cliente(), {})

    def test_cancelar_criado_sem_id_abre_chamado(self):
        # o worker pode marcar CRIADO sem ainda ter achado o #PE (o timer do
        # WMS amarra depois) -- o cancelamento não pode decidir só pelo
        # stokki_id e deixar o #PE vivo sem chamado nenhum.
        _confirmar(self)
        self.conn.execute("UPDATE portal_entradas SET stokki_status = 'CRIADO', stokki_id = NULL, stokki_codigo = NULL")
        self.conn.commit()
        chamado_falso = mock.Mock()
        chamado_falso.criar_chamado.return_value = {"id": 55}
        chamado_falso.conectar.return_value = mock.MagicMock()
        with mock.patch.object(entradas, "_chamados", return_value=chamado_falso):
            r = entradas.cancelar_entrada(self.conn, entradas.buscar_entrada(self.conn, 1, CNPJ), "cliente", self._cliente(), {})
        self.assertTrue(r["aplicado"])
        self.assertEqual(r["chamado_id"], 55)
        self.assertIn("NF 41221", chamado_falso.criar_chamado.call_args.kwargs["assunto"])
        self.assertEqual(entradas.buscar_entrada(self.conn, 1, CNPJ)["status"], "CANCELADO")

    def test_cancelar_recusa_se_mudou_de_estado_entre_ler_e_gravar(self):
        _confirmar(self)
        e = entradas.buscar_entrada(self.conn, 1, CNPJ)
        # mudou por fora depois que a rota leu `e`
        self.conn.execute("UPDATE portal_entradas SET status = 'CHEGOU'")
        self.conn.commit()
        with self.assertRaises(ep.ErroEnvio):
            entradas.cancelar_entrada(self.conn, e, "cliente", self._cliente(), {})
        self.assertEqual(entradas.buscar_entrada(self.conn, 1, CNPJ)["status"], "CHEGOU")

    def test_cancelar_nao_mexe_no_wms_com_enderecamento_ou_fora_de_esperado(self):
        # cenário artificial (a entrada em si segue ANUNCIADO e é cancelada
        # nos dois casos), só pra provar o guard SQL do UPDATE em
        # wms_recebimentos: nunca mexe fora de ESPERADO sem endereçamento.
        _confirmar(self)
        self.conn.execute("CREATE TABLE wms_recebimentos (id INTEGER PRIMARY KEY, id_stokki INTEGER, codigo TEXT, estado TEXT, atualizado_em TEXT)")
        self.conn.execute("CREATE TABLE wms_recebimento_itens (id INTEGER PRIMARY KEY, recebimento_id INTEGER, linha INTEGER, sku TEXT, "
                          "qtd_embalagem REAL, qtd_un REAL, produto_id INTEGER, qtd_enderecada REAL DEFAULT 0, falta_un REAL DEFAULT 0)")
        self.conn.execute("INSERT INTO wms_recebimentos VALUES (7, 2497, '#PE-2497', 'ESPERADO', '')")
        self.conn.execute("INSERT INTO wms_recebimento_itens VALUES (1, 7, 1, 'NUU001FD', 20, 20, 1, 5, 0)")  # já endereçou parte
        self.conn.execute("UPDATE portal_entradas SET wms_recebimento_id = 7")
        self.conn.commit()
        entradas.cancelar_entrada(self.conn, entradas.buscar_entrada(self.conn, 1, CNPJ), "cliente", self._cliente(), {})
        self.assertEqual(self.conn.execute("SELECT estado FROM wms_recebimentos WHERE id = 7").fetchone()[0], "ESPERADO")
        self.assertEqual(entradas.buscar_entrada(self.conn, 1, CNPJ)["status"], "CANCELADO")

        _confirmar(self, nf="41222")
        self.conn.execute("UPDATE portal_entradas SET wms_recebimento_id = 7 WHERE id = 2")
        self.conn.execute("UPDATE wms_recebimentos SET estado = 'DIVERGENCIA'")
        self.conn.commit()
        entradas.cancelar_entrada(self.conn, entradas.buscar_entrada(self.conn, 2, CNPJ), "cliente", self._cliente(), {})
        self.assertEqual(self.conn.execute("SELECT estado FROM wms_recebimentos WHERE id = 7").fetchone()[0], "DIVERGENCIA")
        self.assertEqual(entradas.buscar_entrada(self.conn, 2, CNPJ)["status"], "CANCELADO")

    def test_cancelar_guarda_motivo_e_manda_pro_chamado(self):
        _confirmar(self)
        self.conn.execute("UPDATE portal_entradas SET stokki_status = 'CRIADO', stokki_id = 2497, stokki_codigo = '#PE-2497', "
                          "observacoes = 'portão 2'")
        self.conn.commit()
        chamado_falso = mock.Mock()
        chamado_falso.criar_chamado.return_value = {"id": 55}
        chamado_falso.conectar.return_value = mock.MagicMock()
        with mock.patch.object(entradas, "_chamados", return_value=chamado_falso):
            entradas.cancelar_entrada(self.conn, entradas.buscar_entrada(self.conn, 1, CNPJ), "cliente", self._cliente(), {},
                                      motivo="pedido cancelado pelo comprador")
        e = entradas.buscar_entrada(self.conn, 1, CNPJ)
        self.assertEqual(e["observacoes"], "portão 2 | Cancelada: pedido cancelado pelo comprador")
        msg = chamado_falso.mensagem_sistema.call_args.args[2]
        self.assertIn("pedido cancelado pelo comprador", msg)


class TestFlagCliente(BasePortal):
    def test_nasce_desligada_e_liga_por_cnpj(self):
        self.assertFalse(entradas.config_entradas_cliente(self.conn, CNPJ)["entradas_ativo"])
        entradas.definir_entradas_ativo(self.conn, CNPJ, True)
        self.assertTrue(entradas.config_entradas_cliente(self.conn, CNPJ)["entradas_ativo"])
        self.assertTrue(entradas.entradas_ativas_para(self.conn, [OUTRO, CNPJ]))
        self.assertFalse(entradas.entradas_ativas_para(self.conn, [OUTRO]))
        entradas.definir_entradas_ativo(self.conn, CNPJ, False)
        self.assertFalse(entradas.config_entradas_cliente(self.conn, CNPJ)["entradas_ativo"])


class TestAmarracaoEStatus(BasePortal):
    """As tabelas do WMS aqui são um recorte mínimo (o teste do timer, em
    test_sincronizar_recebimentos_wms.py, usa as reais)."""

    def setUp(self):
        super().setUp()
        self.conn.execute("CREATE TABLE wms_recebimentos (id INTEGER PRIMARY KEY, id_stokki INTEGER, codigo TEXT, stkkc_id TEXT, "
                          "situacao TEXT DEFAULT '', estado TEXT DEFAULT 'ESPERADO', observacao_divergencia TEXT DEFAULT '', "
                          "encerrado_em TEXT DEFAULT '', portal_entrada_id INTEGER, data_prevista TEXT DEFAULT '', atualizado_em TEXT DEFAULT '')")
        self.conn.execute("CREATE TABLE wms_recebimento_itens (id INTEGER PRIMARY KEY, recebimento_id INTEGER, linha INTEGER, sku TEXT, "
                          "qtd_embalagem REAL, qtd_un REAL, produto_id INTEGER, qtd_enderecada REAL DEFAULT 0, falta_un REAL DEFAULT 0)")
        self.conn.commit()

    def _rec(self, rid=7, id_stokki=2497, estado="ESPERADO", situacao="Em transito", enderecada=0):
        self.conn.execute("INSERT INTO wms_recebimentos (id, id_stokki, codigo, stkkc_id, situacao, estado) VALUES (?,?,?,?,?,?)",
                          (rid, id_stokki, f"#PE-{id_stokki}", "48", situacao, estado))
        self.conn.execute("INSERT INTO wms_recebimento_itens (recebimento_id, linha, sku, qtd_embalagem, qtd_un, produto_id, qtd_enderecada) "
                          "VALUES (?, 1, 'NUU001FD', 20, 20, 1, ?)", (rid, enderecada))
        self.conn.commit()

    def test_amarra_por_stokki_id(self):
        _confirmar(self)
        self.conn.execute("UPDATE portal_entradas SET stokki_status = 'CRIADO', stokki_id = 2497")
        self.conn.commit()
        self._rec()
        eid = entradas.amarrar_recebimento(self.conn, 7, 2497, "#PE-2497", "48", chave_nfe="", ref_pedido="")
        self.assertEqual(eid, 1)
        rec = dict(self.conn.execute("SELECT * FROM wms_recebimentos WHERE id = 7").fetchone())
        self.assertEqual((rec["portal_entrada_id"], rec["data_prevista"]), (1, AMANHA))
        e = entradas.buscar_entrada(self.conn, 1, CNPJ)
        self.assertEqual((e["wms_recebimento_id"], e["stokki_codigo"]), (7, "#PE-2497"))

    def test_amarra_pela_chave_nfe_quando_o_pe_foi_criado_a_mao(self):
        criados = _confirmar(self)
        chave = entradas.buscar_entrada(self.conn, criados[0]["id"], CNPJ)["chave_nfe"]
        self._rec()
        eid = entradas.amarrar_recebimento(self.conn, 7, 2497, "#PE-2497", "48", chave_nfe=chave, ref_pedido="41221")
        self.assertEqual(eid, 1)
        e = entradas.buscar_entrada(self.conn, 1, CNPJ)
        self.assertEqual((e["stokki_id"], e["stokki_codigo"], e["stokki_status"]), (2497, "#PE-2497", "CRIADO"))

    def test_amarrar_por_chave_nao_rouba_entrada_ja_ligada_a_outro_pe(self):
        # entrada já ligada ao #PE 2497 (stokki_id): um #PE NOVO (2600) com a
        # MESMA chave de NF-e não pode roubar essa entrada -- caso raro
        # (chave duplicada), mas o guard tem que recusar em vez de
        # sobrescrever o stokki_id de quem já estava ligado.
        criados = _confirmar(self)
        chave = entradas.buscar_entrada(self.conn, criados[0]["id"], CNPJ)["chave_nfe"]
        self.conn.execute("UPDATE portal_entradas SET stokki_id = 2497, stokki_codigo = '#PE-2497', stokki_status = 'CRIADO'")
        self.conn.commit()
        self._rec(rid=8, id_stokki=2600)
        eid = entradas.amarrar_recebimento(self.conn, 8, 2600, "#PE-2600", "48", chave_nfe=chave, ref_pedido="")
        self.assertIsNone(eid)
        e = entradas.buscar_entrada(self.conn, 1, CNPJ)
        self.assertEqual((e["stokki_id"], e["stokki_codigo"]), (2497, "#PE-2497"))

    def test_amarra_planilha_pela_referencia_do_mesmo_embarcador(self):
        entradas.garantir_tabelas(self.conn)
        self.conn.execute("INSERT INTO portal_entradas (cnpj_embarcador, origem, chave_nfe, referencia, data_prevista, arquivo_path, status, "
                          "criado_em, atualizado_em) VALUES (?, 'planilha', 'PLANILHA-ENTRADA-x', 'REM-9', ?, 'x', 'ANUNCIADO', ?, ?)",
                          (CNPJ, AMANHA, _agora_txt(), _agora_txt()))
        self.conn.execute("INSERT INTO portal_entradas (cnpj_embarcador, origem, chave_nfe, referencia, data_prevista, arquivo_path, status, "
                          "criado_em, atualizado_em) VALUES (?, 'planilha', 'PLANILHA-ENTRADA-y', 'REM-9', ?, 'x', 'ANUNCIADO', ?, ?)",
                          (OUTRO, AMANHA, _agora_txt(), _agora_txt()))
        self.conn.commit()
        self._rec()
        eid = entradas.amarrar_recebimento(self.conn, 7, 2497, "#PE-2497", "48", chave_nfe="", ref_pedido="REM-9")
        self.assertEqual(eid, 1)   # a do cliente 48, não a do OUTRO (stkkc 77)

    def test_sem_par_nao_amarra_e_nao_inventa(self):
        _confirmar(self)
        self._rec()
        self.assertIsNone(entradas.amarrar_recebimento(self.conn, 7, 2497, "#PE-2497", "48", chave_nfe="0" * 44, ref_pedido="999"))
        self.assertIsNone(entradas.buscar_entrada(self.conn, 1, CNPJ)["wms_recebimento_id"])

    def test_sincronizar_status_anda_so_pra_frente(self):
        _confirmar(self)
        self._rec(situacao="Em transito")
        self.conn.execute("UPDATE portal_entradas SET wms_recebimento_id = 7, stokki_id = 2497")
        self.conn.commit()
        self.assertEqual(entradas.sincronizar_status(self.conn), 0)
        self.assertEqual(entradas.buscar_entrada(self.conn, 1, CNPJ)["status"], "ANUNCIADO")
        self.conn.execute("UPDATE wms_recebimentos SET situacao = 'Recebido'")
        self.conn.commit()
        self.assertEqual(entradas.sincronizar_status(self.conn), 1)
        self.assertEqual(entradas.buscar_entrada(self.conn, 1, CNPJ)["status"], "CHEGOU")
        self.conn.execute("UPDATE wms_recebimentos SET estado = 'ENDERECADO'")
        self.conn.commit()
        entradas.sincronizar_status(self.conn)
        self.assertEqual(entradas.buscar_entrada(self.conn, 1, CNPJ)["status"], "ENDERECADO")
        self.conn.execute("UPDATE wms_recebimentos SET estado = 'ESPERADO', situacao = 'Em transito'")
        self.conn.commit()
        entradas.sincronizar_status(self.conn)
        self.assertEqual(entradas.buscar_entrada(self.conn, 1, CNPJ)["status"], "ENDERECADO")   # nunca volta

    def test_primeiro_enderecamento_tambem_vira_chegou(self):
        _confirmar(self)
        self._rec(situacao="Em transito", enderecada=5)
        self.conn.execute("UPDATE portal_entradas SET wms_recebimento_id = 7")
        self.conn.commit()
        entradas.sincronizar_status(self.conn)
        self.assertEqual(entradas.buscar_entrada(self.conn, 1, CNPJ)["status"], "CHEGOU")

    def test_divergencia_e_cancelado(self):
        _confirmar(self)
        self._rec(estado="DIVERGENCIA", situacao="Recebido")
        self.conn.execute("UPDATE portal_entradas SET wms_recebimento_id = 7")
        self.conn.commit()
        entradas.sincronizar_status(self.conn)
        self.assertEqual(entradas.buscar_entrada(self.conn, 1, CNPJ)["status"], "DIVERGENCIA")
        self.conn.execute("UPDATE portal_entradas SET status = 'CANCELADO'")
        self.conn.execute("UPDATE wms_recebimentos SET estado = 'ENDERECADO'")
        self.conn.commit()
        entradas.sincronizar_status(self.conn)
        self.assertEqual(entradas.buscar_entrada(self.conn, 1, CNPJ)["status"], "CANCELADO")   # CANCELADO nunca reabre


class TestCli(BasePortal):
    def test_entradas_ativar_avisa_quando_nao_e_o_piloto(self):
        import io
        import contextlib
        import gerenciar_clientes as cli
        saida = io.StringIO()
        # sem patch em envios.conectar: ep.DB_PATH já aponta pro banco temporário
        # (BasePortal) e o comando chama garantir_tabelas antes de tudo
        with mock.patch.object(cli, "_config", return_value={"wms": {"embarcador_piloto_id": "48"}}), \
             contextlib.redirect_stdout(saida):
            cli.main(["entradas", OUTRO, "--ativar"])
            cli.main(["entradas", CNPJ, "--ativar"])
            cli.main(["entradas", CNPJ])
        texto = saida.getvalue()
        self.assertIn("ATENÇÃO", texto)                 # OUTRO (stkkc 77) não é o piloto 48
        self.assertIn("Pedidos de Entrada: ATIVO", texto)
        self.assertTrue(entradas.config_entradas_cliente(self.conn, CNPJ)["entradas_ativo"])
        self.assertTrue(entradas.config_entradas_cliente(self.conn, OUTRO)["entradas_ativo"])   # avisa, mas obedece


if __name__ == "__main__":
    unittest.main()
