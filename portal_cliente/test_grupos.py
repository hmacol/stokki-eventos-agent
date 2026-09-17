# -*- coding: utf-8 -*-
"""
Testes do login de grupo do portal do cliente (decisao do Hugo, 17/09: um
CNPJ + PIN por grupo economico, com os pedidos de todas as empresas do
grupo na mesma tela -- Marchef tem 10 CNPJs, Grupo Trigo 2 remetentes).

    py -3.11 -m unittest portal_cliente.test_grupos
"""
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

_AQUI = Path(__file__).parent
_RAIZ = _AQUI.parent
for _p in (_RAIZ, _AQUI):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import auth_cliente as auth  # noqa: E402

PRINCIPAL, MEMBRO_A, MEMBRO_B, AVULSO, SEM_SENDER = (
    "54993021000184", "48654566000163", "16967388000134", "22310186000118", "19661785000171")


class _BancoTemporario(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self._patch = mock.patch.object(auth, "DB_PATH", Path(self._tmp.name) / "dados.db")
        self._patch.start()
        self.conn = auth.conectar()
        self.conn.execute("CREATE TABLE interno (cnpj_embarcador TEXT, sender_id INTEGER, nome_remetente TEXT, "
                          "apelido TEXT, email TEXT)")
        self.conn.executemany("INSERT INTO interno VALUES (?, ?, ?, NULL, ?)", [
            (PRINCIPAL, 100, "MARCHEF", "log@marchef.com"), (MEMBRO_A, 101, "ITAUEIRA", "log@marchef.com"),
            (MEMBRO_B, 102, "GOURMAR", "log@marchef.com"), (AVULSO, 200, "VIDAVEG", "log@vidaveg.com"),
            (SEM_SENDER, None, "ZANKY", "")])
        self.conn.commit()

    def tearDown(self):
        self.conn.close()
        self._patch.stop()
        self._tmp.cleanup()


class GruposTest(_BancoTemporario):
    def test_sem_grupo_a_empresa_e_so_ela(self):
        empresas = auth.empresas_do_login(self.conn, AVULSO)
        self.assertEqual([e["cnpj"] for e in empresas], [AVULSO])

    def test_grupo_traz_o_principal_primeiro_e_os_membros_por_nome(self):
        auth.definir_grupo(self.conn, PRINCIPAL, [MEMBRO_A, MEMBRO_B])
        empresas = auth.empresas_do_login(self.conn, PRINCIPAL)
        self.assertEqual([e["cnpj"] for e in empresas], [PRINCIPAL, MEMBRO_B, MEMBRO_A])
        self.assertEqual([e["sender_id"] for e in empresas], [100, 102, 101])

    def test_membro_logando_sozinho_nao_enxerga_o_grupo(self):
        auth.definir_grupo(self.conn, PRINCIPAL, [MEMBRO_A, MEMBRO_B])
        self.assertEqual([e["cnpj"] for e in auth.empresas_do_login(self.conn, MEMBRO_A)], [MEMBRO_A])

    def test_definir_grupo_substitui_a_lista_e_ignora_o_proprio_principal(self):
        auth.definir_grupo(self.conn, PRINCIPAL, [MEMBRO_A, MEMBRO_B])
        auth.definir_grupo(self.conn, PRINCIPAL, [PRINCIPAL, "48.654.566/0001-63"])
        self.assertEqual([e["cnpj"] for e in auth.empresas_do_login(self.conn, PRINCIPAL)], [PRINCIPAL, MEMBRO_A])

    def test_definir_grupo_recusa_cnpj_fora_do_cadastro_sem_gravar_nada(self):
        with self.assertRaises(ValueError):
            auth.definir_grupo(self.conn, PRINCIPAL, [MEMBRO_A, SEM_SENDER])
        with self.assertRaises(ValueError):
            auth.definir_grupo(self.conn, "00000000000000", [MEMBRO_A])
        self.assertEqual(len(auth.empresas_do_login(self.conn, PRINCIPAL)), 1)

    def test_membro_que_saiu_do_cadastro_some_do_grupo(self):
        auth.definir_grupo(self.conn, PRINCIPAL, [MEMBRO_A, MEMBRO_B])
        self.conn.execute("UPDATE interno SET sender_id = NULL WHERE cnpj_embarcador = ?", (MEMBRO_A,))
        self.conn.commit()
        self.assertEqual([e["cnpj"] for e in auth.empresas_do_login(self.conn, PRINCIPAL)], [PRINCIPAL, MEMBRO_B])

    def test_listar_grupos(self):
        auth.definir_grupo(self.conn, PRINCIPAL, [MEMBRO_A])
        self.assertEqual(auth.listar_grupos(self.conn), {PRINCIPAL: [MEMBRO_A]})

    def test_sessao_valida_carrega_as_empresas_do_grupo(self):
        auth.definir_grupo(self.conn, PRINCIPAL, [MEMBRO_A])
        conta = auth.definir_pin(self.conn, PRINCIPAL, "123456")
        cliente = auth.sessao_valida(self.conn, PRINCIPAL, auth.versao_conta(conta))
        self.assertEqual(cliente["sender_id"], 100)
        self.assertEqual(cliente["sender_ids"], [100, 101])
        self.assertEqual([e["nome"] for e in cliente["empresas"]], ["MARCHEF", "ITAUEIRA"])
        self.assertEqual(cliente["empresas"][1]["cnpj_formatado"], "48.654.566/0001-63")


# ── Dados consolidados (dados_cliente) ─────────────────────────────────────────

import dados_cliente as dados  # noqa: E402


def _servico(sid, sender, code, status="on_route", **kw):
    return {"id": sid, "sender_id": sender, "code": code, "status": status, "customer": {"name": f"Dest {sid}"},
            "address": "Rua X, 1", **kw}


class DadosConsolidadosTest(unittest.TestCase):
    def setUp(self):
        p = mock.patch.object(dados, "nf_por_codigo", return_value={})
        p.start()
        self.addCleanup(p.stop)
        self.rotas = [{"id": 1, "name": "Rota 1", "status": "started", "start_at": "2026-09-17 06:00:00",
                       "services": [_servico(1, 100, "PS-1"), _servico(2, 101, "PS-2"), _servico(3, 999, "PS-3"),
                                    _servico(4, 102, "PS-4", status="done", status_done="success",
                                             completed_at="2026-09-17 12:00:00")]}]

    def test_um_remetente_continua_funcionando_com_int(self):
        linhas = dados.linhas_das_rotas(self.rotas, 100, {}, {})
        self.assertEqual([l["codigo"] for l in linhas], ["#PS-1"])
        self.assertEqual(linhas[0]["sender_id"], 100)

    def test_varios_remetentes_trazem_os_pedidos_de_todos_e_so_deles(self):
        linhas = dados.linhas_das_rotas(self.rotas, (100, 101, 102), {}, {})
        self.assertEqual([(l["codigo"], l["sender_id"]) for l in linhas], [("#PS-1", 100), ("#PS-2", 101), ("#PS-4", 102)])
        self.assertEqual(linhas[1]["detalhe"], "Parada 2 de 4 · 1 já feitas")   # ordem conta a rota inteira

    def test_historico_soma_os_remetentes_e_recalcula_o_percentual(self):
        partes = [{"entregues": 8, "insucessos": 1, "primeira": 8}, {"entregues": 2, "insucessos": 3, "primeira": 0}, None]
        h = dados.somar_historico(partes)
        self.assertEqual((h["entregues"], h["insucessos"], h["primeira_tentativa_pct"]), (10, 4, 80))
        self.assertIsNone(dados.somar_historico([None, None]))

    def test_cache_e_invalidacao_por_conjunto_de_remetentes(self):
        hoje = dados.date.today()
        with mock.patch.object(dados, "_montar_dia", side_effect=lambda s, d, c: {"remetentes": s}) as montar:
            dados.montar_dia((101, 100), hoje, {})
            dados.montar_dia([100, 101], hoje, {})          # mesma chave, ordem nao importa
            self.assertEqual(montar.call_count, 1)
            dados.montar_dia(100, hoje, {})                  # login avulso e outra chave
            self.assertEqual(montar.call_count, 2)
            dados.invalidar_cache(101)                       # derruba so o grupo que contem o 101
            dados.montar_dia((100, 101), hoje, {})
            dados.montar_dia(100, hoje, {})
            self.assertEqual(montar.call_count, 3)


class AtencaoPorEmpresaTest(_BancoTemporario):
    def test_insucessos_agrupam_por_empresa_e_motivo(self):
        self.conn.execute("CREATE TABLE insucessos_aguardando_resposta (service_id INTEGER, code TEXT, sender_id INTEGER, "
                          "failed_reason_id INTEGER, status TEXT, primeira_notificacao_em TEXT)")
        self.conn.executemany("INSERT INTO insucessos_aguardando_resposta VALUES (?, ?, ?, ?, 'PENDENTE', ?)", [
            (1, "PS-1", 100, 7, "2026-09-17 09:00"), (2, "PS-2", 101, 7, "2026-09-17 09:10"),
            (3, "PS-3", 101, 7, "2026-09-17 09:20"), (4, "PS-4", 999, 7, "2026-09-17 09:30")])
        self.conn.commit()
        with mock.patch.object(dados, "_conectar", side_effect=auth.conectar), \
             mock.patch.object(dados, "texto_do_motivo", return_value="Fechado"), \
             mock.patch.object(dados, "pergunta_do_motivo", return_value="?"):
            grupos = dados._coletar_atencao((100, 101))
        self.assertEqual([(g["sender_id"], g["failed_reason_id"], g["codigos"]) for g in grupos],
                         [(100, 7, ["#PS-1"]), (101, 7, ["#PS-2", "#PS-3"])])


# ── Rotas do app com login de grupo ────────────────────────────────────────────

class RotasGrupoTest(_BancoTemporario):
    HTTPS = "https://localhost"   # o cookie de sessao e Secure

    @classmethod
    def setUpClass(cls):
        import app as portal
        cls.portal = portal

    def setUp(self):
        super().setUp()
        auth.definir_grupo(self.conn, PRINCIPAL, [MEMBRO_A])
        self.cli = self._logar(PRINCIPAL)

    def _logar(self, cnpj):
        conta = auth.definir_pin(self.conn, cnpj, "123456")
        cli = self.portal.app.test_client()
        with cli.session_transaction(base_url=self.HTTPS) as s:
            s["cnpj"], s["v"] = cnpj, auth.versao_conta(conta)
        return cli

    def test_api_dia_consulta_todas_as_empresas_e_devolve_os_nomes(self):
        with mock.patch.object(self.portal.dados, "montar_dia", return_value={"pedidos": []}) as montar:
            r = self.cli.get("/api/dia", base_url=self.HTTPS)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(list(montar.call_args.args[0]), [100, 101])
        self.assertEqual([(e["sender_id"], e["nome"]) for e in r.get_json()["empresas"]],
                         [(100, "MARCHEF"), (101, "ITAUEIRA")])

    def test_login_avulso_segue_com_uma_empresa_so(self):
        cli = self._logar(AVULSO)
        with mock.patch.object(self.portal.dados, "montar_dia", return_value={"pedidos": []}) as montar:
            r = cli.get("/api/dia", base_url=self.HTTPS)
        self.assertEqual(list(montar.call_args.args[0]), [200])
        self.assertEqual(len(r.get_json()["empresas"]), 1)

    def test_canhoto_de_empresa_do_grupo_passa_e_de_fora_da_404(self):
        with mock.patch.object(self.portal.dados, "buscar_servico", return_value={"sender_id": 101, "code": "PS-1"}), \
             mock.patch.object(self.portal.dados, "checklist_id_do_servico", return_value=None):
            r = self.cli.get("/api/canhoto/1", base_url=self.HTTPS)
            self.assertIn("Comprovante ainda não disponível", r.get_data(as_text=True))
        with mock.patch.object(self.portal.dados, "buscar_servico", return_value={"sender_id": 102, "code": "PS-2"}):
            self.assertEqual(self.cli.get("/api/canhoto/2", base_url=self.HTTPS).status_code, 404)
            self.assertNotIn("Comprovante", self.cli.get("/api/canhoto/2", base_url=self.HTTPS).get_data(as_text=True))

    def _responder(self, **corpo):
        return self.cli.post("/api/responder", json={"failed_reason_id": 7, "acao": "manter", **corpo}, base_url=self.HTTPS)

    def test_responder_insucesso_usa_a_empresa_informada(self):
        with mock.patch.object(self.portal, "buscar_pendentes_por_grupo", return_value=[1]) as pend, \
             mock.patch.object(self.portal.logica_insucesso, "aplicar_decisao", return_value=[]) as aplicar:
            self.assertEqual(self._responder(sender_id=101).status_code, 200)
            self.assertEqual(pend.call_args.args, (101, 7))
            self.assertEqual(aplicar.call_args.args[:3], (101, 7, "manter"))
            self.assertEqual(self._responder().status_code, 200)       # sem sender_id = empresa do login
            self.assertEqual(aplicar.call_args.args[0], 100)

    def test_responder_insucesso_de_empresa_de_fora_e_recusado(self):
        with mock.patch.object(self.portal, "buscar_pendentes_por_grupo", return_value=[1]), \
             mock.patch.object(self.portal.logica_insucesso, "aplicar_decisao", return_value=[]) as aplicar:
            self.assertEqual(self._responder(sender_id=102).status_code, 403)
            self.assertEqual(self._responder(sender_id="abc").status_code, 400)
            aplicar.assert_not_called()

    def test_exportacao_do_grupo_tem_coluna_empresa(self):
        import io
        import openpyxl
        dia = {"pedidos": [{"codigo": "#PS-1", "sender_id": 101}], "agendados_futuros": []}
        with mock.patch.object(self.portal.dados, "montar_dia", return_value=dia):
            grupo = openpyxl.load_workbook(io.BytesIO(self.cli.get("/exportar.xlsx", base_url=self.HTTPS).data)).active
            avulso = openpyxl.load_workbook(io.BytesIO(self._logar(AVULSO).get("/exportar.xlsx", base_url=self.HTTPS).data)).active
        self.assertEqual([c.value for c in grupo[1]][:2], ["Empresa", "Pedido"])
        self.assertEqual([c.value for c in grupo[2]][:2], ["ITAUEIRA", "#PS-1"])
        self.assertEqual(avulso[1][0].value, "Pedido")

    def test_exportacao_filtrada_por_empresa_do_grupo(self):
        import io
        import openpyxl
        dia = {"pedidos": [{"codigo": "#PS-1", "sender_id": 101}, {"codigo": "#PS-2", "sender_id": 100}], "agendados_futuros": []}
        with mock.patch.object(self.portal.dados, "montar_dia", return_value=dia):
            so_101 = openpyxl.load_workbook(io.BytesIO(self.cli.get("/exportar.xlsx?empresa=101", base_url=self.HTTPS).data)).active
            de_fora = openpyxl.load_workbook(io.BytesIO(self.cli.get("/exportar.xlsx?empresa=999", base_url=self.HTTPS).data)).active
        self.assertEqual([r[1].value for r in so_101.iter_rows(min_row=2)], ["#PS-1"])
        self.assertEqual(de_fora.max_row, 3)   # empresa que nao e do grupo = ignora o filtro

    # ── envio de pedidos: acontece numa empresa do grupo por vez ──────────────

    def _listar_envios(self, **kw):
        cfg = {"regra_xml": "nenhuma", "envio_ativo": True, "client_id": 61}
        with mock.patch.object(self.portal.envios, "listar_envios", return_value=[]) as listar, \
             mock.patch.object(self.portal.envios, "config_stokki_cliente", return_value=cfg):
            r = self.cli.get("/api/envios", base_url=self.HTTPS, **kw)
        return r, listar

    def test_envios_sem_empresa_informada_valem_pro_cnpj_do_login(self):
        r, listar = self._listar_envios()
        self.assertEqual(r.status_code, 200)
        self.assertEqual(listar.call_args.args[1], PRINCIPAL)

    def test_envios_na_empresa_do_grupo_escolhida(self):
        r, listar = self._listar_envios(headers={"X-Portal-Empresa": "48.654.566/0001-63"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(listar.call_args.args[1], MEMBRO_A)
        r, listar = self._listar_envios(query_string={"empresa": MEMBRO_A})   # links de download
        self.assertEqual(listar.call_args.args[1], MEMBRO_A)

    def test_envios_em_empresa_de_fora_do_grupo_sao_recusados(self):
        r, listar = self._listar_envios(headers={"X-Portal-Empresa": MEMBRO_B})   # existe, mas nao e deste grupo
        self.assertEqual(r.status_code, 403)
        listar.assert_not_called()


class ValidarItemDoGrupoTest(unittest.TestCase):
    def setUp(self):
        import envio_pedidos as ep
        self.ep = ep
        self._tmp = tempfile.TemporaryDirectory()
        p = mock.patch.object(ep, "DB_PATH", Path(self._tmp.name) / "dados.db")
        p.start()
        self.addCleanup(self._tmp.cleanup)
        self.addCleanup(p.stop)
        self.conn = ep.conectar()
        self.addCleanup(self.conn.close)
        self.item = {"chave_nfe": "1" * 44, "emitente_cnpj": MEMBRO_A, "emitente_nome": "ITAUEIRA",
                     "destinatario_doc": "98765432000110", "destinatario_endereco": "Rua X"}

    def test_xml_de_outra_empresa_do_grupo_diz_pra_qual_trocar(self):
        v = self.ep.validar_item(self.conn, self.item, PRINCIPAL, outras_empresas={MEMBRO_A: "ITAUEIRA CAMAROES"})
        self.assertFalse(v["ok"])
        self.assertIn("ITAUEIRA CAMAROES", v["erros"][0])
        self.assertIn("seletor", v["erros"][0])

    def test_xml_de_empresa_de_fora_segue_com_o_erro_de_sempre(self):
        v = self.ep.validar_item(self.conn, self.item, PRINCIPAL, outras_empresas={MEMBRO_B: "GOURMAR"})
        self.assertIn("não é o da sua empresa", v["erros"][0])
        self.assertTrue(self.ep.validar_item(self.conn, self.item, MEMBRO_A)["ok"])


class EnviarLinkParaTest(_BancoTemporario):
    """`enviar-link --para`: o convite vai pra quem o Hugo indicar (ele repassa
    ao cliente), nunca pros e-mails do cadastro."""

    def _rodar(self, *args):
        import gerenciar_clientes as cli
        cfg = {"portal_cliente": {"secret_key": "segredo-de-teste", "url_base": "https://exemplo/cliente"}, "email": {}}
        with mock.patch.object(cli, "_config", return_value=cfg), \
             mock.patch("email_utils.enviar_email", return_value=True) as enviar:
            codigo = cli.main(list(args))
        return codigo, enviar

    def test_sem_para_vai_pros_emails_do_cadastro(self):
        codigo, enviar = self._rodar("enviar-link", AVULSO)
        self.assertEqual(codigo, 0)
        self.assertEqual(enviar.call_args.args[0], ["log@vidaveg.com"])

    def test_com_para_vai_so_pro_destino_indicado_e_avisa_de_quem_e_o_convite(self):
        codigo, enviar = self._rodar("enviar-link", AVULSO, "--para", "hugo@freshlogbr.com")
        self.assertEqual(codigo, 0)
        destinos, assunto, corpo = enviar.call_args.args[:3]
        self.assertEqual(destinos, ["hugo@freshlogbr.com"])
        self.assertIn("VIDAVEG", assunto)
        self.assertIn("log@vidaveg.com", corpo)            # diz pra quem repassar
        self.assertIn("https://exemplo/cliente/definir-pin/", corpo)

    def test_para_invalido_nao_envia(self):
        codigo, enviar = self._rodar("enviar-link", AVULSO, "--para", "nao-e-email")
        self.assertEqual(codigo, 2)
        enviar.assert_not_called()


if __name__ == "__main__":
    unittest.main()
