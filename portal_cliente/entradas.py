# -*- coding: utf-8 -*-
"""
portal_cliente/entradas.py

Aba "Pedidos de Entrada" do portal do cliente (pedido do Hugo, 24/09/2026):
o embarcador anuncia a mercadoria que vai chegar no galpão (XML da NF-e de
remessa emitida por ele, ou planilha no modelo Fresh Log), o worker
enviar_entradas_stokki.py cria o recebimento (#PE) na Stokki pelo wizard
de XML múltiplo, o timer stokki-wms-recebimentos amarra o #PE a esta
entrada pela chave da NF-e, e o cliente acompanha até o galpão endereçar.

Spec: docs/superpowers/specs/2026-09-24-pedidos-de-entrada-portal-design.md

Tabelas (dados/dados.db):
  portal_entradas        -- uma por NF-e/remessa (chave_nfe única): dados da
                            nota, data prevista, status do ciclo, fila da
                            Stokki, ligação com wms_recebimentos.
  portal_entrada_itens   -- as linhas da nota/planilha como vieram.
  portal_clientes_envio  -- ganha entradas_ativo (só o piloto, D3).

Ciclo (portal_entradas.status), só anda pra frente:
  ANUNCIADO -> CHEGOU -> ENDERECADO | DIVERGENCIA ; CANCELADO (só a partir de ANUNCIADO)
Fila da Stokki (portal_entradas.stokki_status):
  NA_FILA -> ENVIANDO -> CRIADO | ERRO

O que este módulo NUNCA faz: escrever em wms_movimentos/wms_saldos. O
estoque entra quando o operador endereça no celular, como hoje.
"""
import hashlib
import json
import logging
import re
import sqlite3
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

_AQUI = Path(__file__).parent
_RAIZ = _AQUI.parent
for _p in (_RAIZ, _AQUI):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import envio_pedidos as ep  # noqa: E402

logger = logging.getLogger(__name__)

DIAS_LISTAGEM = 30
ORIGEM_XML = ep.ORIGEM_XML
ORIGEM_PLANILHA = ep.ORIGEM_PLANILHA

STATUS_ANUNCIADO = "ANUNCIADO"
STATUS_CHEGOU = "CHEGOU"
STATUS_ENDERECADO = "ENDERECADO"
STATUS_DIVERGENCIA = "DIVERGENCIA"
STATUS_CANCELADO = "CANCELADO"
STATUS_ABERTOS = (STATUS_ANUNCIADO, STATUS_CHEGOU)
ROTULOS_STATUS = {
    STATUS_ANUNCIADO: "Anunciado",
    STATUS_CHEGOU: "Chegou no galpão",
    STATUS_ENDERECADO: "Endereçado",
    STATUS_DIVERGENCIA: "Com divergência",
    STATUS_CANCELADO: "Cancelado",
}
STOKKI_NA_FILA = "NA_FILA"
STOKKI_ENVIANDO = "ENVIANDO"
STOKKI_CRIADO = "CRIADO"
STOKKI_ERRO = "ERRO"
ROTULOS_STOKKI = {
    STOKKI_NA_FILA: "Na fila pra Stokki",
    STOKKI_ENVIANDO: "Criando na Stokki",
    STOKKI_CRIADO: "Recebimento criado na Stokki",
    STOKKI_ERRO: "Erro na Stokki",
}

ErroEnvio = ep.ErroEnvio
_so_digitos = ep._so_digitos
_agora = ep._agora
formatar_documento = ep.formatar_documento


# ── Banco ──────────────────────────────────────────────────────────────────────

_DDL = """
CREATE TABLE IF NOT EXISTS portal_entradas (
    id                 INTEGER PRIMARY KEY AUTOINCREMENT,
    cnpj_embarcador    TEXT NOT NULL,
    origem             TEXT NOT NULL,
    chave_nfe          TEXT NOT NULL UNIQUE,
    numero_nf          TEXT,
    serie              TEXT,
    emitida_em         TEXT,
    referencia         TEXT,
    data_prevista      TEXT NOT NULL,
    volumes            INTEGER,
    peso_kg            REAL,
    valor_nf           REAL,
    arquivo_path       TEXT NOT NULL,
    status             TEXT NOT NULL,
    stokki_status      TEXT NOT NULL DEFAULT 'NA_FILA',
    stokki_id          INTEGER,
    stokki_codigo      TEXT,
    stokki_erro        TEXT,
    stokki_tentativas  INTEGER NOT NULL DEFAULT 0,
    wms_recebimento_id INTEGER,
    observacoes        TEXT,
    enviado_por        TEXT,
    criado_em          TEXT NOT NULL,
    atualizado_em      TEXT NOT NULL,
    cancelado_em       TEXT,
    cancelado_por      TEXT,
    chamado_id         INTEGER
);
CREATE INDEX IF NOT EXISTS idx_portal_entradas_emb ON portal_entradas (cnpj_embarcador, criado_em);
CREATE INDEX IF NOT EXISTS idx_portal_entradas_status ON portal_entradas (status);
CREATE TABLE IF NOT EXISTS portal_entrada_itens (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    entrada_id  INTEGER NOT NULL REFERENCES portal_entradas(id),
    linha       INTEGER NOT NULL,
    sku         TEXT NOT NULL,
    ean         TEXT NOT NULL DEFAULT '',
    descricao   TEXT NOT NULL DEFAULT '',
    quantidade  REAL NOT NULL,
    unidade     TEXT NOT NULL DEFAULT '',
    UNIQUE (entrada_id, linha)
);
"""


def garantir_tabelas(conn: sqlite3.Connection) -> None:
    """Idempotente. Chamada por conectar() e também pelo timer do WMS
    (que abre a conexão dele em wms_pedidos.conectar) antes de amarrar."""
    conn.executescript(_DDL)
    if conn.execute("SELECT name FROM sqlite_master WHERE type = 'table' AND name = 'portal_clientes_envio'").fetchone():
        ep._garantir_colunas(conn, "portal_clientes_envio", {"entradas_ativo": "INTEGER NOT NULL DEFAULT 0"})
    conn.commit()


def conectar() -> sqlite3.Connection:
    """A conexão de envio_pedidos (cria portal_envios/portal_clientes_envio)
    mais as tabelas desta aba."""
    conn = ep.conectar()
    garantir_tabelas(conn)
    return conn


def _pasta_entradas() -> Path:
    # função (não constante) porque os testes trocam ep._RAIZ
    return ep._RAIZ / "dados" / "portal_entradas"


def cnpj_freshlog(config: dict | None) -> str:
    return _so_digitos(((config or {}).get("portal_entradas") or {}).get("cnpj_freshlog") or "")


# ── Leitura ────────────────────────────────────────────────────────────────────

def ler_nfe_entrada(conteudo: bytes, nome_arquivo: str = "") -> dict:
    """NF-e de remessa pra armazenagem (D1): aceita tpNF 0 ou 1; a
    validação de quem emitiu fica em validar_entrada."""
    nfe = ep.ler_nfe(conteudo, nome_arquivo, permitir_entrada=True)
    nfe["origem"] = ORIGEM_XML
    nfe["referencia"] = ""
    nfe["data_prevista"] = ""
    return nfe


def rotulo_entrada(e: dict) -> str:
    if e.get("numero_nf"):
        return f"NF {e['numero_nf']}"
    return f"Remessa {e.get('referencia') or '?'}"


def _quando_br(valor: str | None) -> str:
    if not valor:
        return ""
    try:
        return datetime.strptime(valor[:19], "%Y-%m-%d %H:%M:%S").strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return valor


def validar_entrada(conn: sqlite3.Connection, item: dict, cnpj_cliente: str, config: dict | None,
                    outras_empresas: dict[str, str] | None = None) -> dict:
    """{ok, erros[], avisos[], existente} -- as 5 regras da spec (7.1),
    menos a data prevista, que só existe na confirmação."""
    erros, avisos = [], []
    emb = _so_digitos(cnpj_cliente)
    if item.get("emitente_cnpj") != emb:
        do_grupo = (outras_empresas or {}).get(item.get("emitente_cnpj"))
        if do_grupo:
            erros.append(f"Essa nota é da {do_grupo} ({formatar_documento(item['emitente_cnpj'])}). "
                         f"Troque a empresa no seletor acima e envie de novo.")
        else:
            erros.append(f"O CNPJ emitente da nota ({formatar_documento(item.get('emitente_cnpj') or '')} · "
                         f"{item.get('emitente_nome') or ''}) não é o da sua empresa.")
    existente = conn.execute("SELECT id, status, criado_em, stokki_codigo FROM portal_entradas WHERE chave_nfe = ?",
                             (item["chave_nfe"],)).fetchone()
    if existente and existente["status"] != STATUS_CANCELADO:
        erros.append(f"{rotulo_entrada(item)} já anunciada em {_quando_br(existente['criado_em'])} "
                     f"({ROTULOS_STATUS.get(existente['status'], existente['status'])}"
                     + (f", {existente['stokki_codigo']}" if existente["stokki_codigo"] else "") + ").")
    elif existente:
        avisos.append("Já esteve anunciada e foi cancelada -- será reaberta.")
    fl = cnpj_freshlog(config)
    if fl and item.get("origem") == ORIGEM_XML and _so_digitos(item.get("destinatario_doc")) != fl:
        avisos.append(f"O destinatário da nota não é a Fresh Log ({item.get('destinatario_nome') or formatar_documento(item.get('destinatario_doc') or '')}). "
                      f"Confira se é mesmo uma remessa pro galpão.")
    try:
        cfg = ep.config_stokki_cliente(conn, emb, config)
    except ErroEnvio as e:
        erros.append(str(e))
        cfg = None
    if cfg is not None:
        erros_sku, avisos_sku = ep.validar_skus(conn, item, cfg)
        erros.extend(erros_sku)
        avisos.extend(avisos_sku)
    if not item.get("itens_lista"):
        erros.append("A nota não tem itens (<det>) -- não dá pra anunciar o que vai chegar.")
    return {"ok": not erros, "erros": erros, "avisos": avisos, "existente": dict(existente) if existente else None}


# ── Planilha de entrada ────────────────────────────────────────────────────────
# Modelo próprio: UMA LINHA POR ITEM; linhas com a mesma "Referência" formam
# uma remessa. Sem endereço/destinatário/valor: o destino é sempre o galpão.

COLUNAS_PLANILHA_ENTRADA: list[tuple[str, str, bool, tuple[str, ...], int, str]] = [
    ("referencia", "Referência (nº do pedido de compra/remessa)", True,
     ("referencia", "pedido", "remessa", "n pedido", "numero do pedido", "po", "ordem", "order"), 30, "REM-1001"),
    ("data_prevista", "Data prevista de chegada (DD/MM/AAAA)", True,
     ("data prevista", "data prevista de chegada", "chegada", "data de chegada", "previsao", "data"), 22, ""),
    ("sku", "SKU do produto", True, ("sku", "codigo", "codigo do produto", "cod produto", "produto", "ean", "gtin", "item", "cod"), 18, "NUU001FD"),
    ("quantidade", "Quantidade", True, ("quantidade", "qtd", "qtde", "qte", "quant"), 12, "10"),
    ("unidade", "Unidade (CX, UN...)", False, ("unidade", "un", "und", "embalagem"), 12, "CX"),
    ("numero_nf", "Nº da NF (opcional)", False, ("nf", "n nf", "numero nf", "numero da nf", "nota", "nota fiscal", "nfe"), 16, ""),
    ("volumes", "Volumes", False, ("volumes", "vol", "qtd volumes", "caixas"), 10, "3"),
    ("peso_kg", "Peso (kg)", False, ("peso", "peso kg", "peso (kg)", "peso bruto"), 10, "12,5"),
    ("observacoes", "Observações", False, ("observacoes", "observacao", "obs"), 30, ""),
]
_APELIDOS_ENTRADA: dict[str, str] = {}
for _c in COLUNAS_PLANILHA_ENTRADA:
    _APELIDOS_ENTRADA[ep._normalizar_texto(_c[1])] = _c[0]
    for _a in _c[3]:
        _APELIDOS_ENTRADA.setdefault(ep._normalizar_texto(_a), _c[0])
MAX_LINHAS_PLANILHA = 2000


def chave_planilha_entrada(cnpj_embarcador: str, referencia: str) -> str:
    """Chave sintética única por embarcador+referência (vai em chave_nfe só
    pra dedupe). Prefixo próprio pra nunca colidir com a de Envios."""
    ref = re.sub(r"[^A-Z0-9]+", "-", str(referencia or "").upper()).strip("-")[:40]
    digest = hashlib.sha1(f"entrada|{_so_digitos(cnpj_embarcador)}|{str(referencia or '').strip().upper()}".encode()).hexdigest()[:10]
    return f"PLANILHA-ENTRADA-{_so_digitos(cnpj_embarcador)}-{ref}-{digest}"


_MENSAGEM_CABECALHO_ENTRADA = ("Não achei o cabeçalho da planilha -- use o modelo Fresh Log de entrada (botão \"Baixar modelo\") "
                               "e mantenha a primeira linha com os nomes das colunas.")


def ler_planilha_entrada(conteudo: bytes, nome: str, cnpj_embarcador: str) -> tuple[list[dict], list[dict]]:
    """(remessas, rejeitados). Cada remessa tem o MESMO formato do dict de
    ler_nfe_entrada (origem, referencia, data_prevista, itens_lista, ...)."""
    linhas = ep._abrir_planilha(conteudo, nome)
    if not linhas:
        raise ErroEnvio(f"{nome}: a planilha está vazia.")
    i_cab, mapa = ep._achar_cabecalho(linhas, _APELIDOS_ENTRADA, _MENSAGEM_CABECALHO_ENTRADA)
    faltando = [c[1] for c in COLUNAS_PLANILHA_ENTRADA if c[2] and c[0] not in mapa.values()]
    if faltando:
        raise ErroEnvio(f"{nome}: faltam colunas obrigatórias no cabeçalho: {', '.join(faltando)}. "
                        f"Baixe o modelo Fresh Log de entrada pra conferir os nomes.")
    emb = _so_digitos(cnpj_embarcador)
    hoje = date.today()
    grupos: dict[str, dict] = {}
    erros_grupo: dict[str, list[str]] = {}
    rejeitados: list[dict] = []
    total = 0
    for n, linha in enumerate(linhas[i_cab + 1:], start=i_cab + 2):
        campos = {chave: (linha[j] if j < len(linha) else None) for j, chave in mapa.items()}
        if not any(ep._celula_texto(v) for v in campos.values()):
            continue
        total += 1
        if total > MAX_LINHAS_PLANILHA:
            raise ErroEnvio(f"{nome}: a planilha tem mais de {MAX_LINHAS_PLANILHA} linhas -- divida em arquivos menores.")
        ref = ep._celula_texto(campos.get("referencia"))
        if not ref:
            rejeitados.append({"arquivo": nome, "rotulo": f"linha {n}", "linha": n, "erro": "Sem a referência da remessa."})
            continue
        chave_grupo = ref.strip().upper()
        erros = erros_grupo.setdefault(chave_grupo, [])
        sku = ep._celula_texto(campos.get("sku"))
        try:
            qtd = ep._numero_br(campos.get("quantidade"))
        except ValueError:
            qtd = None
            erros.append(f"linha {n}: quantidade inválida ({ep._celula_texto(campos.get('quantidade'))}).")
        if not sku:
            erros.append(f"linha {n}: sem SKU.")
        elif qtd is None or qtd <= 0:
            erros.append(f"linha {n}: quantidade precisa ser maior que zero.")
        data_txt = ep._celula_texto(campos.get("data_prevista"))
        try:
            data_prev = ep._data_planilha(campos.get("data_prevista"))
        except ValueError:
            data_prev = None
            erros.append(f"linha {n}: data prevista inválida ({data_txt}) -- use DD/MM/AAAA.")
        g = grupos.get(chave_grupo)
        if g is None:
            if not data_prev and not data_txt:
                erros.append("Sem a data prevista de chegada (obrigatória).")
            elif data_prev and data_prev < hoje:
                erros.append(f"Data prevista {data_prev.strftime('%d/%m/%Y')} já passou.")
            volumes = peso = None
            try:
                volumes = ep._numero_br(campos.get("volumes"))
                peso = ep._numero_br(campos.get("peso_kg"))
            except ValueError:
                erros.append("Volumes ou peso com número inválido.")
            g = grupos[chave_grupo] = {
                "chave_nfe": chave_planilha_entrada(emb, ref), "origem": ORIGEM_PLANILHA, "referencia": ref[:60],
                "numero_nf": _so_digitos(campos.get("numero_nf"))[:20], "serie": "", "emitida_em": "",
                "emitente_cnpj": emb, "emitente_nome": "", "destinatario_doc": "", "destinatario_nome": "",
                "data_prevista": data_prev.isoformat() if data_prev else "",
                "volumes": int(volumes) if volumes else 0, "peso_kg": round(peso or 0.0, 3), "valor_nf": 0.0,
                "observacoes": ep._celula_texto(campos.get("observacoes"))[:500],
                "itens_lista": [], "linhas": [], "nome_arquivo": nome,
            }
        elif data_prev and g["data_prevista"] and data_prev.isoformat() != g["data_prevista"]:
            erros.append(f"linha {n}: data prevista diferente da primeira linha da remessa {ref}.")
        g["itens_lista"].append({"linha": len(g["itens_lista"]) + 1, "sku": sku, "ean": "",
                                 "descricao": "", "quantidade": qtd or 0,
                                 "unidade": ep._celula_texto(campos.get("unidade")).upper()[:10], "valor_unitario": 0.0})
        g["linhas"].append(n)

    pedidos = []
    for chave_grupo, g in grupos.items():
        erros = erros_grupo.get(chave_grupo) or []
        if not any(i["sku"] and i["quantidade"] > 0 for i in g["itens_lista"]):
            erros.append("Nenhum item válido (SKU + quantidade).")
        if erros:
            rejeitados.append({"arquivo": nome, "rotulo": f"Remessa {g['referencia']}", "linha": g["linhas"][0] if g["linhas"] else None,
                               "erro": " ".join(dict.fromkeys(erros))})
            continue
        g["itens"] = len(g["itens_lista"])
        if not g["volumes"]:
            g["volumes"] = 1
        pedidos.append(g)
    return pedidos, rejeitados


def gerar_modelo_planilha_entrada(nome_cliente: str = "") -> bytes:
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Entradas"
    cab = Font(bold=True, color="FFFFFF")
    fundo_ob, fundo_op = PatternFill("solid", fgColor="0EA575"), PatternFill("solid", fgColor="6B7280")
    for j, (chave, rotulo, obrig, _ap, larg, _ex) in enumerate(COLUNAS_PLANILHA_ENTRADA, start=1):
        c = ws.cell(row=1, column=j, value=rotulo)
        c.font, c.fill = cab, (fundo_ob if obrig else fundo_op)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = larg
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    amanha = (date.today() + timedelta(days=1)).strftime("%d/%m/%Y")
    exemplos = [
        {c[0]: c[5] for c in COLUNAS_PLANILHA_ENTRADA} | {"data_prevista": amanha},
        {"referencia": "REM-1001", "sku": "SKU-B", "quantidade": "4", "unidade": "UN"},
        {"referencia": "REM-1002", "data_prevista": amanha, "sku": "NUU001FD", "quantidade": "2", "unidade": "CX", "volumes": "1"},
    ]
    for i, ex in enumerate(exemplos, start=2):
        for j, col in enumerate(COLUNAS_PLANILHA_ENTRADA, start=1):
            v = ex.get(col[0], "")
            if v != "":
                ws.cell(row=i, column=j, value=v).font = Font(italic=True, color="9CA3AF")
    wi = wb.create_sheet("Instruções")
    wi.column_dimensions["A"].width = 110
    texto = [
        f"Modelo Fresh Log de anúncio de mercadoria (Pedidos de Entrada){(' · ' + nome_cliente) if nome_cliente else ''}",
        "",
        "1. Cada LINHA é um ITEM (SKU + quantidade). Linhas com a mesma \"Referência\" formam uma remessa só.",
        "2. Colunas em verde são obrigatórias; em cinza, opcionais. Mantenha os nomes do cabeçalho (a ordem pode mudar).",
        "3. Data prevista de chegada em DD/MM/AAAA, hoje ou futura -- é ela que organiza a fila do galpão.",
        "4. SKU é o código do produto cadastrado na Fresh Log/Stokki. Quantidade maior que zero.",
        "5. Não precisa de endereço nem destinatário: a mercadoria vem sempre pro galpão da Fresh Log.",
        "6. Apague as linhas de exemplo (em cinza) antes de enviar. Limite: 2.000 linhas por arquivo.",
        "",
        "Ao subir a planilha no portal, você vê uma prévia por remessa, confirma, e o recebimento é criado na Stokki automaticamente.",
    ]
    for i, t in enumerate(texto, start=1):
        c = wi.cell(row=i, column=1, value=t)
        if i == 1:
            c.font = Font(bold=True, size=13)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Confirmação ────────────────────────────────────────────────────────────────

def _data_prevista_valida(valor) -> str:
    try:
        d = date.fromisoformat(str(valor or ""))
    except ValueError:
        raise ErroEnvio("Informe a data prevista de chegada (AAAA-MM-DD).")
    if d < date.today():
        raise ErroEnvio(f"A data prevista {d.strftime('%d/%m/%Y')} já passou -- precisa ser hoje ou depois.")
    return d.isoformat()


def _caminho_definitivo_xml(cnpj: str, chave: str) -> Path:
    pasta = _pasta_entradas() / _so_digitos(cnpj)
    pasta.mkdir(parents=True, exist_ok=True)
    return pasta / f"{chave}.xml"


def _caminho_definitivo_planilha(cnpj: str, nome_original: str) -> Path:
    import secrets
    pasta = _pasta_entradas() / _so_digitos(cnpj) / "planilhas"
    pasta.mkdir(parents=True, exist_ok=True)
    seguro = re.sub(r"[^A-Za-z0-9._-]+", "_", Path(nome_original or "planilha.xlsx").name)[:80] or "planilha.xlsx"
    return pasta / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{secrets.token_hex(3)}_{seguro}"


def confirmar_entradas(conn: sqlite3.Connection, cnpj_embarcador: str, itens: list[dict], enviado_por: str,
                       config: dict | None) -> list[dict]:
    """Grava as entradas ANUNCIADO / NA_FILA a partir dos temporários da
    prévia. Revalida tudo (o catálogo pode ter mudado entre analisar e
    confirmar). Uma entrada que já existia CANCELADO é regravada na mesma
    linha (UPDATE), com os itens antigos apagados. Commit por entrada."""
    emb = _so_digitos(cnpj_embarcador)
    criados = []
    planilhas_movidas: dict[str, Path] = {}
    for item in itens:
        caminho = ep._caminho_temporario(item.get("token", ""))
        conteudo = caminho.read_bytes()
        if caminho.suffix == ".json":
            nfe = json.loads(conteudo.decode("utf-8"))
            if nfe.get("origem") != ORIGEM_PLANILHA or _so_digitos(nfe.get("emitente_cnpj")) != emb:
                raise ErroEnvio("Remessa temporária inválida -- envie a planilha de novo.")
        else:
            nfe = ler_nfe_entrada(conteudo)
        v = validar_entrada(conn, nfe, emb, config)
        if not v["ok"]:
            raise ErroEnvio(f"{rotulo_entrada(nfe)}: " + " ".join(v["erros"]))
        data_prevista = _data_prevista_valida(item.get("data_prevista") or nfe.get("data_prevista"))

        planilha = nfe.get("origem") == ORIGEM_PLANILHA
        if planilha:
            tk = nfe.get("arquivo_token", "")
            destino = planilhas_movidas.get(tk)
            if destino is None:
                origem_plan = ep._caminho_temporario(tk)
                destino = _caminho_definitivo_planilha(emb, nfe.get("nome_arquivo") or origem_plan.name)
                destino.write_bytes(origem_plan.read_bytes())
                planilhas_movidas[tk] = destino
                try:
                    origem_plan.unlink()
                except OSError:
                    pass
        else:
            destino = _caminho_definitivo_xml(emb, nfe["chave_nfe"])
            destino.write_bytes(conteudo)
        try:
            caminho.unlink()
        except OSError:
            pass

        agora = _agora()
        campos = {
            "cnpj_embarcador": emb, "origem": nfe["origem"], "chave_nfe": nfe["chave_nfe"],
            "numero_nf": nfe.get("numero_nf") or None, "serie": nfe.get("serie") or None,
            "emitida_em": nfe.get("emitida_em") or None, "referencia": nfe.get("referencia") or None,
            "data_prevista": data_prevista, "volumes": nfe.get("volumes") or 1, "peso_kg": nfe.get("peso_kg") or 0,
            "valor_nf": nfe.get("valor_nf") or 0, "arquivo_path": str(destino.relative_to(ep._RAIZ)),
            "status": STATUS_ANUNCIADO, "stokki_status": STOKKI_NA_FILA, "stokki_id": None, "stokki_codigo": None,
            "stokki_erro": None, "stokki_tentativas": 0, "wms_recebimento_id": None,
            "observacoes": (item.get("observacoes") or nfe.get("observacoes") or "")[:500],
            "enviado_por": enviado_por, "criado_em": agora, "atualizado_em": agora,
            "cancelado_em": None, "cancelado_por": None, "chamado_id": None,
        }
        existente = v["existente"]
        if existente:
            sets = ", ".join(f"{k} = ?" for k in campos if k != "chave_nfe")
            conn.execute(f"UPDATE portal_entradas SET {sets} WHERE id = ?",
                         (*[x for k, x in campos.items() if k != "chave_nfe"], existente["id"]))
            entrada_id = existente["id"]
            conn.execute("DELETE FROM portal_entrada_itens WHERE entrada_id = ?", (entrada_id,))
        else:
            cols = ", ".join(campos)
            conn.execute(f"INSERT INTO portal_entradas ({cols}) VALUES ({','.join('?' * len(campos))})", tuple(campos.values()))
            entrada_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
        for it in nfe.get("itens_lista") or []:
            conn.execute("INSERT INTO portal_entrada_itens (entrada_id, linha, sku, ean, descricao, quantidade, unidade) VALUES (?,?,?,?,?,?,?)",
                         (entrada_id, int(it["linha"]), str(it["sku"])[:60], str(it.get("ean") or "")[:20],
                          str(it.get("descricao") or "")[:200], float(it["quantidade"]), str(it.get("unidade") or "")[:10]))
        conn.commit()
        criados.append({"id": entrada_id, "numero_nf": nfe.get("numero_nf"), "referencia": nfe.get("referencia"),
                        "origem": nfe["origem"], "data_prevista": data_prevista, "status": STATUS_ANUNCIADO})
    return criados


# ── Listagem ───────────────────────────────────────────────────────────────────

def _tem_tabela(conn, nome: str) -> bool:
    return conn.execute("SELECT name FROM sqlite_master WHERE type = 'table' AND name = ?", (nome,)).fetchone() is not None


def _conferencia_do_galpao(conn, recebimento_id: int | None) -> list[dict]:
    """Anunciado × recebido × falta POR SKU, em UN, do wms_recebimento_itens
    ligado. Vazio sem ligação ou sem as tabelas do WMS (banco só do portal)."""
    if not recebimento_id or not _tem_tabela(conn, "wms_recebimento_itens"):
        return []
    if _tem_tabela(conn, "wms_produtos"):
        sql = ("SELECT i.sku, COALESCE(p.descricao, '') AS descricao, COALESCE(p.unidade, 'UN') AS unidade, "
               "SUM(COALESCE(i.qtd_un, 0)) AS anunciado, SUM(COALESCE(i.qtd_enderecada, 0)) AS recebido, "
               "SUM(COALESCE(i.falta_un, 0)) AS falta FROM wms_recebimento_itens i LEFT JOIN wms_produtos p ON p.id = i.produto_id "
               "WHERE i.recebimento_id = ? GROUP BY i.sku ORDER BY MIN(i.linha)")
    else:
        sql = ("SELECT i.sku, '' AS descricao, 'UN' AS unidade, SUM(COALESCE(i.qtd_un, 0)) AS anunciado, "
               "SUM(COALESCE(i.qtd_enderecada, 0)) AS recebido, SUM(COALESCE(i.falta_un, 0)) AS falta "
               "FROM wms_recebimento_itens i WHERE i.recebimento_id = ? GROUP BY i.sku ORDER BY MIN(i.linha)")
    saida = []
    for r in conn.execute(sql, (int(recebimento_id),)):
        saida.append({"sku": r["sku"], "descricao": r["descricao"], "unidade": r["unidade"],
                      "anunciado": round(float(r["anunciado"]), 3), "recebido": round(float(r["recebido"]), 3),
                      "falta": round(float(r["falta"]), 3)})
    return saida


def _linha(conn, r: sqlite3.Row, itens: dict[int, list[dict]]) -> dict:
    d = dict(r)
    hoje = date.today().isoformat()
    rec = None
    if d.get("wms_recebimento_id") and _tem_tabela(conn, "wms_recebimentos"):
        rec = conn.execute("SELECT * FROM wms_recebimentos WHERE id = ?", (d["wms_recebimento_id"],)).fetchone()
    d.update({
        "rotulo": rotulo_entrada(d),
        "status_rotulo": ROTULOS_STATUS.get(d["status"], d["status"]),
        "stokki_rotulo": ROTULOS_STOKKI.get(d["stokki_status"], d["stokki_status"]),
        "criado_em_br": _quando_br(d["criado_em"]),
        "data_prevista_br": "/".join(reversed(d["data_prevista"].split("-"))) if d.get("data_prevista") else "",
        "atrasada": d["status"] == STATUS_ANUNCIADO and bool(d.get("data_prevista")) and d["data_prevista"] < hoje,
        "itens_lista": itens.get(d["id"], []),
        "conferencia": _conferencia_do_galpao(conn, d.get("wms_recebimento_id")),
        "observacao_divergencia": (rec["observacao_divergencia"] if rec is not None and "observacao_divergencia" in rec.keys() else "") or "",
        "encerrado_em_br": _quando_br(rec["encerrado_em"]) if rec is not None and "encerrado_em" in rec.keys() else "",
        "pode_cancelar": d["status"] == STATUS_ANUNCIADO and d["stokki_status"] != STOKKI_ENVIANDO,
        "arquivo_ext": Path(d.get("arquivo_path") or "").suffix.lstrip(".").lower() or "xml",
    })
    d.pop("arquivo_path", None)
    return d


def listar_entradas(conn: sqlite3.Connection, cnpj_embarcador: str, dias: int = DIAS_LISTAGEM) -> list[dict]:
    desde = (datetime.now() - timedelta(days=dias)).strftime("%Y-%m-%d 00:00:00")
    emb = _so_digitos(cnpj_embarcador)
    rows = conn.execute("SELECT * FROM portal_entradas WHERE cnpj_embarcador = ? AND (criado_em >= ? OR status IN ('ANUNCIADO','CHEGOU')) "
                        "ORDER BY data_prevista, criado_em DESC, id DESC", (emb, desde)).fetchall()
    itens: dict[int, list[dict]] = {}
    if rows:
        ids = [r["id"] for r in rows]
        for it in conn.execute(f"SELECT * FROM portal_entrada_itens WHERE entrada_id IN ({','.join('?' * len(ids))}) ORDER BY entrada_id, linha", ids):
            itens.setdefault(it["entrada_id"], []).append({k: it[k] for k in ("linha", "sku", "ean", "descricao", "quantidade", "unidade")})
    return [_linha(conn, r, itens) for r in rows]


def resumo_entradas(lista: list[dict]) -> dict:
    hoje = date.today().strftime("%d/%m/%Y")
    return {
        "total": len(lista),
        "anunciados": sum(1 for e in lista if e["status"] == STATUS_ANUNCIADO),
        "atrasados": sum(1 for e in lista if e["atrasada"]),
        # "Chegaram": tudo que já está no galpão (conferindo, endereçado ou com
        # divergência) -- o nome da chave é histórico do desenho, o tile diz "no galpão"
        "chegaram_hoje": sum(1 for e in lista if e["status"] in (STATUS_CHEGOU, STATUS_ENDERECADO, STATUS_DIVERGENCIA)),
        "divergencias": sum(1 for e in lista if e["status"] == STATUS_DIVERGENCIA),
        "erros_stokki": sum(1 for e in lista if e["stokki_status"] == STOKKI_ERRO and e["status"] == STATUS_ANUNCIADO),
    }


def buscar_entrada(conn: sqlite3.Connection, entrada_id: int, cnpj_embarcador: str) -> dict | None:
    r = conn.execute("SELECT * FROM portal_entradas WHERE id = ? AND cnpj_embarcador = ?",
                     (entrada_id, _so_digitos(cnpj_embarcador))).fetchone()
    return dict(r) if r else None


def caminho_arquivo(entrada: dict) -> Path:
    return ep._RAIZ / entrada["arquivo_path"]


# ── Cancelar ───────────────────────────────────────────────────────────────────

def _chamados():
    import chamados as ch
    return ch


def _abrir_chamado_cancelamento(entrada: dict, cliente: dict, por: str) -> int | None:
    """Com #PE já criado, a Stokki não é cancelada sozinha (spec D5): abre
    chamado pra equipe, o mesmo desenho de bloqueio_area.abrir_bloqueio.
    Falha do chat não desfaz o cancelamento no portal."""
    try:
        ch = _chamados()
        conn_ch = ch.conectar()
        try:
            chamado = ch.criar_chamado(conn_ch, cliente, ch.ORIGEM_SISTEMA, ch.STATUS_AGUARDANDO_FL,
                                       assunto=f"Cancelar recebimento {entrada.get('stokki_codigo') or entrada.get('stokki_id')} na Stokki",
                                       area="entradas", pedido_ref=rotulo_entrada(entrada))
            quem = "pela equipe Fresh Log" if str(por).startswith("equipe") else "pelo cliente"
            ch.mensagem_sistema(conn_ch, chamado,
                                f"{rotulo_entrada(entrada)} foi cancelada {quem} na aba Pedidos de Entrada. O recebimento "
                                f"{entrada.get('stokki_codigo') or ''} já existe na Stokki e precisa ser cancelado lá à mão.")
            return chamado["id"]
        finally:
            conn_ch.close()
    except Exception as e:  # noqa: BLE001
        logger.warning(f"entradas: não abriu chamado do cancelamento da entrada {entrada.get('id')}: {e}")
        return None


def cancelar_entrada(conn: sqlite3.Connection, entrada: dict, por: str, cliente: dict, config: dict | None) -> dict:
    """{aplicado, mensagem, chamado_id}. Só a partir de ANUNCIADO e nunca
    enquanto o worker está criando o #PE."""
    if entrada["status"] != STATUS_ANUNCIADO:
        raise ErroEnvio("Só uma entrada ainda anunciada (que não chegou) pode ser cancelada.")
    if entrada["stokki_status"] == STOKKI_ENVIANDO:
        raise ErroEnvio("Essa entrada está sendo criada na Stokki agora -- tente de novo em alguns instantes.")
    agora = _agora()
    conn.execute("UPDATE portal_entradas SET status = ?, cancelado_em = ?, cancelado_por = ?, atualizado_em = ? WHERE id = ?",
                 (STATUS_CANCELADO, agora, por, agora, entrada["id"]))
    if entrada.get("wms_recebimento_id") and _tem_tabela(conn, "wms_recebimentos"):
        conn.execute("UPDATE wms_recebimentos SET estado = 'CANCELADO', atualizado_em = ? WHERE id = ? AND estado = 'ESPERADO' "
                     "AND NOT EXISTS (SELECT 1 FROM wms_recebimento_itens i WHERE i.recebimento_id = wms_recebimentos.id AND i.qtd_enderecada > 0)",
                     (agora, entrada["wms_recebimento_id"]))
    conn.commit()
    chamado_id = None
    if entrada.get("stokki_id"):
        chamado_id = _abrir_chamado_cancelamento(entrada, cliente, por)
        if chamado_id:
            conn.execute("UPDATE portal_entradas SET chamado_id = ? WHERE id = ?", (chamado_id, entrada["id"]))
            conn.commit()
        mensagem = ("Entrada cancelada. O recebimento já existia na Stokki: a Fresh Log vai cancelá-lo lá e você acompanha pelo chat."
                    if chamado_id else "Entrada cancelada. O recebimento já existia na Stokki -- avise a Fresh Log pelo chat pra cancelar lá.")
    else:
        mensagem = "Entrada cancelada -- não será criada na Stokki."
    return {"aplicado": True, "mensagem": mensagem, "chamado_id": chamado_id}


# ── Flag por cliente (D3: só o piloto) ─────────────────────────────────────────

def config_entradas_cliente(conn: sqlite3.Connection, cnpj_embarcador: str) -> dict:
    r = conn.execute("SELECT entradas_ativo FROM portal_clientes_envio WHERE cnpj = ?", (_so_digitos(cnpj_embarcador),)).fetchone()
    return {"entradas_ativo": bool(r["entradas_ativo"]) if r else False}


def definir_entradas_ativo(conn: sqlite3.Connection, cnpj_embarcador: str, ativo: bool) -> None:
    emb = _so_digitos(cnpj_embarcador)
    if not conn.execute("SELECT 1 FROM portal_clientes_envio WHERE cnpj = ?", (emb,)).fetchone():
        ep.definir_parametros_cliente(conn, emb)   # cria a linha com os padrões de Envios
    conn.execute("UPDATE portal_clientes_envio SET entradas_ativo = ?, atualizado_em = ? WHERE cnpj = ?", (int(bool(ativo)), _agora(), emb))
    conn.commit()


def entradas_ativas_para(conn: sqlite3.Connection, cnpjs: list[str]) -> bool:
    return any(config_entradas_cliente(conn, c)["entradas_ativo"] for c in cnpjs or [])
