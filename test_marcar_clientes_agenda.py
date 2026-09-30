# -*- coding: utf-8 -*-
"""
Testes do marcar_clientes_agenda.py (rotina + fila de autorizacao do Hugo).

Rodar: py -3.11 -m unittest test_marcar_clientes_agenda
"""
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

import marcar_clientes_agenda as mca


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.addCleanup(self._tmp.cleanup)
        pasta = Path(self._tmp.name)
        self.planilha = pasta / "BD_CLIENTES.xlsx"
        self.banco = pasta / "dados.db"

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Destinatário - Código", "Destinatário - Nome ", "ALERTAS CLIENTES"])
        ws.append([49749452000160, "VERO PANE", None])
        ws.append([21590391000111, "NUTRICAR", "AGENDAMENTO"])
        wb.save(self.planilha)

        conn = sqlite3.connect(self.banco)
        conn.execute(
            "CREATE TABLE agendamentos_pedido (pedido TEXT, cnpj_destinatario TEXT, "
            "nome_destinatario TEXT, cnpj_embarcador TEXT, status TEXT, data_agendada TEXT, origem TEXT)")
        conn.commit()
        conn.close()

        patch = mock.patch.object(mca, "DB_PATH", self.banco)
        patch.start()
        self.addCleanup(patch.stop)

    def _agendar(self, pedido, cnpj, nome, status="RESPONDIDO", data="30/09/2026", origem=None,
                 embarcador="11111111000111"):
        conn = sqlite3.connect(self.banco)
        conn.execute("INSERT INTO agendamentos_pedido VALUES (?, ?, ?, ?, ?, ?, ?)",
                     (pedido, cnpj, nome, embarcador, status, data, origem))
        conn.commit()
        conn.close()

    def _alertas(self):
        wb = openpyxl.load_workbook(self.planilha, read_only=True, data_only=True)
        alertas = {str(l[0]): l[2] for l in wb.active.iter_rows(min_row=2, values_only=True)}
        wb.close()
        return alertas

    def _situacoes(self):
        return {r["documento"]: r["situacao"] for r in mca.listar()}


class TestCandidatos(Base):
    def test_so_entra_quem_tem_data_informada_e_documento(self):
        self._agendar("PS-1", "49.749.452/0001-60", "VERO PANE")
        self._agendar("PS-2", "59708718000180", "REAL BREAD", status="PENDENTE", data=None)
        self._agendar("PS-3", "", "", origem="PLANILHA_NUU")
        self._agendar("PS-4", "49749452000160", "VERO PANE PRODUTOS", origem="PLANILHA_NUU")
        candidatos = mca.buscar_candidatos()
        self.assertEqual(list(candidatos), ["49749452000160"])
        self.assertEqual(candidatos["49749452000160"]["nome"], "VERO PANE")
        self.assertEqual(candidatos["49749452000160"]["pedidos"], [
            {"codigo": "PS-1", "data": "30/09/2026", "origem": "portal ou e-mail", "embarcador": "11111111000111"},
            {"codigo": "PS-4", "data": "30/09/2026", "origem": "planilha NUU", "embarcador": "11111111000111"},
        ])


class TestRotina(Base):
    def test_rotina_deixa_pendente_e_nao_grava_na_planilha(self):
        self._agendar("PS-1", "49749452000160", "VERO PANE")
        resultado = mca.processar(self.planilha, modo_teste=False)
        self.assertEqual(resultado["pendentes_novos"], ["49749452000160"])
        self.assertEqual(resultado["total_pendentes"], 1)
        self.assertIsNone(self._alertas()["49749452000160"])
        self.assertEqual(self._situacoes(), {"49749452000160": "PENDENTE"})

    def test_quem_ja_e_agenda_na_planilha_nao_vira_pendente(self):
        self._agendar("PS-1", "21590391000111", "NUTRICAR")
        resultado = mca.processar(self.planilha, modo_teste=False)
        self.assertEqual(resultado["pendentes_novos"], [])
        self.assertEqual(resultado["ja_marcados"], ["21590391000111"])
        self.assertEqual(self._situacoes(), {"21590391000111": "AUTORIZADO"})

    def test_segunda_rodada_nao_repete_o_pendente(self):
        self._agendar("PS-1", "49749452000160", "VERO PANE")
        mca.processar(self.planilha, modo_teste=False)
        self._agendar("PS-2", "49749452000160", "VERO PANE")
        self._agendar("PS-3", "59708718000180", "REAL BREAD")
        resultado = mca.processar(self.planilha, modo_teste=False)
        self.assertEqual(resultado["pendentes_novos"], ["59708718000180"])
        self.assertEqual(resultado["total_pendentes"], 2)

    def test_modo_teste_nao_registra_nada(self):
        self._agendar("PS-1", "49749452000160", "VERO PANE")
        resultado = mca.processar(self.planilha, modo_teste=True)
        self.assertEqual(resultado["pendentes_novos"], ["49749452000160"])
        self.assertEqual(self._situacoes(), {})

    def test_linhas_antigas_da_tabela_ganham_situacao(self):
        # antes de 30/09 a tabela so tinha documento/nome; "(não marcar)" era a exclusao do Hugo
        conn = sqlite3.connect(self.banco)
        conn.execute("CREATE TABLE clientes_agenda_marcados (documento TEXT PRIMARY KEY, nome TEXT, "
                     "tratado_em TEXT DEFAULT (datetime('now','localtime')))")
        conn.execute("INSERT INTO clientes_agenda_marcados (documento, nome) VALUES ('48178686000131', '(não marcar)')")
        conn.execute("INSERT INTO clientes_agenda_marcados (documento, nome) VALUES ('21590391000111', 'NUTRICAR')")
        conn.commit()
        conn.close()
        self.assertEqual(self._situacoes(), {"48178686000131": "NAO_MARCAR", "21590391000111": "AUTORIZADO"})
        self._agendar("PS-1", "48178686000131", "SABIA")
        self.assertEqual(mca.processar(self.planilha, modo_teste=False)["pendentes_novos"], [])


class TestDecisao(Base):
    def setUp(self):
        super().setUp()
        self._agendar("PS-1", "49749452000160", "VERO PANE")
        self._agendar("PS-2", "59708718000180", "REAL BREAD")
        mca.processar(self.planilha, modo_teste=False)

    def test_autorizar_grava_agenda_na_planilha(self):
        resultado = mca.decidir("49749452000160", autorizar=True, caminho_planilha=self.planilha)
        self.assertEqual(resultado["situacao"], "AUTORIZADO")
        self.assertEqual(self._alertas()["49749452000160"], "AGENDA")
        self.assertEqual(self._situacoes()["49749452000160"], "AUTORIZADO")
        self.assertEqual([p["documento"] for p in mca.listar("PENDENTE")], ["59708718000180"])

    def test_autorizar_quem_nao_esta_na_planilha_inclui_linha(self):
        mca.decidir("59708718000180", autorizar=True, caminho_planilha=self.planilha)
        self.assertEqual(self._alertas()["59708718000180"], "AGENDA")

    def test_nao_marcar_so_tira_da_fila(self):
        resultado = mca.decidir("49749452000160", autorizar=False, caminho_planilha=self.planilha)
        self.assertEqual(resultado["situacao"], "NAO_MARCAR")
        self.assertIsNone(self._alertas()["49749452000160"])
        # nova data informada nao traz o cliente de volta
        self._agendar("PS-9", "49749452000160", "VERO PANE")
        self.assertEqual(mca.processar(self.planilha, modo_teste=False)["pendentes_novos"], [])

    def test_decidir_quem_nao_esta_pendente_e_erro(self):
        mca.decidir("49749452000160", autorizar=False, caminho_planilha=self.planilha)
        with self.assertRaises(ValueError):
            mca.decidir("49749452000160", autorizar=True, caminho_planilha=self.planilha)
        with self.assertRaises(ValueError):
            mca.decidir("00000000000000", autorizar=True, caminho_planilha=self.planilha)

    def test_listar_traz_os_pedidos_que_motivaram(self):
        pendentes = mca.listar("PENDENTE")
        self.assertEqual(pendentes[0]["documento"], "49749452000160")
        self.assertEqual(pendentes[0]["pedidos"][0]["codigo"], "PS-1")
        self.assertEqual(pendentes[0]["documento_formatado"], "49.749.452/0001-60")


class TestNaoMarcarPelaLinhaDeComando(Base):
    def test_registra_como_nao_marcar(self):
        self.assertEqual(mca.registrar_nao_marcar(["49.749.452/0001-60", "abc"]), ["49749452000160"])
        self.assertEqual(self._situacoes(), {"49749452000160": "NAO_MARCAR"})
        self._agendar("PS-1", "49749452000160", "VERO PANE")
        self.assertEqual(mca.processar(self.planilha, modo_teste=False)["pendentes_novos"], [])


class TestLink(unittest.TestCase):
    def test_link_padrao_e_configuravel(self):
        self.assertEqual(mca.link_tela({}), "https://app.freshhub.com.br/painel/clientes-agenda")
        self.assertEqual(mca.link_tela({"painel_agentes": {"url_base": "http://x.test/painel/"}}),
                         "http://x.test/painel/clientes-agenda")


if __name__ == "__main__":
    unittest.main()
