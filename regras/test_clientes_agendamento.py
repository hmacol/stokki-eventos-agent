# -*- coding: utf-8 -*-
"""
Testes de regras/clientes_agendamento.py: leitura do alerta AGENDA e
marcacao automatica na BD_CLIENTES (marcar_agendamento).

Rodar: py -3.11 -m unittest regras.test_clientes_agendamento
"""
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent.parent))

from regras.clientes_agendamento import (
    carregar_clientes_agendamento, marcar_agendamento, tem_agendamento,
)

CABECALHO = ["Destinatário - Código", "Destinatário - Nome ", "Destinatário - Cidade ",
             "ALERTAS CLIENTES", "Classificação Dificuldade"]


class PlanilhaTemporaria(unittest.TestCase):
    def setUp(self):
        # a leitura (read_only) deixa a planilha aberta; no Windows isso
        # impede apagar a pasta no fim do teste
        self._tmp = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.addCleanup(self._tmp.cleanup)
        self.caminho = Path(self._tmp.name) / "BD_CLIENTES.xlsx"

    def _criar(self, linhas):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(CABECALHO)
        for linha in linhas:
            ws.append(linha)
        wb.save(self.caminho)

    def _linhas(self):
        wb = openpyxl.load_workbook(self.caminho, read_only=True, data_only=True)
        linhas = [list(l) for l in wb.active.iter_rows(min_row=2, values_only=True)]
        wb.close()
        return linhas

    def _backups(self):
        return list(Path(self._tmp.name).glob("BD_CLIENTES_backup_*_pre_agenda.xlsx"))


class TestLeitura(PlanilhaTemporaria):
    def test_codigo_numerico_sem_zero_a_esquerda_casa_com_cnpj_completo(self):
        # a planilha guarda o codigo como numero: 04972092003148 vira 4972092003148
        self._criar([[4972092003148, "GRUPO FARTURA", "SAO PAULO", "AGENDA", 2]])
        conjunto = carregar_clientes_agendamento(self.caminho)
        self.assertTrue(tem_agendamento("04972092003148", conjunto))
        self.assertTrue(tem_agendamento("04.972.092/0031-48", conjunto))

    def test_outro_alerta_nao_e_agendamento(self):
        self._criar([[21590391000111, "NUTRICAR", "SAO PAULO", "ENTREGA PRIORITÁRIA", 2]])
        conjunto = carregar_clientes_agendamento(self.caminho)
        self.assertFalse(tem_agendamento("21590391000111", conjunto))


class TestMarcarAgendamento(PlanilhaTemporaria):
    def test_alerta_vazio_recebe_agenda(self):
        self._criar([[49749452000160, "VERO PANE", "SAO PAULO", None, 2]])
        resultado = marcar_agendamento(self.caminho, {"49749452000160": "VERO PANE"})
        self.assertEqual(resultado["marcados"], ["49749452000160"])
        self.assertEqual(resultado["incluidos"], [])
        self.assertEqual(self._linhas()[0][3], "AGENDA")
        self.assertTrue(tem_agendamento("49749452000160", carregar_clientes_agendamento(self.caminho)))

    def test_outro_alerta_e_mantido(self):
        self._criar([[49749452000160, "VERO PANE", "SAO PAULO", "ENTREGA PRIORITÁRIA", 2]])
        marcar_agendamento(self.caminho, {"49749452000160": "VERO PANE"})
        self.assertEqual(self._linhas()[0][3], "ENTREGA PRIORITÁRIA / AGENDA")

    def test_cnpj_repetido_marca_todas_as_linhas(self):
        self._criar([
            [49749452000160, "VERO PANE", "SAO PAULO", None, 2],
            [11111111000111, "OUTRO", "SAO PAULO", None, 1],
            [49749452000160, "VERO PANE FILIAL", "OSASCO", None, 3],
        ])
        marcar_agendamento(self.caminho, {"49749452000160": "VERO PANE"})
        linhas = self._linhas()
        self.assertEqual([l[3] for l in linhas], ["AGENDA", None, "AGENDA"])

    def test_cnpj_que_nao_existe_ganha_linha_nova(self):
        self._criar([[11111111000111, "OUTRO", "SAO PAULO", None, 1]])
        resultado = marcar_agendamento(self.caminho, {"59708718000180": "REAL BREAD LTDA"})
        self.assertEqual(resultado["marcados"], [])
        self.assertEqual(resultado["incluidos"], ["59708718000180"])
        linhas = self._linhas()
        self.assertEqual(len(linhas), 2)
        self.assertEqual(linhas[1][:4], [59708718000180, "REAL BREAD LTDA", None, "AGENDA"])
        self.assertTrue(tem_agendamento("59708718000180", carregar_clientes_agendamento(self.caminho)))

    def test_cnpj_com_zero_a_esquerda_acha_a_linha_existente(self):
        self._criar([[4972092003148, "GRUPO FARTURA", "SAO PAULO", None, 2]])
        resultado = marcar_agendamento(self.caminho, {"04972092003148": "GRUPO FARTURA"})
        self.assertEqual(resultado["marcados"], ["04972092003148"])
        linhas = self._linhas()
        self.assertEqual(len(linhas), 1)
        self.assertEqual(linhas[0][3], "AGENDA")

    def test_quem_ja_tem_agenda_nao_muda_e_nao_gera_backup(self):
        self._criar([[21590391000111, "NUTRICAR", "SAO PAULO", "AGENDAMENTO", 2]])
        resultado = marcar_agendamento(self.caminho, {"21590391000111": "NUTRICAR"})
        self.assertEqual(resultado, {"marcados": [], "incluidos": [], "backup": None})
        self.assertEqual(self._linhas()[0][3], "AGENDAMENTO")
        self.assertEqual(self._backups(), [])

    def test_backup_guarda_a_planilha_de_antes(self):
        self._criar([[49749452000160, "VERO PANE", "SAO PAULO", None, 2]])
        resultado = marcar_agendamento(self.caminho, {"49749452000160": "VERO PANE"})
        backups = self._backups()
        self.assertEqual(len(backups), 1)
        self.assertEqual(Path(resultado["backup"]), backups[0])
        wb = openpyxl.load_workbook(backups[0], read_only=True, data_only=True)
        self.assertIsNone(list(wb.active.iter_rows(min_row=2, values_only=True))[0][3])
        wb.close()

    def test_sem_gravar_so_informa_quem_seria_marcado(self):
        self._criar([[49749452000160, "VERO PANE", "SAO PAULO", None, 2]])
        resultado = marcar_agendamento(
            self.caminho, {"49749452000160": "VERO PANE", "59708718000180": "REAL BREAD"}, gravar=False)
        self.assertEqual(resultado, {"marcados": ["49749452000160"], "incluidos": ["59708718000180"],
                                     "backup": None})
        self.assertEqual(self._linhas(), [[49749452000160, "VERO PANE", "SAO PAULO", None, 2]])
        self.assertEqual(self._backups(), [])

    def test_demais_colunas_ficam_intactas(self):
        self._criar([[49749452000160, "VERO PANE", "SAO PAULO", None, 2]])
        marcar_agendamento(self.caminho, {"49749452000160": "VERO PANE"})
        self.assertEqual(self._linhas()[0], [49749452000160, "VERO PANE", "SAO PAULO", "AGENDA", 2])


if __name__ == "__main__":
    unittest.main()
