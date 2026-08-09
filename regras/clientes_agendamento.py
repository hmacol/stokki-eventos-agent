# -*- coding: utf-8 -*-
"""
regras/clientes_agendamento.py

Cadastro de quais DESTINATÁRIOS (por CNPJ/CPF) exigem agendamento real
de entrega — usado para só preencher scheduled_start/scheduled_end no
VUUPT quando existe de fato uma exigência de agendamento (pedido do
Hugo, 29/07).

Fonte: a MESMA planilha de complexidade de entrega (BD_CLIENTES —
config.yaml complexidade_entrega.planilha / clientes_agendamento.planilha,
apontando pro mesmo arquivo), coluna "ALERTAS CLIENTES" (coluna O).

Essa coluna NÃO é um flag limpo — é um campo de alerta de texto livre
com vários tipos de conteúdo (confirmado com dado real, 29/07):
    'ENTREGA PRIORITÁRIA', 'REDE TAPI', 'AGENDAMENTO', 'AGENDADA', ...
Só uma fração dos clientes tem QUALQUER alerta (2 de 5773 tinham algo
relacionado a agendamento no arquivo de teste). Por isso a detecção é
por PALAVRA-CHAVE ("contém AGENDA", normalizado — sem acento/caixa),
não por valor exato — cobre 'AGENDAMENTO', 'AGENDADA', 'AGENDAR' etc.
sem precisar listar cada grafia.

Cliente cujo CNPJ/CPF não tem alerta de agendamento: NÃO tem
agendamento — scheduled_start/scheduled_end ficam de FORA do payload.

Duplicatas de CNPJ na planilha (confirmado ao carregar complexidade):
se QUALQUER linha do mesmo CNPJ tiver o alerta, o cliente é marcado
como tem_agendamento=True (OR entre as ocorrências — mais seguro do
que arriscar perder um agendamento real por causa de uma linha
duplicada/incompleta).
"""
import logging
import re
import unicodedata
from pathlib import Path

import openpyxl

logger = logging.getLogger(__name__)

COLUNAS_DOCUMENTO = {
    "CNPJCPF", "CPFCNPJ", "DOCUMENTO", "CNPJ", "CPF",
    "DESTINATARIOCODIGO", "CODIGO",
}
COLUNAS_ALERTA = {
    "ALERTASCLIENTES", "ALERTACLIENTE", "ALERTAS", "ALERTA",
}

# Palavra-chave normalizada (sem acento/caixa) que identifica um alerta
# de agendamento entre os vários tipos possíveis na coluna.
PALAVRA_CHAVE_AGENDAMENTO = "AGENDA"


def _so_digitos(s) -> str:
    return "".join(c for c in str(s or "") if c.isdigit())


def _normalizar_cabecalho(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def _normalizar_texto(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or ""))
    return "".join(c for c in s if not unicodedata.combining(c)).upper().strip()


def carregar_clientes_agendamento(caminho: str | Path) -> set[str]:
    """
    Lê a planilha (a mesma de complexidade de entrega) e retorna o
    conjunto de documentos (só dígitos) dos destinatários cujo campo
    "ALERTAS CLIENTES" contém uma palavra relacionada a agendamento.

    Caminho vazio ou arquivo inexistente: retorna set() — nenhum
    cliente tem agendamento (comportamento seguro por padrão).
    Coluna de alerta ausente: também retorna set() com aviso (mais
    seguro do que travar o pipeline por uma coluna opcional ausente —
    diferente da coluna de documento, que é obrigatória).
    """
    if not caminho:
        return set()

    caminho = Path(caminho)
    if not caminho.exists():
        logger.warning(
            f"Planilha de clientes com agendamento não encontrada: {caminho} — "
            f"nenhum pedido terá agendamento preenchido."
        )
        return set()

    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb.active

    linha_cabecalho = next(ws.iter_rows(min_row=1, max_row=1))
    cabecalho = [_normalizar_cabecalho(c.value) for c in linha_cabecalho]

    idx_doc = None
    for i, c in enumerate(cabecalho):
        if c in COLUNAS_DOCUMENTO:
            idx_doc = i
            break
    if idx_doc is None:
        raise ValueError(
            f"Planilha de clientes com agendamento ({caminho}) sem coluna de "
            f"CNPJ/CPF. Cabeçalho encontrado: {[c.value for c in linha_cabecalho]}"
        )

    idx_alerta = None
    for i, c in enumerate(cabecalho):
        if c in COLUNAS_ALERTA:
            idx_alerta = i
            break
    if idx_alerta is None:
        logger.warning(
            f"Planilha de clientes com agendamento ({caminho}) sem coluna de "
            f"alertas (ex: 'ALERTAS CLIENTES') — nenhum pedido terá agendamento "
            f"preenchido. Cabeçalho encontrado: {[c.value for c in linha_cabecalho]}"
        )
        return set()

    documentos: set[str] = set()
    linhas_com_alerta_nao_agendamento = 0
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if linha is None or all(v is None for v in linha):
            continue
        doc = _so_digitos(linha[idx_doc])
        alerta = _normalizar_texto(linha[idx_alerta])
        if not doc or not alerta:
            continue
        if PALAVRA_CHAVE_AGENDAMENTO in alerta:
            documentos.add(doc)
        else:
            linhas_com_alerta_nao_agendamento += 1

    if linhas_com_alerta_nao_agendamento:
        logger.info(
            f"{linhas_com_alerta_nao_agendamento} linha(s) com outro tipo de alerta "
            f"(não relacionado a agendamento) — ignoradas para este cadastro."
        )
    logger.info(f"Clientes com agendamento carregados: {len(documentos)} documento(s).")
    return documentos


def tem_agendamento(documento: str, conjunto: set[str]) -> bool:
    """True se o CNPJ/CPF do destinatário exige agendamento real de entrega."""
    doc = _so_digitos(documento)
    return bool(doc) and doc in conjunto
