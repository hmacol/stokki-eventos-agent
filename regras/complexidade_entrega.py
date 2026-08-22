# -*- coding: utf-8 -*-
"""
regras/complexidade_entrega.py

Nível de complexidade de entrega (1 a 4) por CNPJ/CPF do DESTINATÁRIO,
usado junto com o tipo de carga do embarcador (ver _buscar_dados_
embarcador_banco em pipeline.py, coluna `habilidade` da tabela `interno`)
para montar o nome da skill do VUUPT: "{TipoCarga}-{Nível}".

Substitui regras/classificacao_clientes.py (28/07) — aquele módulo
tratava tipo de carga e nível como uma coisa só vinda de uma planilha
por CNPJ/CPF do destinatário; Hugo corrigiu: o tipo de carga é do
EMBARCADOR (já disponível em interno.habilidade), só o nível de
complexidade é mesmo do destinatário. Arquivo antigo pode ser apagado.

Fonte: planilha Excel mantida pelo Hugo, com pelo menos estas colunas
(nomes tolerantes a variação de acento/espaço/caixa):
    CNPJ/CPF   -- documento do destinatário (com ou sem máscara)
    Nível      -- 1 (mais fácil) a 4 (possível mais de 3h de espera)

Destinatário cujo CPF não está na planilha: aplica NIVEL_PADRAO_CPF (1)
direto, sem exigir revisão — CPF (pessoa física) é considerado baixo
risco por padrão. Destinatário cujo CNPJ não está na planilha: aplica
NIVEL_PADRAO_CNPJ (2) como valor provisório e marca requer_revisao=True —
CNPJ exige classificação manual na planilha antes de confiar nesse nível
(pedido do Hugo, 10/08).
"""
import datetime
import logging
import re
import sqlite3
import unicodedata
from pathlib import Path

import openpyxl

logger = logging.getLogger(__name__)

_DB_PATH = Path(__file__).parent.parent / "dados" / "dados.db"

# Aplicados quando o documento do destinatário não está na planilha —
# nunca fica sem nível nenhum (mesmo espírito do padrão "Seco-1" original,
# aqui só a metade do nível). CPF usa o padrão mais baixo direto; CNPJ usa
# um padrão mais conservador (2) e fica marcado para revisão manual, já
# que carga jurídica tende a ter operação mais complexa (pedido do Hugo,
# 10/08).
NIVEL_PADRAO_CPF = 1
NIVEL_PADRAO_CNPJ = 2

NIVEIS_VALIDOS = {1, 2, 3, 4}

# Nomes de coluna aceitos, já normalizados (maiúsculas, sem acento, sem
# espaço/pontuação) — tolera variações comuns de como a planilha pode vir.
# Inclui os nomes reais da planilha BD_CLIENTES do Hugo (Destinatário -
# Código / Classificação Dificuldade) além de variações mais genéricas.
COLUNAS_DOCUMENTO = {
    "CNPJCPF", "CPFCNPJ", "DOCUMENTO", "CNPJ", "CPF",
    "DESTINATARIOCODIGO", "CODIGO",
}
COLUNAS_NIVEL = {
    "NIVEL", "NIVELDECOMPLEXIDADE", "NIVELCOMPLEXIDADE", "COMPLEXIDADE",
    "CLASSIFICACAODIFICULDADE", "DIFICULDADE",
}
# CEP é opcional — só usado para desempatar duplicatas (linha completa vs
# linha incompleta de exportações antigas). Ausência não é erro.
COLUNAS_CEP = {"CEP", "DESTINATARIOCEP"}

# Horário padrão de atendimento (recebimento) do destinatário — mesma
# planilha do nível, colunas "Destinatário - Horário de atendimento -
# início/fim" (Hugo, 22/08: não confundir com o agendamento pontual de
# UM pedido, que é editado por outra tela e grava scheduled_start/end na
# VUUPT -- isto aqui é o horário PADRÃO do cliente).
COLUNAS_HORARIO_INICIO = {
    "DESTINATARIOHORARIODEATENDIMENTOINICIO", "HORARIODEATENDIMENTOINICIO",
    "HORARIOATENDIMENTOINICIO", "HORARIOINICIO",
}
COLUNAS_HORARIO_FIM = {
    "DESTINATARIOHORARIODEATENDIMENTOFIM", "HORARIODEATENDIMENTOFIM",
    "HORARIOATENDIMENTOFIM", "HORARIOFIM",
}


def _so_digitos(s) -> str:
    return "".join(c for c in str(s or "") if c.isdigit())


def _e_cnpj(doc: str) -> bool:
    """CNPJ tem 14 dígitos; CPF tem 11. Fora disso não dá pra afirmar."""
    return len(doc) == 14


def _normalizar_cabecalho(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def _normalizar_nivel(v) -> int | None:
    """Aceita int, float (2.0) ou string ('2', '2.0') — None se inválido."""
    try:
        n = int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None
    return n if n in NIVEIS_VALIDOS else None


def _normalizar_horario(v) -> str | None:
    """Aceita datetime.time/datetime.datetime (célula Excel formatada como
    hora) ou string 'HH:MM'/'HH:MM:SS' — devolve sempre 'HH:MM', ou None
    se vazio/não reconhecível."""
    if v is None or v == "":
        return None
    if isinstance(v, (datetime.time, datetime.datetime)):
        return v.strftime("%H:%M")
    texto = str(v).strip()
    m = re.match(r"^(\d{1,2}):(\d{2})", texto)
    if not m:
        return None
    hora, minuto = int(m.group(1)), int(m.group(2))
    if not (0 <= hora <= 23 and 0 <= minuto <= 59):
        return None
    return f"{hora:02d}:{minuto:02d}"


def carregar_niveis(caminho: str | Path) -> dict[str, int]:
    """
    Lê a planilha de complexidade e retorna {documento_só_dígitos: nível}.

    Colunas detectadas por nome (tolerante a variação de acento/espaço/
    caixa). Linhas com documento ou nível inválido/vazio são ignoradas
    (contadas e avisadas no log), sem travar a carga do restante.

    Duplicatas de documento (confirmado na planilha real do Hugo — mesmo
    CNPJ aparece várias vezes, geralmente com linhas incompletas de
    exportações antigas misturadas com a linha completa atual):
      - Se todas as ocorrências têm o MESMO nível → usa direto, sem drama.
      - Se divergem: prioriza a(s) linha(s) com CEP preenchido (linha
        incompleta sem endereço é o resíduo, não a fonte confiável).
      - Se ainda divergir mesmo entre linhas com CEP (conflito genuíno —
        confirmado que existe, inclusive com o MESMO CEP e nível
        diferente, ou seja, não dá pra desempatar por endereço) → usa o
        MAIOR nível (mais conservador: melhor alocar tempo a mais do que
        a menos numa entrega difícil) e avisa no log para revisão manual.

    Caminho vazio ou arquivo inexistente: retorna {} (todo destinatário
    cai no padrão por tipo de documento — CPF ou CNPJ) — não é erro
    fatal, só um aviso no log.
    """
    if not caminho:
        return {}

    caminho = Path(caminho)
    if not caminho.exists():
        logger.warning(
            f"Planilha de complexidade de entrega não encontrada: {caminho} — "
            f"todos os destinatários usarão o nível padrão por tipo de documento "
            f"(CPF={NIVEL_PADRAO_CPF}, CNPJ={NIVEL_PADRAO_CNPJ})."
        )
        return {}

    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb.active

    linha_cabecalho = next(ws.iter_rows(min_row=1, max_row=1))
    cabecalho = [_normalizar_cabecalho(c.value) for c in linha_cabecalho]

    def _achar_coluna(candidatos):
        for i, c in enumerate(cabecalho):
            if c in candidatos:
                return i
        return None

    idx_doc = _achar_coluna(COLUNAS_DOCUMENTO)
    idx_niv = _achar_coluna(COLUNAS_NIVEL)
    idx_cep = _achar_coluna(COLUNAS_CEP)

    faltantes = [
        nome for nome, idx in [("CNPJ/CPF", idx_doc), ("Nível", idx_niv)]
        if idx is None
    ]
    if faltantes:
        raise ValueError(
            f"Planilha de complexidade ({caminho}) sem a(s) coluna(s): "
            f"{', '.join(faltantes)}. Cabeçalho encontrado: "
            f"{[c.value for c in linha_cabecalho]}"
        )

    # Coleta todas as ocorrências por documento primeiro (doc -> [(cep, nivel), ...])
    ocorrencias: dict[str, list[tuple[str, int]]] = {}
    ignoradas = 0
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if linha is None or all(v is None for v in linha):
            continue  # linha totalmente vazia (comum no fim de planilhas)

        doc   = _so_digitos(linha[idx_doc])
        nivel = _normalizar_nivel(linha[idx_niv])
        cep   = _so_digitos(linha[idx_cep]) if idx_cep is not None else ""

        if not doc or nivel is None:
            ignoradas += 1
            continue

        ocorrencias.setdefault(doc, []).append((cep, nivel))

    # Resolve cada documento para um único nível, tratando duplicatas
    mapa: dict[str, int] = {}
    conflitos_resolvidos = []
    for doc, ocs in ocorrencias.items():
        niveis_distintos = {n for _, n in ocs}
        if len(niveis_distintos) == 1:
            mapa[doc] = ocs[0][1]
            continue

        # Diverge: prioriza linhas com CEP preenchido
        com_cep = [(c, n) for c, n in ocs if c]
        niveis_com_cep = {n for _, n in com_cep}
        if len(niveis_com_cep) == 1:
            mapa[doc] = com_cep[0][1]
            continue

        # Conflito genuíno (mesmo entre linhas com CEP, ou nenhuma tem CEP)
        # -> mais conservador: usa o maior nível
        candidatos = niveis_com_cep if niveis_com_cep else niveis_distintos
        maior = max(candidatos)
        mapa[doc] = maior
        conflitos_resolvidos.append((doc, sorted(niveis_distintos), maior))

    if ignoradas:
        logger.warning(
            f"{ignoradas} linha(s) da planilha de complexidade ignorada(s) "
            f"(documento ou nível ausente/inválido)."
        )
    if conflitos_resolvidos:
        logger.warning(
            f"{len(conflitos_resolvidos)} documento(s) com nível conflitante entre "
            f"linhas duplicadas — resolvido pelo maior valor (mais conservador). "
            f"Primeiros exemplos: {conflitos_resolvidos[:5]}"
        )
    logger.info(f"Complexidade de entrega carregada: {len(mapa)} documento(s) únicos "
               f"({len(ocorrencias)} documento(s) na planilha, "
               f"{sum(len(v) for v in ocorrencias.values())} linha(s) válida(s) no total).")
    return mapa


def carregar_horarios(caminho: str | Path) -> dict[str, tuple[str, str]]:
    """
    Lê a MESMA planilha de complexidade e retorna {documento_só_dígitos:
    (horario_inicio, horario_fim)} das colunas "Destinatário - Horário de
    atendimento - início/fim" (horário PADRÃO de recebimento do cliente,
    Hugo 22/08 -- não confundir com agendamento pontual de um pedido).

    Diferente de carregar_niveis: colunas ausentes NÃO são erro fatal (só
    aviso no log, retorna {}) -- é uma leitura opcional, mais nova, e sua
    ausência não pode quebrar o carregamento do nível que já funciona.

    Duplicata de documento com horários divergentes: usa a janela mais
    ABRANGENTE (menor início, maior fim) -- mais conservador pra quem lê
    o horário (melhor considerar a janela mais larga do que arriscar
    achar que o cliente não recebe fora de um horário que na real recebe).
    """
    if not caminho:
        return {}

    caminho = Path(caminho)
    if not caminho.exists():
        return {}

    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb.active

    linha_cabecalho = next(ws.iter_rows(min_row=1, max_row=1))
    cabecalho = [_normalizar_cabecalho(c.value) for c in linha_cabecalho]

    def _achar_coluna(candidatos):
        for i, c in enumerate(cabecalho):
            if c in candidatos:
                return i
        return None

    idx_doc = _achar_coluna(COLUNAS_DOCUMENTO)
    idx_ini = _achar_coluna(COLUNAS_HORARIO_INICIO)
    idx_fim = _achar_coluna(COLUNAS_HORARIO_FIM)

    if idx_doc is None or idx_ini is None or idx_fim is None:
        logger.warning(
            f"Planilha de complexidade ({caminho}) sem coluna(s) de horário de "
            f"atendimento -- todo destinatário usará o horário padrão (00:00-23:59)."
        )
        return {}

    ocorrencias: dict[str, list[tuple[str, str]]] = {}
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if linha is None or all(v is None for v in linha):
            continue

        doc = _so_digitos(linha[idx_doc])
        inicio = _normalizar_horario(linha[idx_ini])
        fim = _normalizar_horario(linha[idx_fim])
        if not doc or inicio is None or fim is None:
            continue

        ocorrencias.setdefault(doc, []).append((inicio, fim))

    mapa: dict[str, tuple[str, str]] = {}
    for doc, ocs in ocorrencias.items():
        inicio = min(i for i, _ in ocs)
        fim = max(f for _, f in ocs)
        mapa[doc] = (inicio, fim)

    logger.info(f"Horário de atendimento carregado: {len(mapa)} documento(s) únicos.")
    return mapa


def classificar_nivel(documento: str, mapa: dict[str, int]) -> tuple[int, bool, bool]:
    """
    Retorna (nível, encontrado_na_planilha, requer_revisao).

    Se o documento estiver no mapa, usa o nível da planilha e
    requer_revisao=False.

    Se não estiver: aplica o padrão conforme o tipo de documento. Para
    CPF, usa NIVEL_PADRAO_CPF (1) e isso é o suficiente
    (requer_revisao=False). Para CNPJ, usa NIVEL_PADRAO_CNPJ (2) e marca
    requer_revisao=True — precisa de classificação manual na planilha
    antes de confiar nesse nível.
    """
    doc = _so_digitos(documento)
    if doc and doc in mapa:
        return mapa[doc], True, False
    if _e_cnpj(doc):
        return NIVEL_PADRAO_CNPJ, False, True
    return NIVEL_PADRAO_CPF, False, False


# ---------------------------------------------------------------------
# Ajuste manual (Hugo, 22/08): correção pontual de nível/horário de UM
# cliente direto pela tela de Planejamento (menu de contexto de um
# pedido), sem precisar editar a planilha Excel. Guardado por
# DOCUMENTO (mesma chave da planilha) -- assim a correção vale pra
# TODOS os pedidos pendentes/futuros daquele cliente, não só o pedido
# em que se clicou. Sempre tem prioridade sobre a planilha.
# ---------------------------------------------------------------------

def _conectar_ajustes() -> sqlite3.Connection:
    _DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(_DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("""
        CREATE TABLE IF NOT EXISTS ajustes_complexidade_cliente (
            documento                   TEXT PRIMARY KEY,
            nivel_dificuldade           INTEGER NOT NULL,
            horario_atendimento_inicio  TEXT NOT NULL,
            horario_atendimento_fim     TEXT NOT NULL,
            atualizado_em               TEXT NOT NULL
        )
    """)
    conn.commit()
    return conn


def carregar_ajustes_manuais() -> dict[str, dict]:
    """{documento: {"nivel_dificuldade": int, "horario_atendimento_inicio":
    str, "horario_atendimento_fim": str}} de todos os ajustes manuais já
    feitos pela tela de Planejamento."""
    conn = _conectar_ajustes()
    try:
        return {
            row["documento"]: {
                "nivel_dificuldade": row["nivel_dificuldade"],
                "horario_atendimento_inicio": row["horario_atendimento_inicio"],
                "horario_atendimento_fim": row["horario_atendimento_fim"],
            }
            for row in conn.execute("SELECT * FROM ajustes_complexidade_cliente")
        }
    finally:
        conn.close()


def definir_ajuste_manual(documento: str, nivel: int, horario_inicio: str, horario_fim: str) -> None:
    """Grava (upsert) o ajuste manual de um cliente -- opção "Nível /
    horário de atendimento" do menu de contexto da tela de Planejamento."""
    doc = _so_digitos(documento)
    if not doc:
        raise ValueError("Documento (CNPJ/CPF) vazio -- não dá pra gravar ajuste manual sem saber de qual cliente.")
    if nivel not in NIVEIS_VALIDOS:
        raise ValueError(f"Nível inválido: {nivel!r} (válidos: {sorted(NIVEIS_VALIDOS)}).")

    conn = _conectar_ajustes()
    try:
        conn.execute("""
            INSERT INTO ajustes_complexidade_cliente (
                documento, nivel_dificuldade, horario_atendimento_inicio,
                horario_atendimento_fim, atualizado_em
            ) VALUES (?, ?, ?, ?, datetime('now', 'localtime'))
            ON CONFLICT(documento) DO UPDATE SET
                nivel_dificuldade = excluded.nivel_dificuldade,
                horario_atendimento_inicio = excluded.horario_atendimento_inicio,
                horario_atendimento_fim = excluded.horario_atendimento_fim,
                atualizado_em = excluded.atualizado_em
        """, (doc, nivel, horario_inicio, horario_fim))
        conn.commit()
    finally:
        conn.close()


def nivel_efetivo(documento: str, mapa_niveis: dict[str, int], ajustes_manuais: dict[str, dict]) -> int:
    """Nível "de verdade" a usar pro documento: ajuste manual > planilha
    > padrão por tipo de documento (ver classificar_nivel)."""
    doc = _so_digitos(documento)
    ajuste = ajustes_manuais.get(doc)
    if ajuste:
        return ajuste["nivel_dificuldade"]
    nivel, _, _ = classificar_nivel(documento, mapa_niveis)
    return nivel


def horario_efetivo(documento: str, mapa_horarios: dict[str, tuple[str, str]],
                     ajustes_manuais: dict[str, dict]) -> tuple[str, str]:
    """Horário de atendimento "de verdade" a usar pro documento: ajuste
    manual > planilha > padrão fixo (00:00-23:59, mesmo piso que a VUUPT
    já usa pro horário de atendimento importado)."""
    doc = _so_digitos(documento)
    ajuste = ajustes_manuais.get(doc)
    if ajuste:
        return ajuste["horario_atendimento_inicio"], ajuste["horario_atendimento_fim"]
    return mapa_horarios.get(doc, ("00:00", "23:59"))
