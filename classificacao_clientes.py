# -*- coding: utf-8 -*-
"""
regras/classificacao_clientes.py

Classificação de clientes por CNPJ/CPF, usada para determinar a skill do
serviço no VUUPT: tipo de carga (Congelado/Seco/Refrigerado) + nível de
complexidade de entrega (1 a 4).

Fonte: planilha Excel mantida pelo Hugo, com pelo menos estas colunas
(nomes tolerantes a variação de acento/espaço/caixa):
    CNPJ/CPF        -- documento do cliente (com ou sem máscara)
    Tipo de Carga   -- Congelado | Seco | Refrigerado
    Nível           -- 1 (mais fácil) a 4 (possível mais de 3h de espera)

O nome da skill no VUUPT segue o padrão confirmado pelo Hugo:
    "{TipoCarga}-{Nível}"   ex: "Congelado-3", "Seco-1", "Refrigerado-4"

Cliente cujo CNPJ/CPF não está na planilha: aplica a classificação
padrão CLASSIFICACAO_PADRAO ("Seco-1", confirmado com o Hugo) — nunca
fica sem skill nenhuma.
"""
import logging
import re
import unicodedata
from pathlib import Path

import openpyxl

logger = logging.getLogger(__name__)

# Aplicada quando o CNPJ/CPF do cliente não está na planilha (confirmado
# com o Hugo — nunca deixar o serviço sem classificação nenhuma).
CLASSIFICACAO_PADRAO = "Seco-1"

TIPOS_CARGA_VALIDOS = {"CONGELADO", "SECO", "REFRIGERADO"}
NIVEIS_VALIDOS = {1, 2, 3, 4}

# Nomes de coluna aceitos, já normalizados (maiúsculas, sem acento, sem
# espaço/pontuação) — tolera variações comuns de como a planilha pode vir.
COLUNAS_DOCUMENTO = {"CNPJCPF", "CPFCNPJ", "DOCUMENTO", "CNPJ", "CPF"}
COLUNAS_TIPO_CARGA = {"TIPODECARGA", "TIPOCARGA", "TIPO", "CARGA"}
COLUNAS_NIVEL = {"NIVEL", "NIVELDECOMPLEXIDADE", "NIVELCOMPLEXIDADE", "COMPLEXIDADE"}


def _so_digitos(s) -> str:
    return "".join(c for c in str(s or "") if c.isdigit())


def _normalizar_cabecalho(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def _normalizar_tipo_carga(s) -> str:
    """Normaliza para um dos 3 tipos válidos (capitalizado) ou "" se inválido."""
    s = unicodedata.normalize("NFKD", str(s or "").strip())
    s = "".join(c for c in s if not unicodedata.combining(c)).upper()
    return s.capitalize() if s in TIPOS_CARGA_VALIDOS else ""


def _normalizar_nivel(v) -> int | None:
    """Aceita int, float (2.0) ou string ('2', '2.0') — None se inválido."""
    try:
        n = int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None
    return n if n in NIVEIS_VALIDOS else None


def carregar_classificacao(caminho: str | Path) -> dict[str, str]:
    """
    Lê a planilha de classificação e retorna {documento_so_digitos: nome_skill}.

    Colunas detectadas por nome (tolerante a variação de acento/espaço/
    caixa) — não depende de posição fixa. Linhas com documento, tipo de
    carga ou nível inválido/vazio são ignoradas (contadas e avisadas no
    log), sem travar a carga do restante da planilha.

    Caminho vazio ou arquivo inexistente: retorna {} (todo cliente cai
    na classificação padrão) — não é erro fatal, só um aviso no log.
    """
    if not caminho:
        return {}

    caminho = Path(caminho)
    if not caminho.exists():
        logger.warning(
            f"Planilha de classificação não encontrada: {caminho} — todos "
            f"os clientes usarão a classificação padrão ({CLASSIFICACAO_PADRAO})."
        )
        return {}

    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = wb.active

    linha_cabecalho = next(ws.iter_rows(min_row=1, max_row=1))
    cabecalho = [_normalizar_cabecalho(c.value) for c in linha_cabecalho]

    def _achar_coluna(candidatos):
        for i, c in enumerate(cabecalho):
            if c in candidatos:
                return i
        return None

    idx_doc  = _achar_coluna(COLUNAS_DOCUMENTO)
    idx_tipo = _achar_coluna(COLUNAS_TIPO_CARGA)
    idx_niv  = _achar_coluna(COLUNAS_NIVEL)

    faltantes = [
        nome for nome, idx in
        [("CNPJ/CPF", idx_doc), ("Tipo de Carga", idx_tipo), ("Nível", idx_niv)]
        if idx is None
    ]
    if faltantes:
        raise ValueError(
            f"Planilha de classificação ({caminho}) sem a(s) coluna(s): "
            f"{', '.join(faltantes)}. Cabeçalho encontrado: "
            f"{[c.value for c in linha_cabecalho]}"
        )

    mapa: dict[str, str] = {}
    ignoradas = 0
    duplicadas = 0
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if linha is None or all(v is None for v in linha):
            continue  # linha totalmente vazia (comum no fim de planilhas)

        doc   = _so_digitos(linha[idx_doc])
        tipo  = _normalizar_tipo_carga(linha[idx_tipo])
        nivel = _normalizar_nivel(linha[idx_niv])

        if not doc or not tipo or nivel is None:
            ignoradas += 1
            continue

        if doc in mapa:
            duplicadas += 1  # última ocorrência prevalece — mesmo padrão de "último vence"

        mapa[doc] = f"{tipo}-{nivel}"

    if ignoradas:
        logger.warning(
            f"{ignoradas} linha(s) da planilha de classificação ignorada(s) "
            f"(documento/tipo de carga/nível ausente ou inválido)."
        )
    if duplicadas:
        logger.warning(
            f"{duplicadas} documento(s) duplicado(s) na planilha de "
            f"classificação — a última linha de cada um prevaleceu."
        )
    logger.info(f"Classificação de clientes carregada: {len(mapa)} documento(s).")
    return mapa


def classificar(documento: str, mapa: dict[str, str]) -> tuple[str, bool]:
    """
    Retorna (nome_skill, encontrado_na_planilha).

    Se o documento vier vazio ou não estiver no mapa, aplica
    CLASSIFICACAO_PADRAO e retorna encontrado_na_planilha=False — nunca
    retorna sem uma skill (o serviço nunca fica sem classificação).
    """
    doc = _so_digitos(documento)
    if doc and doc in mapa:
        return mapa[doc], True
    return CLASSIFICACAO_PADRAO, False
