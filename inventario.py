"""
inventario.py

Lista todos os pedidos que seriam importados pelo pipeline sem
buscar o detalhe de nenhum -- apenas a listagem do Stokki.
Gera um Excel com todos os pedidos por fonte e embarcador.

Execute:
  py -3.11 inventario.py
  py -3.11 inventario.py --embarcador 86
"""
import argparse, logging, re, sqlite3, sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import yaml

_RAIZ = Path(__file__).parent
sys.path.insert(0, str(_RAIZ))

(_RAIZ / "dados").mkdir(parents=True, exist_ok=True)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("inventario")

from stokki.auth import StokkiSession
from stokki import pedidos as stokki_pedidos
from pipeline import (
    CONFIG_PATH, TRANSPORTADORAS, DB_PATH,
    STATUS_AGUARDANDO_TRANSPORTADOR, STATUSES_EM_ABERTO,
    EMBARCADORES_IMPORTAR_ABERTOS,
    _buscar_dados_embarcador_banco,
)
from fingerprint_importacao import DB_PATH as FP_DB_PATH
from regras.transportadoras import CatalogoTransportadoras


def _embarcador_prioritario(linha) -> bool:
    """True se a linha pertence a um embarcador prioritario."""
    import unicodedata
    def norm(s):
        s = re.sub(r"<[^>]+>", "", s).upper().strip()
        s = unicodedata.normalize("NFKD", s)
        return "".join(c for c in s if not unicodedata.combining(c))
    nome = norm(str(linha.get("client", "") if isinstance(linha, dict) else ""))
    return any(
        norm(n) in nome or nome in norm(n)
        for n in EMBARCADORES_IMPORTAR_ABERTOS.values()
    )


def _codigos_ja_importados() -> set:
    """
    Retorna o conjunto de codigos PS que ja foram importados no VUUPT
    (presentes na tabela fingerprint_importacao_vuupt).
    """
    if not FP_DB_PATH.exists():
        return set()
    try:
        conn = sqlite3.connect(FP_DB_PATH)
        rows = conn.execute("SELECT codigo FROM fingerprint_importacao_vuupt").fetchall()
        conn.close()
        return {r[0] for r in rows}
    except Exception as e:
        logger.warning(f"Nao foi possivel ler fingerprint: {e}")
        return set()

COR_HEADER = "141428"
COR_ACENTO = "00C896"
CORES_FONTE = {
    "Prioritario":          "E6FBF5",
    "Aguardando Transp.":   "EEF2F7",
    "Estacao Impressao":    "FFF3E0",
}


def _limpar(html: str) -> str:
    return re.sub(r"<[^>]+>", " ", html).strip().split("\n")[0].strip()


def _extrair_campos(linha) -> dict:
    if not isinstance(linha, dict):
        return {}
    return {
        "id_stokki":   stokki_pedidos.extrair_id_da_linha(linha),
        "codigo_ps":   stokki_pedidos.extrair_codigo_ps_da_linha(linha),
        "referencia":  stokki_pedidos.extrair_referencia_da_linha(linha),
        "embarcador":  _limpar(str(linha.get("client", "")))[:50],
        "destino":     _limpar(str(linha.get("destination", "")))[:50],
        "transportadora": _limpar(str(linha.get("carrier", "")))[:40],
        "data":        str(linha.get("expedition_date", "")),
        "marker":      _limpar(str(linha.get("marker", "")))[:30],
        "state":       _limpar(str(linha.get("state", "")))[:20],
        "situacao":    "",  # preenchido apos cruzamento com fingerprint
        "tipo_transp": "",  # preenchido apos cruzamento com catalogo
    }


def coletar_pedidos(sess, filtro_cliente=""):
    todos: dict[int, dict] = {}  # id_stokki -> dados

    # Fonte 1: Embarcadores prioritarios
    for id_emb, nome_emb in EMBARCADORES_IMPORTAR_ABERTOS.items():
        cliente = filtro_cliente or id_emb
        if filtro_cliente and filtro_cliente != id_emb:
            continue
        for status in STATUSES_EM_ABERTO:
            for linha in stokki_pedidos.iterar_todos_pedidos(
                sess, status=status, cliente=cliente, pausa_entre_paginas=0.2
            ):
                campos = _extrair_campos(linha)
                id_ = campos.get("id_stokki")
                if id_ and id_ not in todos:
                    campos["fonte"] = "Prioritario"
                    todos[id_] = campos
        if filtro_cliente:
            break

    if not filtro_cliente:
        # Fonte 2: Aguardando Transportador (outros embarcadores)
        for linha in stokki_pedidos.iterar_todos_pedidos(
            sess, status=STATUS_AGUARDANDO_TRANSPORTADOR, pausa_entre_paginas=0.2
        ):
            campos = _extrair_campos(linha)
            id_ = campos.get("id_stokki")
            if id_ and id_ not in todos and not _embarcador_prioritario(linha):
                campos["fonte"] = "Aguardando Transp."
                todos[id_] = campos

    logger.info(f"Total coletado: {len(todos)} pedido(s)")

    # Carrega catalogo de transportadoras para identificar tipo
    catalogo = CatalogoTransportadoras.carregar(TRANSPORTADORAS)

    # Cruza com fingerprint e catalogo
    ja_importados = _codigos_ja_importados()
    logger.info(f"Fingerprint: {len(ja_importados)} codigos ja importados anteriormente")

    for campos in todos.values():
        codigo = campos.get("codigo_ps", "")
        campos["situacao"] = "Ja importado" if codigo in ja_importados else "Novo"
        # Resolve tipo da transportadora
        nome_transp = campos.get("transportadora", "")
        if nome_transp:
            r = catalogo.resolver(nome_transp)
            campos["tipo_transp"] = r.tipo or ""

    # Filtra RETIRADA -- nao vao pro VUUPT
    retirada = [c for c in todos.values() if c.get("tipo_transp") == "RETIRADA"]
    if retirada:
        logger.info(f"Ignorados (RETIRADA): {len(retirada)} pedido(s)")
        todos = {k: v for k, v in todos.items() if v.get("tipo_transp") != "RETIRADA"}

    return list(todos.values())


def gerar_excel(pedidos: list, caminho: Path, somente_novos: bool = False):
    if somente_novos:
        pedidos = [p for p in pedidos if p.get("situacao") == "Novo"]
        logger.info(f"Filtrado: {len(pedidos)} pedido(s) novos (ainda nao importados)")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario"

    headers = [
        ("codigo_ps",      "Pedido",        12),
        ("situacao",       "Situacao",       13),
        ("referencia",     "Referencia",    14),
        ("data",           "Data",          11),
        ("fonte",          "Fonte",         18),
        ("embarcador",     "Embarcador",    35),
        ("destino",        "Destino",       35),
        ("transportadora", "Transportadora",30),
        ("state",          "Status",        18),
        ("marker",         "Marker",        25),
    ]

    # Titulo
    ws.merge_cells(f"A1:{get_column_letter(len(headers))}1")
    cel = ws["A1"]
    cel.value     = f"Inventario VUUPT — {datetime.now().strftime('%d/%m/%Y %H:%M')} — {len(pedidos)} pedido(s)"
    cel.font      = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    cel.fill      = PatternFill("solid", fgColor=COR_ACENTO)
    cel.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Cabecalhos
    for col, (_, titulo, largura) in enumerate(headers, 1):
        c = ws.cell(row=2, column=col, value=titulo)
        c.font      = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill      = PatternFill("solid", fgColor=COR_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = largura
    ws.row_dimensions[2].height = 22

    # Dados
    for i, p in enumerate(sorted(pedidos, key=lambda x: (x.get("fonte",""), x.get("embarcador",""))), 3):
        cor = CORES_FONTE.get(p.get("fonte",""), "FFFFFF")
        fill = PatternFill("solid", fgColor=cor)
        for col, (campo, _, _) in enumerate(headers, 1):
            val = p.get(campo, "")
            c = ws.cell(row=i, column=col, value=val)
            c.font      = Font(name="Arial", size=10)
            c.fill      = fill
            c.alignment = Alignment(vertical="center")
            # Coluna Situacao: verde para Novo, cinza para Ja importado
            if campo == "situacao":
                if val == "Novo":
                    c.font = Font(name="Arial", size=10, color="1B5E3F", bold=True)
                else:
                    c.font = Font(name="Arial", size=10, color="888888")

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{len(pedidos)+2}"

    # Aba resumo
    ws2 = wb.create_sheet("Resumo")
    ws2["A1"] = "Fonte"
    ws2["B1"] = "Qtd"
    for c in ["A1","B1"]:
        ws2[c].font = Font(name="Arial", bold=True, color="FFFFFF")
        ws2[c].fill = PatternFill("solid", fgColor=COR_HEADER)
    contagem = defaultdict(int)
    for p in pedidos:
        contagem[p.get("fonte","?")] += 1
    for i, (fonte, qtd) in enumerate(sorted(contagem.items()), 2):
        ws2.cell(row=i, column=1, value=fonte).font = Font(name="Arial", size=10)
        ws2.cell(row=i, column=2, value=qtd).font   = Font(name="Arial", size=10)
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 8

    wb.save(caminho)
    logger.info(f"Excel salvo em: {caminho.resolve()}")


def main(filtro_embarcador="", somente_novos=False):
    config = yaml.safe_load(open(CONFIG_PATH, encoding="utf-8")) or {}
    sess   = StokkiSession(config)

    filtro_cliente = ""
    if filtro_embarcador:
        if filtro_embarcador.isdigit():
            filtro_cliente = filtro_embarcador
        else:
            import unicodedata as _ud
            def norm(s):
                s = s.upper()
                s = _ud.normalize("NFKD", s)
                return "".join(c for c in s if not _ud.combining(c))
            conn = sqlite3.connect(DB_PATH)
            conn.row_factory = sqlite3.Row
            rows = conn.execute("SELECT stkkc_id, nome_remetente FROM interno WHERE stkkc_id IS NOT NULL").fetchall()
            conn.close()
            busca = norm(filtro_embarcador)
            match = next((r for r in rows if busca in norm(r["nome_remetente"] or "")), None)
            if match:
                filtro_cliente = str(match["stkkc_id"])
                logger.info(f"Embarcador: {match['nome_remetente']} (stkkc_id={filtro_cliente})")
            else:
                raise SystemExit(f"Embarcador '{filtro_embarcador}' nao encontrado no banco.")

    pedidos = coletar_pedidos(sess, filtro_cliente)

    ts  = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = _RAIZ / "dados" / f"inventario_{ts}.xlsx"
    gerar_excel(pedidos, out, somente_novos=somente_novos)

    # Resumo no terminal
    novos = sum(1 for p in pedidos if p.get("situacao") == "Novo")
    importados = len(pedidos) - novos
    contagem = defaultdict(int)
    lista = [p for p in pedidos if p.get("situacao") == "Novo"] if somente_novos else pedidos
    for p in lista:
        contagem[p.get("fonte","?")] += 1
    logger.info("=" * 50)
    for fonte, qtd in sorted(contagem.items()):
        logger.info(f"  {fonte}: {qtd} pedido(s)")
    logger.info(f"  Novos: {novos} | Ja importados: {importados}")
    logger.info(f"  TOTAL listado: {len(lista)} pedido(s)")
    logger.info("=" * 50)
    logger.info(f"Arquivo salvo em: {out.resolve()}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--embarcador", metavar="ID_OU_NOME", default="")
    parser.add_argument("--somente-novos", action="store_true",
                        help="Exporta apenas pedidos que ainda nao foram importados no VUUPT")
    args = parser.parse_args()
    main(filtro_embarcador=args.embarcador, somente_novos=args.somente_novos)
