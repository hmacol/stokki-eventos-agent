# -*- coding: utf-8 -*-
"""
revisar_complexidade_entrega.py

Compara o nível de dificuldade ATUAL de cada cliente (planilha
BD_CLIENTES.xlsx, ver regras/complexidade_entrega.py) contra o tempo
médio de permanência REAL que o próprio VUUPT já mede por cliente
(campo `stat_average_time_on_site` do customer, em segundos) -- gera
uma lista de DIVERGÊNCIAS pra revisão manual do Hugo.

NÃO edita a planilha sozinho -- só aponta candidatos, com o valor
observado ao lado do nível atual, pra revisão humana decidir.

Validado em produção, 22/08: cruzando 21 dias de rotas reais contra a
planilha, a mediana de stat_average_time_on_site SOBE monotonicamente
do nível 1 ao 4 (5,3min / 14,0min / 16,6min / 32,4min) -- correlação
real, sem precisar de nenhum modelo novo.

Ressalva importante (mesmo achado, 22/08): parte das confirmações de
entrega no VUUPT chega em LOTE (started_at/arrived_at/completed_at do
mesmo motorista, minutos ou segundos entre si, fora de ordem) -- ou
seja, um tempo de permanência muito BAIXO pode ser ruído de
confirmação tardia em vez de entrega genuinamente rápida. Um tempo
muito ALTO não sofre desse viés (confirmação em lote nunca INFLA o
tempo, só zera) -- por isso a seção "subir de nível" é mais confiável
que "descer de nível" neste relatório; a segunda vem marcada como tal.

Execute:
  py -3.11 revisar_complexidade_entrega.py
  py -3.11 revisar_complexidade_entrega.py --dias 60
"""
import argparse
import logging
import statistics
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import yaml

_RAIZ = Path(__file__).parent
sys.path.insert(0, str(_RAIZ))
sys.path.insert(0, str(_RAIZ / "roteirizacao"))
sys.path.insert(0, str(_RAIZ / "painel_agentes"))

(_RAIZ / "dados").mkdir(parents=True, exist_ok=True)

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("revisar_complexidade_entrega")

from rotas_client import listar_rotas
from mapa_util import extrair_servicos_da_rota
from regras.complexidade_entrega import carregar_niveis, classificar_nivel, carregar_ajustes_manuais, nivel_efetivo

CONFIG_PATH = _RAIZ / "config.yaml"

# Mínimo de DIAS DISTINTOS observados (não serviços -- vários serviços
# no mesmo dia/rota podem compartilhar o mesmo horário de confirmação
# em lote, então contá-los separadamente infla artificialmente a
# "amostra" sem reduzir o risco de ruído) pra um cliente entrar na
# lista de divergências. Filtro fraco (não elimina o viés de lote
# quando ele é sistemático, só descarta o caso mais óbvio de "vi esse
# cliente 1x só").
AMOSTRA_MINIMA_DIAS_CLIENTE = 2

# Mínimo de clientes com nível já classificado (fora do padrão) por
# nível de dificuldade pra confiar na MEDIANA daquele nível -- com
# poucos clientes a mediana é só ruído (ex: nível 4 costuma ter poucos
# clientes cadastrados).
AMOSTRA_MINIMA_NIVEL = 15

# Cliente é sinalizado como divergente quando o tempo observado é pelo
# menos essa quantidade de vezes MAIOR (candidato a subir) ou no máximo
# essa fração MENOR (candidato a descer) que a mediana do próprio nível
# atual dele. Limiares largos de propósito -- o objetivo é uma lista
# curta de casos óbvios pro Hugo revisar, não recalcular a planilha
# inteira.
LIMIAR_RAZAO_SUBIR = 2.5
LIMIAR_RAZAO_DESCER = 0.3

COR_HEADER = "141428"
COR_ACENTO = "00C896"
COR_AVISO = "B45309"


def _carregar_config() -> dict:
    with open(CONFIG_PATH, encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def coletar_clientes(token: str, mapa_niveis: dict, ajustes_manuais: dict, dias: int) -> dict:
    """Varre as rotas dos últimos `dias` dias e devolve, por customer_id
    único: nível EFETIVO (ajuste manual > planilha > padrão -- mesma
    prioridade de roteirizacao/criar_rotas_diarias.py, ver
    regras.complexidade_entrega.nivel_efetivo; comparar contra a
    planilha crua geraria falso positivo pra cliente que já foi
    corrigido manualmente pela tela de Planejamento), tempo médio de
    permanência (VUUPT), e quantos DIAS DISTINTOS esse cliente apareceu
    na janela."""
    fim_exclusiva = date.today()
    inicio = fim_exclusiva - timedelta(days=dias)
    filtro = [
        {"field": "start_at", "operator": "gte", "value": inicio.strftime("%Y-%m-%d") + " 00:00:00"},
        {"field": "start_at", "operator": "lt", "value": fim_exclusiva.strftime("%Y-%m-%d") + " 00:00:00"},
    ]
    logger.info(f"Buscando rotas de {inicio} a {fim_exclusiva} (exclusive) na VUUPT...")
    rotas = listar_rotas(token, include=["services.customer"], filtro=filtro)
    logger.info(f"{len(rotas)} rota(s) retornada(s) (inclui canceladas).")

    clientes: dict[int, dict] = {}
    dias_por_cliente: dict[int, set] = defaultdict(set)

    for rota in rotas:
        if rota.get("status") == "canceled":
            continue
        data_rota = str(rota.get("start_at", ""))[:10]
        for s in extrair_servicos_da_rota(rota):
            if s.get("status") == "canceled":
                continue
            cust = s.get("customer") or {}
            cid = cust.get("id")
            if cid is None:
                continue
            if data_rota:
                dias_por_cliente[cid].add(data_rota)
            if cid in clientes:
                continue
            cnpj = cust.get("code") or ""
            doc_digitos = "".join(ch for ch in cnpj if ch.isdigit())
            _, encontrado_planilha, _ = classificar_nivel(cnpj, mapa_niveis)
            tem_ajuste = doc_digitos in ajustes_manuais
            clientes[cid] = {
                "nome": cust.get("name") or "",
                "documento": cnpj,
                "nivel": nivel_efetivo(cnpj, mapa_niveis, ajustes_manuais),
                "fonte_nivel": "Ajuste manual" if tem_ajuste else ("Planilha" if encontrado_planilha else "Padrão"),
                # elegível pra comparação se o nível efetivo vem de uma
                # classificação DE VERDADE (ajuste manual ou planilha),
                # não do padrão por tipo de documento (esse não é uma
                # "decisão" a ser questionada, é só o fallback)
                "encontrado_planilha": encontrado_planilha or tem_ajuste,
                "stat_segundos": cust.get("stat_average_time_on_site"),
            }

    for cid, c in clientes.items():
        c["dias_distintos"] = len(dias_por_cliente.get(cid, ()))

    return clientes


def montar_divergencias(clientes: dict) -> tuple[list[dict], list[dict], dict[int, float], dict[int, int]]:
    """Filtra clientes elegíveis (classificados na planilha, com tempo
    observado > 0 e amostra mínima de dias), calcula a mediana por
    nível, e separa em duas listas (subir / descer) por LIMIAR_RAZAO_*.
    Devolve (subir, descer, medianas_por_nivel_em_minutos, contagem_por_nivel)."""
    elegiveis = [
        c for c in clientes.values()
        if c["encontrado_planilha"]
        and c["stat_segundos"] and c["stat_segundos"] > 0
        and c["dias_distintos"] >= AMOSTRA_MINIMA_DIAS_CLIENTE
    ]

    por_nivel = defaultdict(list)
    for c in elegiveis:
        por_nivel[c["nivel"]].append(c["stat_segundos"] / 60)

    contagens = {nivel: len(valores) for nivel, valores in por_nivel.items()}
    medianas = {
        nivel: statistics.median(valores)
        for nivel, valores in por_nivel.items()
        if len(valores) >= AMOSTRA_MINIMA_NIVEL
    }
    for nivel, valores in por_nivel.items():
        if nivel not in medianas:
            logger.info(f"Nível {nivel}: só {len(valores)} cliente(s) na amostra (< {AMOSTRA_MINIMA_NIVEL}) -- mediana não confiável, ignorado nas comparações.")

    subir, descer = [], []
    for c in elegiveis:
        mediana = medianas.get(c["nivel"])
        if mediana is None or mediana <= 0:
            continue
        tempo_min = c["stat_segundos"] / 60
        razao = tempo_min / mediana
        linha = {**c, "tempo_min": tempo_min, "mediana_nivel_min": mediana, "razao": razao}
        if razao >= LIMIAR_RAZAO_SUBIR:
            subir.append(linha)
        elif razao <= LIMIAR_RAZAO_DESCER:
            descer.append(linha)

    subir.sort(key=lambda c: -c["razao"])
    descer.sort(key=lambda c: c["razao"])
    return subir, descer, medianas, contagens


def _cabecalho(ws, titulo: str, cor_fundo: str, headers: list[tuple]):
    ws.merge_cells(f"A1:{get_column_letter(len(headers))}1")
    cel = ws["A1"]
    cel.value = titulo
    cel.font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    cel.fill = PatternFill("solid", fgColor=cor_fundo)
    cel.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    for col, (_, rotulo, largura) in enumerate(headers, 1):
        c = ws.cell(row=2, column=col, value=rotulo)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=COR_HEADER)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = largura
    ws.row_dimensions[2].height = 22


def _preencher_linhas(ws, linhas: list[dict], headers: list[tuple]):
    for i, linha in enumerate(linhas, 3):
        for col, (campo, _, _) in enumerate(headers, 1):
            c = ws.cell(row=i, column=col, value=linha.get(campo, ""))
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(vertical="center")
    if linhas:
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{len(linhas) + 2}"


def gerar_excel(subir: list[dict], descer: list[dict], medianas: dict, contagens: dict, dias: int, caminho: Path):
    wb = openpyxl.Workbook()

    headers = [
        ("documento", "CNPJ/CPF", 18),
        ("nome", "Cliente", 40),
        ("nivel", "Nível atual", 12),
        ("fonte_nivel", "Fonte do nível atual", 18),
        ("tempo_min", "Tempo médio observado (min)", 24),
        ("mediana_nivel_min", "Mediana do nível atual (min)", 24),
        ("razao", "Razão (observado / mediana)", 22),
        ("dias_distintos", "Dias distintos observados", 20),
    ]
    for linha in subir + descer:
        linha["tempo_min"] = round(linha["tempo_min"], 1)
        linha["mediana_nivel_min"] = round(linha["mediana_nivel_min"], 1)
        linha["razao"] = round(linha["razao"], 2)

    ws1 = wb.active
    ws1.title = "Subir de nivel"
    _cabecalho(
        ws1,
        f"Candidatos a SUBIR de nível -- tempo observado ≥{LIMIAR_RAZAO_SUBIR}x a mediana ({len(subir)})",
        COR_ACENTO, headers,
    )
    _preencher_linhas(ws1, subir, headers)
    nota1 = ws1.cell(row=len(subir) + 4, column=1,
                     value="'Fonte do nível atual' = Planilha: editar o BD_CLIENTES.xlsx já resolve. "
                           "= Ajuste manual: já existe uma correção pontual feita pela tela de Planejamento "
                           "(tem prioridade sobre a planilha) -- editar o Excel NÃO muda nada nesse cliente até "
                           "o ajuste manual também ser atualizado/removido.")
    nota1.font = Font(name="Arial", italic=True, size=9, color=COR_HEADER)

    ws2 = wb.create_sheet("Descer de nivel")
    ws2.merge_cells(f"A1:{get_column_letter(len(headers))}1")
    ws2["A2"] = None
    _cabecalho(
        ws2,
        f"Candidatos a DESCER de nível -- tempo observado ≤{LIMIAR_RAZAO_DESCER}x a mediana ({len(descer)})",
        COR_AVISO, headers,
    )
    aviso = ws2.cell(row=len(descer) + 4, column=1,
                     value="ATENÇÃO: tempo muito baixo pode ser confirmação em lote (motorista marcou várias "
                           "entregas 'concluídas' de uma vez), não entrega genuinamente rápida -- confira "
                           "'dias distintos observados' antes de rebaixar o nível. Ver docstring do script.")
    aviso.font = Font(name="Arial", italic=True, size=9, color=COR_AVISO)
    _preencher_linhas(ws2, descer, headers)

    ws3 = wb.create_sheet("Medianas por nivel")
    ws3["A1"] = "Nível"
    ws3["B1"] = "Mediana observada (min)"
    ws3["C1"] = "Amostra (clientes)"
    for cel in ("A1", "B1", "C1"):
        ws3[cel].font = Font(name="Arial", bold=True, color="FFFFFF")
        ws3[cel].fill = PatternFill("solid", fgColor=COR_HEADER)
    for i, nivel in enumerate(sorted(medianas), 2):
        ws3.cell(row=i, column=1, value=nivel).font = Font(name="Arial", size=10)
        ws3.cell(row=i, column=2, value=round(medianas[nivel], 1)).font = Font(name="Arial", size=10)
        ws3.cell(row=i, column=3, value=contagens.get(nivel, 0)).font = Font(name="Arial", size=10)
    ws3.cell(row=1, column=5, value=f"Janela: últimos {dias} dias -- gerado {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    for col, largura in [("A", 10), ("B", 24), ("C", 16), ("E", 45)]:
        ws3.column_dimensions[col].width = largura

    wb.save(caminho)
    logger.info(f"Excel salvo em: {caminho.resolve()}")


def main(dias: int):
    config = _carregar_config()
    token = config.get("vuupt_api", {}).get("token", "")
    caminho_niveis = config.get("complexidade_entrega", {}).get("planilha", "")

    mapa_niveis = carregar_niveis(caminho_niveis)
    ajustes_manuais = carregar_ajustes_manuais()
    logger.info(f"Planilha de complexidade: {len(mapa_niveis)} documento(s) classificado(s) "
               f"({len(ajustes_manuais)} com ajuste manual, que tem prioridade sobre a planilha).")

    clientes = coletar_clientes(token, mapa_niveis, ajustes_manuais, dias)
    logger.info(f"{len(clientes)} cliente(s) único(s) visto(s) em rotas na janela.")

    subir, descer, medianas, contagens = montar_divergencias(clientes)

    logger.info("=" * 60)
    for nivel in sorted(medianas):
        logger.info(f"  Nível {nivel}: mediana {medianas[nivel]:.1f}min ({contagens[nivel]} cliente(s))")
    logger.info(f"  Candidatos a SUBIR de nível: {len(subir)}")
    logger.info(f"  Candidatos a DESCER de nível (possível ruído de lote): {len(descer)}")
    logger.info("=" * 60)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = _RAIZ / "dados" / f"revisao_complexidade_{ts}.xlsx"
    gerar_excel(subir, descer, medianas, contagens, dias, out)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--dias", type=int, default=45,
                        help="Janela de rotas históricas a considerar (padrão: 45 dias)")
    args = parser.parse_args()
    main(dias=args.dias)
